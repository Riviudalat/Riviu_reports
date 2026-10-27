"""Workbook side of a scan: column mapping, rows to scan, TỔNG rows, the result sheet and result writes."""
import asyncio
import re
from urllib.parse import urlparse

import openpyxl

from riviu.workbook_utils import (
    COLUMN_ALIASES,
    TTBD_RESOLVED_URL_HEADER,
    TTBD_SCAN_STATUS_HEADER,
    TTBD_SOURCE_URL_HEADER,
    clean_text,
    format_display_datetime,
    format_excel_sheet_datetime,
    is_scrapable_tiktok_url,
    normalize_tiktok_url,
    result_sheet_display_name,
    save_workbook_atomic,
    set_cell_literal,
    workbook_data_sheet_names,
    worksheet_ensure_column,
    worksheet_find_column_index,
    worksheet_find_link_column_index,
    worksheet_headers,
    worksheet_partner_column_indexes,
    worksheet_row_partners,
)
from riviu.platforms.tiktok import channels
from riviu.platforms.tiktok.constants import LAST_UPDATE_HEADER, METRIC_HEADERS, METRIC_KEYS
from riviu.platforms.tiktok.lifecycle import finish_pending_task
from riviu.platforms.tiktok.parsing import extract_media_id, metric_or_none, normalize_text
from riviu.platforms.tiktok.statuses import should_clear_stale_metrics


RESULT_SHEET_HEADERS = [
    "Stt",
    "Ngày",
    "Link",
    "Tên Kênh",
    "LƯỢT XEM",
    "TIM",
    "BÌNH LUẬN",
    "LƯỢT LƯU",
    "CHIA SẺ",
    "Đối tác",
    LAST_UPDATE_HEADER,
    TTBD_SCAN_STATUS_HEADER,
    TTBD_RESOLVED_URL_HEADER,
    TTBD_SOURCE_URL_HEADER,
]


def normalize_selected_partners(selected_partner=None, selected_partners=None):
    values = []
    if isinstance(selected_partners, (list, tuple, set)):
        values.extend(selected_partners)
    elif selected_partners:
        values.append(selected_partners)
    if isinstance(selected_partner, (list, tuple, set)):
        values.extend(selected_partner)
    elif selected_partner:
        values.append(selected_partner)

    partners = []
    seen = set()
    for value in values:
        name = clean_text(value)
        key = name.casefold()
        if name and key not in seen:
            partners.append(name)
            seen.add(key)
    return partners


def selected_data_sheet_names(workbook, sheet_name=None):
    target_sheet = clean_text(sheet_name)
    if target_sheet:
        return [target_sheet] if target_sheet in workbook.sheetnames else []
    return workbook_data_sheet_names(workbook)


def build_sheet_contexts(workbook, sheet_name=None):
    contexts = {}
    for sheet_name in selected_data_sheet_names(workbook, sheet_name):
        sheet = workbook[sheet_name]
        contexts[sheet_name] = {
            "worksheet": sheet,
            "columns": ensure_columns(sheet),
        }
    return contexts


INTERNAL_COLUMN_HEADERS = {
    "scan_status": TTBD_SCAN_STATUS_HEADER,
    "resolved_url": TTBD_RESOLVED_URL_HEADER,
    "source_url": TTBD_SOURCE_URL_HEADER,
}
CHANNEL_HEADER = "Tên Kênh"


TIKTOK_SHORT_LINK_HOSTS = {"vm.tiktok.com", "vt.tiktok.com"}


def is_tiktok_post_url(value):
    """True for a TikTok video/photo URL with a media id, or a vm/vt short link.

    Profile URLs (https://www.tiktok.com/@kenh) are not posts, so a "Link kênh"
    column is never mistaken for the link column.
    """
    url = normalize_tiktok_url(value)
    if not url:
        return False
    parsed = urlparse(url)
    host = (parsed.hostname or "").casefold()
    path = parsed.path or ""
    if host in TIKTOK_SHORT_LINK_HOSTS:
        return bool(re.fullmatch(r"/(?:t/)?[A-Za-z0-9_-]+/?", path))
    if re.fullmatch(r"/t/[A-Za-z0-9_-]+/?", path):
        return True
    return bool(re.search(r"/(?:video|photo)/\d+(?:/|$)", path, flags=re.IGNORECASE))


def _find_link_column_by_content(sheet, skip_columns):
    max_column = sheet.max_column or 0
    for row in range(2, min(sheet.max_row or 1, 25) + 1):
        for col in range(1, max_column + 1):
            if col in skip_columns:
                continue
            if is_tiktok_post_url(sheet.cell(row=row, column=col).value):
                return col
    return None


def detect_columns(sheet, headers=None):
    """Map scanner columns by exact (case/diacritic-insensitive) header aliases.

    No positional guesses: a column that is not found stays None and
    ensure_columns appends it, so notes or channel-link columns are never
    mistaken for metric columns. The header row is read once (see `headers`
    in riviu.workbook.columns).
    """
    if headers is None:
        headers = worksheet_headers(sheet)
    column_map = {
        key: worksheet_find_column_index(sheet, (header,), headers=headers)
        for key, header in INTERNAL_COLUMN_HEADERS.items()
    }
    column_map["url"] = worksheet_find_link_column_index(sheet, headers=headers) or _find_link_column_by_content(
        sheet, {index for index in column_map.values() if index}
    )
    column_map["channel"] = worksheet_find_column_index(sheet, COLUMN_ALIASES[CHANNEL_HEADER], headers=headers)
    column_map["date"] = worksheet_find_column_index(sheet, COLUMN_ALIASES["date"], headers=headers)
    column_map["last_update"] = worksheet_find_column_index(sheet, COLUMN_ALIASES[LAST_UPDATE_HEADER], headers=headers)
    for key, header in METRIC_HEADERS.items():
        column_map[key] = worksheet_find_column_index(sheet, (header,), headers=headers)
    return column_map


def ensure_columns(sheet):
    headers = worksheet_headers(sheet)
    column_map = detect_columns(sheet, headers=headers)
    column_map["channel"] = worksheet_ensure_column(sheet, CHANNEL_HEADER, headers=headers)
    for key, header in METRIC_HEADERS.items():
        column_map[key] = worksheet_ensure_column(sheet, header, (header,), headers=headers)
    column_map["last_update"] = worksheet_ensure_column(sheet, LAST_UPDATE_HEADER, headers=headers)
    for key, header in INTERNAL_COLUMN_HEADERS.items():
        column_map[key] = worksheet_ensure_column(sheet, header, (header,), hidden=True, headers=headers)
    return column_map


def trusted_resolved_media_id(url, stored_source_url, stored_resolved_url):
    """Media id a short link resolved to last time, only when it still describes this link.

    The stored id is trusted only when the hidden source-URL cell names the
    current link. Legacy rows without that cell cannot prove the link was not
    replaced, so they get no expectation (a link that carries its own media id
    is already guarded by the redirect check).
    """
    resolved_media_id = extract_media_id(stored_resolved_url)
    if not resolved_media_id:
        return ""
    if stored_source_url:
        return resolved_media_id if stored_source_url.casefold() == url.casefold() else ""
    return resolved_media_id if extract_media_id(url) == resolved_media_id else ""


def collect_rows(workbook, selected_partner=None, selected_partners=None, sheet_name=None):
    selected_names = normalize_selected_partners(selected_partner, selected_partners)
    selected_keys = {partner.casefold() for partner in selected_names}
    rows = []

    for sheet_name in selected_data_sheet_names(workbook, sheet_name):
        worksheet = workbook[sheet_name]
        headers = worksheet_headers(worksheet)
        partner_columns = worksheet_partner_column_indexes(worksheet, headers=headers)
        columns = detect_columns(worksheet, headers=headers)
        url_column = columns["url"]
        if not url_column:
            continue

        for row_index in range(2, worksheet.max_row + 1):
            raw_url = worksheet.cell(row=row_index, column=url_column).value
            if not is_scrapable_tiktok_url(raw_url):
                continue
            url = normalize_tiktok_url(raw_url)
            stored_source_url = normalize_tiktok_url(
                worksheet.cell(row=row_index, column=columns["source_url"]).value
            ) if columns.get("source_url") else ""
            stored_resolved_url = normalize_tiktok_url(
                worksheet.cell(row=row_index, column=columns["resolved_url"]).value
            ) if columns.get("resolved_url") else ""
            expected_media_id = trusted_resolved_media_id(url, stored_source_url, stored_resolved_url)

            partners = []
            if partner_columns:
                partners = worksheet_row_partners(worksheet, row_index, partner_columns)

            if selected_keys:
                if not partners:
                    continue
                if not selected_keys.intersection({partner.casefold() for partner in partners}):
                    continue

            rows.append({
                "sequence": len(rows) + 1,
                "sheet_name": sheet_name,
                "row": row_index,
                "url": url,
                "partners": partners,
                "expected_media_id": expected_media_id,
            })

    return rows


def is_total_row(sheet, row_index, url_column):
    value = clean_text(sheet.cell(row=row_index, column=url_column).value)
    return normalize_text(value) == "TỔNG" or normalize_text(value) == "TONG"


def clear_existing_total_rows(workbook, sheet_name=None):
    """Delete TỔNG rows; return the names of sheets that had one."""
    cleared = []
    for sheet_name in selected_data_sheet_names(workbook, sheet_name):
        sheet = workbook[sheet_name]
        url_column = detect_columns(sheet)["url"]
        if not url_column:
            continue
        for row_index in range(sheet.max_row, 1, -1):
            if is_total_row(sheet, row_index, url_column):
                sheet.delete_rows(row_index, 1)
                if sheet_name not in cleared:
                    cleared.append(sheet_name)
    return cleared


def last_occupied_row(sheet):
    """Last row below the header with any value, or 1.

    max_row/max_column scan every cell on each access, so read them once and
    walk up from the bottom: one pass at most, usually a single row.
    """
    max_column = sheet.max_column
    for row_index in range(sheet.max_row, 1, -1):
        row_values = next(sheet.iter_rows(
            min_row=row_index, max_row=row_index, max_col=max_column, values_only=True,
        ))
        if any(value is not None for value in row_values):
            return row_index
    return 1


def append_sheet_total_rows(workbook, sheet_name=None):
    for sheet_name in selected_data_sheet_names(workbook, sheet_name):
        sheet = workbook[sheet_name]
        columns = detect_columns(sheet)
        url_column = columns.get("url")
        if not url_column:
            continue

        link_rows = [
            row_index for row_index in range(2, sheet.max_row + 1)
            if is_scrapable_tiktok_url(sheet.cell(row=row_index, column=url_column).value)
        ]
        if not link_rows:
            continue
        columns = ensure_columns(sheet)
        total_row = last_occupied_row(sheet) + 1
        sheet.cell(row=total_row, column=url_column).value = "TỔNG"
        # Sum only TikTok rows; contiguous ranges keep large-sheet formulas short.
        ranges = []
        start = end = link_rows[0]
        for row_index in link_rows[1:]:
            if row_index == end + 1:
                end = row_index
            else:
                ranges.append((start, end))
                start = end = row_index
        ranges.append((start, end))
        for metric_key in METRIC_HEADERS:
            column_index = columns.get(metric_key)
            if column_index:
                col_letter = openpyxl.utils.get_column_letter(column_index)
                references = ",".join(f"{col_letter}{start}:{col_letter}{end}" for start, end in ranges)
                sheet.cell(row=total_row, column=column_index).value = f"=SUM({references})"


def build_result_sheet(workbook, rows_to_process, summary_update_time):
    sheet_name = result_sheet_display_name(format_excel_sheet_datetime())
    if sheet_name in workbook.sheetnames:
        base_name = sheet_name[:28]
        suffix = 2
        while f"{base_name}-{suffix}" in workbook.sheetnames:
            suffix += 1
        sheet_name = f"{base_name}-{suffix}"

    worksheet = workbook.create_sheet(title=sheet_name)
    try:
        _fill_result_sheet(workbook, worksheet, rows_to_process, summary_update_time)
    except Exception:
        # Never leave a half-built result tab behind.
        workbook.remove(worksheet)
        raise
    return sheet_name


def _fill_result_sheet(workbook, worksheet, rows_to_process, summary_update_time):
    for column_index, header in enumerate(RESULT_SHEET_HEADERS, start=1):
        worksheet.cell(row=1, column=column_index).value = header

    row_lookup = {(item["sheet_name"], item["row"]) for item in rows_to_process}

    output_row = 2
    for source_sheet_name in workbook_data_sheet_names(workbook):
        if source_sheet_name == worksheet.title:
            continue
        source_sheet = workbook[source_sheet_name]
        columns = ensure_columns(source_sheet)
        url_column = columns.get("url")
        if not url_column:
            continue

        def source_value(key, row_index):
            column_index = columns.get(key)
            return source_sheet.cell(row=row_index, column=column_index).value if column_index else None

        partner_columns = worksheet_partner_column_indexes(source_sheet)
        for source_row_index in range(2, source_sheet.max_row + 1):
            if (source_sheet_name, source_row_index) not in row_lookup:
                continue

            url = clean_text(source_sheet.cell(row=source_row_index, column=url_column).value)
            if not url or is_total_row(source_sheet, source_row_index, url_column):
                continue

            partner_names = worksheet_row_partners(source_sheet, source_row_index, partner_columns) if partner_columns else []
            last_update = (
                clean_text(source_value("last_update", source_row_index))
                if columns.get("last_update") else summary_update_time
            )
            values = [
                output_row - 1,
                source_value("date", source_row_index) if columns.get("date") else "",
                url,
                clean_text(source_value("channel", source_row_index)),
                # Blank stays blank; text like "1.234" is read as a number, never crashes.
                *(metric_or_none(source_value(key, source_row_index)) for key in METRIC_HEADERS),
                "\n".join(partner_names),
                last_update,
                *(clean_text(source_value(key, source_row_index)) for key in INTERNAL_COLUMN_HEADERS),
            ]
            for column_index, value in enumerate(values, start=1):
                set_cell_literal(worksheet.cell(row=output_row, column=column_index), value)
            output_row += 1

    widths = [10, 16, 72, 24, 14, 12, 14, 14, 12, 28, 20]
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[openpyxl.utils.get_column_letter(index)].width = width

    if output_row > 2:
        total_row = output_row
        worksheet.cell(row=total_row, column=3).value = "TỔNG"
        for column_index in range(5, 10):
            letter = openpyxl.utils.get_column_letter(column_index)
            worksheet.cell(row=total_row, column=column_index).value = f"=SUM({letter}2:{letter}{total_row - 1})"

    worksheet.freeze_panes = "A2"
    for column_index in (12, 13, 14):
        worksheet.column_dimensions[openpyxl.utils.get_column_letter(column_index)].hidden = True


def write_result(
    sheet_contexts,
    item,
    data,
    channel_name,
    status,
    *,
    resolved_url="",
    channel_cache=None,
    channel_overrides=None,
    cache_lock=None,
):
    """Write one scan result into its row. Pure workbook work: no network I/O."""
    context = sheet_contexts[item["sheet_name"]]
    sheet = context["worksheet"]
    columns = context["columns"]
    row_index = item["row"]
    update_time = format_display_datetime()

    # Success writes confirmed numbers (unknown ones blank). Terminal "no data"
    # statuses clear stale numbers to blank; other failures keep existing cells.
    if status == "Success" or should_clear_stale_metrics(status):
        for metric_key, data_key in METRIC_KEYS.items():
            column_index = columns.get(metric_key)
            if column_index:
                value = metric_or_none((data or {}).get(data_key)) if status == "Success" else None
                sheet.cell(row=row_index, column=column_index).value = value

    if columns.get("scan_status"):
        set_cell_literal(sheet.cell(row=row_index, column=columns["scan_status"]), clean_text(status))
    if status == "Success" and resolved_url and columns.get("resolved_url"):
        sheet.cell(row=row_index, column=columns["resolved_url"]).value = normalize_tiktok_url(resolved_url)
    if status == "Success" and columns.get("source_url"):
        sheet.cell(row=row_index, column=columns["source_url"]).value = normalize_tiktok_url(item["url"])

    if columns.get("channel"):
        channel_value = channels.channel_name_for_sheet(
            item["url"],
            channel_name,
            resolved_url=resolved_url if status == "Success" else "",
            status=status,
            channel_cache=channel_cache,
            channel_overrides=channel_overrides,
            cache_lock=cache_lock,
        )
        existing = clean_text(sheet.cell(row=row_index, column=columns["channel"]).value)
        if status != "Success" and channels.channel_name_quality(existing, item["url"]) > 0:
            channel_value = existing
        elif channels.channel_name_quality(existing, item["url"]) > channels.channel_name_quality(channel_value, item["url"]):
            channel_value = existing
        set_cell_literal(sheet.cell(row=row_index, column=columns["channel"]), channel_value)

    # Chỉ cập nhật timestamp khi quét thành công — tránh hiểu nhầm đã quét OK.
    if status == "Success" and columns.get("last_update"):
        sheet.cell(row=row_index, column=columns["last_update"]).value = update_time


async def save_workbook(workbook, file_path, websocket_manager=None):
    # Shield the complete save + replace transaction so cancellation cannot
    # close/mutate the workbook while its background writer still uses it.
    # to_thread runs on the loop's default executor, never a run's request pool.
    write_task = asyncio.create_task(asyncio.to_thread(save_workbook_atomic, workbook, file_path))
    try:
        await finish_pending_task(write_task)
        return True, None
    except PermissionError:
        if websocket_manager:
            await websocket_manager.broadcast_log("CẢNH BÁO: File Excel đang mở, không thể ghi đè. Vui lòng đóng file rồi quét lại hoặc chờ lần lưu tiếp theo.")
        return False, "permission"
    except Exception as error:
        if websocket_manager:
            await websocket_manager.broadcast_log(f"CẢNH BÁO: Lỗi lưu file ({str(error)})")
        return False, "other"
