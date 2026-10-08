"""Public Threads post checks, kept separate from TikTok's scraper."""

import asyncio
import contextlib
import json
import os
import re
import time
import urllib.error
import urllib.request
from collections import deque
from decimal import Decimal
from html.parser import HTMLParser
from urllib.parse import parse_qs, urljoin, urlparse

from riviu import proxy_utils
import openpyxl
from riviu.platforms import threads_session
from playwright.async_api import async_playwright

from riviu.platforms.tiktok import close_browser_bounded, finish_pending_task, playwright_session
from riviu.workbook_utils import (
    LAST_UPDATE_COLUMN,
    THREADS_SCAN_STATUS_HEADER,
    clean_text,
    format_display_datetime,
    is_threads_link,
    normalize_threads_url,
    rebuild_summary_sheet,
    save_workbook_atomic,
    summary_sheet_title_for_data_sheet,
    workbook_data_sheet_names,
    worksheet_ensure_column,
    worksheet_find_link_column_index,
    worksheet_partner_column_indexes,
    worksheet_row_partners,
    set_cell_literal,
    write_json_atomic,
)


METRICS = {
    "views": "LƯỢT XEM",
    "likes": "TIM",
    "comments": "BÌNH LUẬN",
    "reposts": "REPOST",
    "shares": "CHIA SẺ",
}
ACTION_TITLES = {"likes": "Like", "comments": "Comment", "reposts": "Repost", "shares": "Share"}
POST_PATH = re.compile(r"^/@([^/]+)/post/([A-Za-z0-9_-]+)/?$", re.IGNORECASE)
SHARE_PATH = re.compile(r"^/share/[A-Za-z0-9_-]+/?$", re.IGNORECASE)
HTTP_HEADERS = {
    "User-Agent": "Mozilla/5.0", "Accept-Language": "en-US,en;q=0.9",
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
}
TERMINAL_ERRORS = {"geo_restricted", "audience", "invalid_post", "redirect", "profile_redirect", "access_denied", "rate_limit", "proxy_auth", "auth_invalid", "auth_checkpoint"}
THREADS_HOSTS = {"threads.com", "www.threads.com", "threads.net", "www.threads.net"}
AUTH_FAILURE_MESSAGE = "Phiên Threads không còn hợp lệ; kiểm tra cookie lại."
BROWSER_POLL_SECONDS = 10
BROWSER_POLL_INTERVAL = 0.25
# Metrics whose absence after HTTP justifies a Hybrid browser fallback. Shares
# are not one: Threads sends reshare_count null to Chromium as well, and the
# rendered Share button then shows no number (docs/threads-scanning.md).
CORE_METRICS = ("views", "likes", "comments", "reposts")
# Hybrid browser fallbacks open at once. Each also takes a route slot, so HTTP
# keeps the other slots while slow Chromium pages load.
BROWSER_FALLBACK_LIMIT = 2


def resolve_threads_proxies(base_dir, proxy_text="", mode="request"):
    text = str(proxy_text).strip() or proxy_utils.load_proxy_list_text(base_dir)
    try:
        for line in text.splitlines():
            if line.strip() and not line.strip().startswith("#") and not proxy_utils.parse_proxy_line(line):
                raise ValueError("Danh sách proxy có dòng không hợp lệ.")
        configs = proxy_utils.parse_proxy_text(text)
    except (ValueError, TypeError):
        raise ValueError("Danh sách proxy có dòng không hợp lệ.") from None
    configs = [config for config in configs if config.get("enabled")]
    if not configs:
        raise ValueError("Bật proxy nhưng chưa có proxy hợp lệ đang bật.")
    if mode != "request" and any(config["type"] == "socks5" and config.get("username") for config in configs):
        raise ValueError("Chromium không hỗ trợ SOCKS5 có tài khoản/mật khẩu. Dùng Request hoặc proxy HTTP cho Browser/Hybrid.")
    return configs


def failure(message, kind):
    return {"channel": "", "metrics": empty_metrics(), "error": message, "error_kind": kind}


def _count(value):
    return value if isinstance(value, int) and not isinstance(value, bool) and value >= 0 else None


def _embedded_post_metrics(content, identity, *, with_sources=False, documents=None):
    pending = list(threads_session.json_script_documents(content) if documents is None else documents)
    views = set()
    posts = []
    view_queries = []
    while pending:
        node = pending.pop()
        if isinstance(node, dict):
            pending.extend(value for value in node.values() if isinstance(value, (dict, list)))
        elif isinstance(node, list):
            # Read only typed payloads, never count fields in captions or
            # recommendations. The target query's direct media has its own identity.
            if len(node) >= 3 and node[0] == "BarcelonaLoggedOutExpansionGating" and isinstance(node[2], dict):
                value = _count(node[2].get("view_counts"))
                if value is not None:
                    views.add(value)
            if (len(node) >= 4 and isinstance(node[0], str)
                    and re.fullmatch(r"RelayPrefetchedStreamCache(?:@[0-9a-fA-F]{32})?", node[0]) and node[1] == "next"
                    and isinstance(node[3], list) and len(node[3]) >= 2):
                query, payload = node[3][:2]
                if isinstance(query, str) and isinstance(payload, dict):
                    bbox = payload.get("__bbox")
                    result = bbox.get("result") if isinstance(bbox, dict) else None
                    data = result.get("data") if isinstance(result, dict) else None
                    media = data.get("media") if isinstance(data, dict) else None
                    if isinstance(media, dict):
                        if query.startswith("adp_BarcelonaPostPageTargetQueryRelayPreloader_"):
                            user = media.get("user")
                            username = user.get("username") if isinstance(user, dict) else None
                            if isinstance(username, str) and identity and (username.casefold(), media.get("code")) == identity:
                                posts.append(media)
                        elif query.startswith("adp_BarcelonaPostViewCountQueryRelayPreloader_"):
                            view_queries.append(media)
            pending.extend(value for value in node if isinstance(value, (dict, list)))
    metrics = empty_metrics()
    metrics["views"] = next(iter(views)) if len(views) == 1 else None
    sources = {"views": "embedded_exact"} if metrics["views"] is not None else {}
    target_ids = {post["id"] for post in posts if isinstance(post.get("id"), str) and post["id"]}
    matched_views = [media for media in view_queries if len(target_ids) == 1 and media.get("id") in target_ids]
    if matched_views:
        exact = set()
        for media in matched_views:
            info = media.get("text_post_app_info")
            exact.add(_count(info.get("impression_count")) if isinstance(info, dict) else None)
        metrics["views"] = next(iter(exact)) if len(exact) == 1 else None
        sources["views"] = "view_query_exact" if metrics["views"] is not None else "view_query_unknown"
    candidates = {key: set() for key in ACTION_TITLES}
    for media in posts:
        # A hidden or unspecified like-count visibility flag is not proof of zero.
        candidates["likes"].add(_count(media.get("like_count"))
                                if media.get("like_and_view_counts_disabled") is False else None)
        info = media.get("text_post_app_info")
        info = info if isinstance(info, dict) else {}
        for key, field in (("comments", "direct_reply_count"), ("reposts", "repost_count"), ("shares", "reshare_count")):
            candidates[key].add(_count(info.get(field)))
    for key, values in candidates.items():
        if len(values) == 1:
            metrics[key] = next(iter(values))
    return (metrics, bool(posts), sources) if with_sources else (metrics, bool(posts))


def post_identity(url):
    normalized = normalize_threads_url(url)
    if not is_threads_link(normalized):
        return None
    match = POST_PATH.fullmatch(urlparse(normalized).path)
    return (match.group(1).casefold(), match.group(2)) if match else None


def same_threads_post(requested_url, final_url):
    requested = post_identity(requested_url)
    final = post_identity(final_url)
    if requested is not None:
        return requested == final
    return bool(is_threads_link(requested_url) and SHARE_PATH.fullmatch(urlparse(normalize_threads_url(requested_url)).path) and final)


def empty_metrics():
    return {key: None for key in METRICS}


def parse_action_count(value):
    text = clean_text(value).replace("\xa0", "").replace(",", "").strip()
    match = re.fullmatch(r"(\d+(?:\.\d+)?)\s*([KMB]?)", text, re.IGNORECASE)
    if not match:
        return None
    scale = {"": 1, "K": 1000, "M": 1000000, "B": 1000000000}[match.group(2).upper()]
    # Decimal keeps displayed counts exact: float("32.3") * 1000 == 32299.999...
    return int(Decimal(match.group(1)) * scale)


class _PermalinkHeaderViews(HTMLParser):
    """Keep a bounded element tree for the permalink's semantic column header."""
    VOID = {"area", "base", "br", "col", "embed", "hr", "img", "input", "link", "meta", "param", "source", "track", "wbr"}

    def __init__(self):
        super().__init__()
        # Flat parent indexes avoid retaining a cycle for every element per poll.
        self.root = {"tag": "document", "attrs": {}, "children": [], "text": [], "parent": None, "index": 0}
        self.nodes = [self.root]
        self.stack = [self.root]
        self.anchors = []
        self.count = 0

    def handle_starttag(self, tag, attrs):
        self.count += 1
        if self.count > 100000:
            raise ValueError("Header tree limit")
        node = {"tag": tag, "attrs": dict(attrs), "children": [], "text": [], "parent": self.stack[-1]["index"], "index": len(self.nodes)}
        self.nodes.append(node)
        self.stack[-1]["children"].append(node)
        if tag == "a" and node["attrs"].get("aria-label") == "Column title":
            self.anchors.append(node)
        if tag not in self.VOID:
            self.stack.append(node)

    def handle_startendtag(self, tag, attrs):
        self.handle_starttag(tag, attrs)
        if tag not in self.VOID:
            self.handle_endtag(tag)

    def handle_endtag(self, tag):
        for index in range(len(self.stack) - 1, 0, -1):
            if self.stack[index]["tag"] == tag:
                del self.stack[index:]
                break

    def handle_data(self, data):
        if self.stack[-1]["tag"] not in {"script", "style"}:
            self.stack[-1]["text"].append(data)


def _header_nodes(root):
    pending = [root]
    while pending:
        node = pending.pop()
        yield node
        pending.extend(node["children"])


def _header_text(root):
    return " ".join(part for node in _header_nodes(root) for part in node["text"])


def _permalink_header_views(content, final_url):
    if len(content) > 8 * 1024 * 1024:
        return None
    parser = _PermalinkHeaderViews()
    try:
        parser.feed(content)
    except (ValueError, RecursionError):
        return None
    wanted = post_identity(final_url)
    values = set()
    for anchor in parser.anchors:
        href = anchor["attrs"].get("href")
        if (not wanted or not isinstance(href, str) or post_identity(urljoin(final_url, href)) != wanted
                or not any(node["tag"] == "h1" for node in _header_nodes(anchor))):
            continue
        ancestor = parser.nodes[anchor["parent"]] if anchor["parent"] is not None else None
        caption_scope = False
        while ancestor is not None:
            if ancestor["tag"] in {"article", "p", "button"} or ancestor["attrs"].get("role") == "button":
                caption_scope = True
                break
            ancestor = parser.nodes[ancestor["parent"]] if ancestor["parent"] is not None else None
        if caption_scope:
            continue
        child = anchor
        for _ in range(3):
            parent = parser.nodes[child["parent"]] if child["parent"] is not None else None
            if parent is None or parent["tag"] in {"document", "html", "body", "main", "article", "p"}:
                break
            nodes = list(_header_nodes(parent))
            # Never cross into caption/action/reply branches or another column.
            if any(node["tag"] in {"article", "p", "button", "svg", "script", "style"}
                   or node["attrs"].get("role") == "button" for node in nodes):
                break
            links = [node for node in nodes if node["tag"] == "a"]
            if any(node is not anchor for node in links):
                break
            found = set()
            for sibling in parent["children"]:
                if sibling is child:
                    continue
                text = clean_text(_header_text(sibling)).replace("\xa0", " ")
                match = re.fullmatch(r"(\d[\d,.]*\s*[KMB]?)\s+views?", text, re.IGNORECASE)
                if match:
                    count = parse_action_count(match.group(1))
                    if count is not None:
                        found.add(count)
            if found:
                values.update(found)
                break
            child = parent
    return next(iter(values)) if len(values) == 1 else None


def _geo_restricted_route(content, final_url, documents=None):
    pending = []
    for document in threads_session.json_script_documents(content) if documents is None else documents:
        required = document.get("require") if isinstance(document, dict) else None
        if isinstance(required, list):
            pending.extend(module for module in required if isinstance(module, list))
    while pending:
        node = pending.pop()
        if isinstance(node, list):
            if (len(node) >= 4 and node[0] == "ScheduledServerJS" and node[1] == "handle"
                    and isinstance(node[3], list) and node[3] and isinstance(node[3][0], dict)):
                bbox = node[3][0].get("__bbox")
                required = bbox.get("require", []) if isinstance(bbox, dict) else []
                required = required if isinstance(required, list) else []
                roots = [module for module in required if isinstance(module, list) and len(module) >= 4
                         and module[0] == "CometPlatformRootClient" and module[1] == "initialize"
                         and isinstance(module[3], list) and module[3] and isinstance(module[3][0], dict)]
            else:
                roots = []
            for root_module in roots:
                info = root_module[3][0].get("initialRouteInfo")
                route = info.get("route") if isinstance(info, dict) else None
                if isinstance(route, dict):
                    route_url = route.get("url")
                    root = route.get("rootView")
                    resource = root.get("resource") if isinstance(root, dict) else None
                    if (isinstance(route_url, str) and same_threads_post(final_url, urljoin(final_url, route_url))
                            and route.get("canonicalRouteName") == "comet.barcelonawebloggedout.BarcelonaGeoBlockRoute"
                            and isinstance(resource, dict) and resource.get("__dr") == "BarcelonaGeoBlockedErrorRoot.react"):
                        return True
    return False


def _location_failure(requested_url, final_url):
    """Classify a URL that is not the requested post; its page numbers are never used."""
    if "invalid_post" in parse_qs(urlparse(final_url).query).get("error", []):
        return failure("Link Threads không khả dụng (invalid_post)", "invalid_post")
    if not same_threads_post(requested_url, final_url):
        parsed = urlparse(final_url)
        if (parsed.hostname or "").lower() in THREADS_HOSTS and re.fullmatch(r"/@[^/]+/?", parsed.path):
            return failure("Link bài Threads không khả dụng (có thể bị ẩn hoặc đã xóa)", "profile_redirect")
        return failure("Bài Threads chuyển hướng sang URL khác", "redirect")
    return None


def parse_threads_http(content, requested_url, final_url, documents=None):
    """documents: threads_session.json_script_documents(content), if already parsed."""
    moved = _location_failure(requested_url, final_url)
    if moved:
        return moved
    if documents is None:
        documents = threads_session.json_script_documents(content)
    if _geo_restricted_route(content, final_url, documents):
        return failure("Bài Threads bị giới hạn theo vị trí hiện tại", "geo_restricted")
    metrics, target, sources = _embedded_post_metrics(content, post_identity(final_url), with_sources=True, documents=documents)
    if metrics["views"] is None and sources.get("views") != "view_query_unknown":
        metrics["views"] = _permalink_header_views(content, final_url)
        if metrics["views"] is not None:
            sources["views"] = "header_display"
    result = {"channel": post_identity(final_url)[0], "metrics": metrics, "error": "", "metricSources": sources}
    # Gate evidence is considered only when no verified target payload exists.
    # Both error-page sentences are required; captions/error strings alone are not gates.
    if not target:
        text = re.sub(r"<script\b[^>]*>.*?</script>|<style\b[^>]*>.*?</style>", "", content, flags=re.DOTALL | re.IGNORECASE)
        text = clean_text(re.sub(r"<[^>]+>", " ", text))
        if ("This content isn't available to everyone" in text
                and "It can't be seen by certain audiences." in text):
            return failure("Bài Threads giới hạn người xem với phiên hiện tại", "audience")
    return result


def missing_core_metrics(result):
    return any(result["metrics"].get(key) is None for key in CORE_METRICS)


def missing_metric_labels(result):
    return [label for key, label in METRICS.items() if result["metrics"].get(key) is None]


def result_status(result):
    if result.get("error"):
        return f"Error: {result['error']}"
    # Zero is a confirmed count. Scan success is separate from metric completeness.
    if any(value is not None for value in result["metrics"].values()):
        return "Success"
    return "Ẩn số liệu: Threads không trả chỉ số"


def merge_results(first, second):
    if (second.get("error_kind") == "auth_unknown"
            or first.get("error_kind") == "auth_unknown" and second.get("auth_verified") is not True):
        return failure("Chưa xác minh được phiên Threads; kiểm tra cookie lại.", "auth_unknown")
    for result in (second, first):
        if result.get("error_kind") in {"auth_invalid", "auth_checkpoint"}:
            return failure("Phiên Threads không còn hợp lệ; kiểm tra cookie lại.", result["error_kind"])
    for result in (second, first):
        if result.get("error_kind") == "geo_restricted":
            return failure("Bài Threads bị giới hạn theo vị trí hiện tại", "geo_restricted")
    merged = {"channel": first.get("channel") or second.get("channel", ""), "metrics": dict(first["metrics"]), "error": ""}
    invalidate_display_views = (first.get("metricSources", {}).get("views") == "header_display"
                                and second.get("metricSources", {}).get("views") == "view_query_unknown")
    if invalidate_display_views:
        merged["metrics"]["views"] = None
    for key, value in second["metrics"].items():
        # Browser is a fallback for missing values, not a reason to replace
        # a value already extracted from the HTTP post response.
        first_source = first.get("metricSources", {}).get(key)
        second_source = second.get("metricSources", {}).get(key)
        if key == "views" and first_source == "view_query_unknown" and second_source != "view_query_exact":
            continue
        if value is not None and (merged["metrics"].get(key) is None
                                  or key == "views" and first_source == "header_display" and second_source == "view_query_exact"):
            merged["metrics"][key] = value
    if not any(value is not None for value in merged["metrics"].values()):
        merged["error"] = second.get("error") or first.get("error", "")
        if merged["error"]:
            merged["error_kind"] = second.get("error_kind") or first.get("error_kind", "payload")
    elif second.get("error") or second.get("diagnostic"):
        merged["diagnostic"] = second.get("error") or second["diagnostic"]
    sources = dict(first.get("metricSources", {}))
    if invalidate_display_views:
        sources["views"] = "view_query_unknown"
    for key, source in second.get("metricSources", {}).items():
        if (first["metrics"].get(key) is None and (sources.get(key) != "view_query_unknown" or source == "view_query_exact")
                or key == "views" and sources.get(key) == "header_display" and source == "view_query_exact"):
            sources[key] = source
    if sources:
        merged["metricSources"] = sources
    if first.get("auth_verified") is True or second.get("auth_verified") is True:
        merged["auth_verified"] = True
    return merged


def _analyze_snapshot(content, requested_url, final_url, *, session=False):
    """Tokenize one HTML snapshot once; derive session evidence and post metrics from it.

    Returns (auth_state, parsed). auth_state is None without a cookie session;
    parsed is None when the snapshot positively rejects the session. CPU-bound:
    browser callers run it through asyncio.to_thread.
    """
    try:
        documents = threads_session.json_script_documents(content)
    except Exception:
        if not session:
            raise
        return "unknown", None
    state = threads_session.auth_evidence(content, final_url, documents) if session else None
    if state in {"invalid", "checkpoint"}:
        return state, None
    return state, parse_threads_http(content, requested_url, final_url, documents)


def fetch_threads_http(url, timeout=20, proxy_config=None, session_cookies=None):
    request = urllib.request.Request(url, headers=HTTP_HEADERS)
    session = session_cookies is not None
    try:
        options = {"cookiejar": threads_session.cookie_jar(session_cookies)} if session else {}
        with proxy_utils.urlopen_with_config(request, proxy_config, timeout=timeout, **options) as response:
            content = response.read().decode("utf-8", errors="replace")
            final_url = response.geturl()
            state, result = _analyze_snapshot(content, url, final_url, session=session)
            if state in {"invalid", "checkpoint"}:
                return failure(AUTH_FAILURE_MESSAGE, "auth_" + state)
            if session and state != "valid":
                return failure("HTTP chưa xác minh được phiên đăng nhập Threads.", "auth_unknown")
            if session:
                result["auth_verified"] = True
            return result
    except threads_session.SessionError as error:
        # The cookie redirect guard refused a hop off HTTPS Threads.
        if error.state == "redirect":
            return failure("Bài Threads chuyển hướng sang URL khác", "redirect")
        return failure("Chưa xác minh được phiên Threads; kiểm tra cookie lại.", "auth_unknown")
    except urllib.error.HTTPError as error:
        kind = {403: "access_denied", 429: "rate_limit", 407: "proxy_auth"}.get(error.code, "http")
        return failure(f"HTTP {error.code}", kind)
    except (urllib.error.URLError, TimeoutError, OSError):
        return failure("Không kết nối được proxy HTTP" if proxy_config else "HTTP: kết nối thất bại hoặc hết thời gian", "transport")
    except Exception:
        return failure("HTTP: không đọc được phản hồi", "http")


# Reads the single rendered card that links to the requested post. Only explicit
# action/views controls count; captions, replies and recommendations never do.
EXTRACT_ACTIONS_JS = r"""({username, shortcode}) => {
    const identity = href => {
        try {
            const url = new URL(href, location.href);
            if (!/^(www\.)?threads\.(com|net)$/i.test(url.hostname)) return null;
            const match = url.pathname.match(/^\/@([^/]+)\/post\/([^/]+)\/?$/);
            return match ? [match[1].toLowerCase(), match[2]] : null;
        } catch { return null; }
    };
    const matches = href => {
        const id = identity(href);
        return id && id[0] === username && id[1] === shortcode;
    };
    const anchors = [...document.querySelectorAll('a[href]')].filter(a => matches(a.href));
    const roots = new Set();
    for (const anchor of anchors) {
        let node = anchor.parentElement;
        while (node && node !== document.body && node !== document.documentElement) {
            const postLinks = [...node.querySelectorAll('a[href]')].filter(a => identity(a.href));
            // Crossing into replies/recommendations makes a container ambiguous.
            if (postLinks.some(a => !matches(a.href))) break;
            if (node.querySelector('svg[title="Like"]') && node.querySelector('svg[title="Comment"]')) {
                roots.add(node);
                break;
            }
            node = node.parentElement;
        }
    }
    // A sticky page title may link to the same post. Prefer its inner
    // card; multiple disjoint cards remain ambiguous.
    const cards = [...roots].filter(root => ![...roots].some(other => other !== root && root.contains(other)));
    if (cards.length !== 1) return null;
    const root = cards[0];
    const values = {};
    for (const [key, title] of Object.entries({likes:'Like', comments:'Comment', reposts:'Repost', shares:'Share'})) {
        const actions = root.querySelectorAll(`svg[title="${title}"]`);
        if (actions.length !== 1) continue;
        const button = actions[0].closest('[role="button"],button');
        if (button && root.contains(button)) values[key] = button.innerText.trim();
    }
    // Only an explicit views control is evidence; never regex caption/body text.
    const viewControls = [...root.querySelectorAll('[role="button"],button,[aria-label]')]
        .filter(node => /\bviews?\b/i.test(node.getAttribute('aria-label') || '')
            || node.querySelector('svg[title="Views"],svg[title="View"]'));
    const viewCounts = new Set();
    for (const control of viewControls) {
        const text = (control.innerText || '').trim();
        const match = text.match(/^(\d[\d,.]*\s*[KMB]?)\s*(?:views?)?$/i);
        if (match) viewCounts.add(match[1]);
    }
    if (viewCounts.size === 1) values.views = [...viewCounts][0];
    return values;
}"""


async def _open_browser_page(browser, proxy_config, session_cookies):
    """Return (owned_context, page, navigation_guard); anonymous direct pages share the browser."""
    if not proxy_config and session_cookies is None:
        return None, await browser.new_page(), None
    options = {"proxy": proxy_utils.playwright_proxy_settings(proxy_config)} if proxy_config else {}
    if session_cookies is not None:
        options["service_workers"] = "block"
    context = await browser.new_context(**options)
    try:
        guard = None
        if session_cookies is not None:
            await context.add_cookies(threads_session.cookies_to_playwright(session_cookies))
            guard = await threads_session.install_session_navigation_guard(context)
        return context, await context.new_page(), guard
    except BaseException:
        await context.close()
        raise


def _dom_result(identity, extracted):
    return {"channel": identity[0], "metrics": {**empty_metrics(), **{key: parse_action_count(value) for key, value in extracted.items()}}, "error": ""}


async def fetch_threads_browser(browser, url, proxy_config=None, session_cookies=None):
    session = session_cookies is not None
    navigation = {"channel": "", "metrics": empty_metrics(), "error": ""}
    verified = False
    context = page = None
    try:
        context, page, guard = await _open_browser_page(browser, proxy_config, session_cookies)
        if session:
            response = await threads_session.navigate_session_page(page, url, guard, timeout=30000)
            # The guard pins the checked HTTP final URL; page.url may still move client-side.
            final_url = guard.get("final_url") or page.url
        else:
            response = await page.goto(url, wait_until="domcontentloaded", timeout=30000)
            final_url = page.url
        if response and response.status in (403, 429, 407):
            return failure(f"Browser: HTTP {response.status}", {403: "access_denied", 429: "rate_limit", 407: "proxy_auth"}[response.status])
        if response:
            state, parsed = await asyncio.to_thread(_analyze_snapshot, await response.text(), url, final_url, session=session)
        else:
            state, parsed = (threads_session.auth_evidence("", final_url) if session else None), None
        if state in {"invalid", "checkpoint"}:
            return failure(AUTH_FAILURE_MESSAGE, "auth_" + state)
        verified = state == "valid"
        if parsed is not None and (not session or verified):
            navigation = parsed
            if verified:
                navigation["auth_verified"] = True
        if navigation.get("error_kind") in TERMINAL_ERRORS:
            return navigation

        deadline = time.monotonic() + BROWSER_POLL_SECONDS
        previous = None
        while True:
            content = await page.content()
            snapshot_url = page.url
            if session:
                # A client-side move (Home, ?error=invalid_post) is never the requested post.
                moved = _location_failure(final_url, snapshot_url) if snapshot_url != final_url else None
                if moved:
                    return moved
                snapshot_url = final_url
            if (content, snapshot_url) != previous:
                # Polls repeat until hydration settles; parse only changed snapshots, off the event loop.
                previous = (content, snapshot_url)
                state, embedded = await asyncio.to_thread(_analyze_snapshot, content, url, snapshot_url, session=session)
                # Once the navigation proved the login, only positive evidence fails it:
                # React may drop the JSON scripts that carried the viewer.
                if state in {"invalid", "checkpoint"}:
                    return failure(AUTH_FAILURE_MESSAGE, "auth_" + state)
                verified = verified or state == "valid"
                if embedded is not None and (not session or verified):
                    if session:
                        embedded["auth_verified"] = True
                    if embedded.get("error_kind") in TERMINAL_ERRORS:
                        return embedded
                    navigation = merge_results(navigation, embedded)
            identity = post_identity(snapshot_url)
            extracted = await page.evaluate(EXTRACT_ACTIONS_JS, {"username": identity[0], "shortcode": identity[1]})
            # Wait for core evidence, including a header that may hydrate later.
            if extracted is not None and (not session or verified):
                navigation = merge_results(navigation, _dom_result(identity, extracted))
            if not missing_core_metrics(navigation) or time.monotonic() >= deadline:
                break
            await asyncio.sleep(BROWSER_POLL_INTERVAL)
        if session and not verified:
            return failure("Browser chưa xác minh được phiên đăng nhập Threads.", "auth_unknown")
        result = failure("Không xác minh được card Threads cần quét", "payload") if extracted is None else _dom_result(identity, extracted)
        if session:
            result["auth_verified"] = True
        # A verified navigation payload remains valid if React removes its script.
        return merge_results(navigation, result)
    except Exception:
        if session and not verified:
            return failure("Chưa xác minh được phiên Threads; kiểm tra cookie lại.", "auth_unknown")
        error = failure("Browser: không kết nối hoặc đọc được trang qua proxy" if proxy_config else "Browser: không kết nối hoặc đọc được trang", "transport")
        return merge_results(navigation, error)
    finally:
        if context:
            await context.close()
        elif page:
            await page.close()


def write_threads_result(sheet, row_index, result):
    # Shared, diacritic-insensitive header matching: never append a duplicate the preview ignores.
    channel = result.get("channel")
    channel_column = worksheet_ensure_column(sheet, "Tên Kênh")
    if channel:
        set_cell_literal(sheet.cell(row=row_index, column=channel_column), channel)
    for key, header in METRICS.items():
        value = result["metrics"].get(key)
        if value is not None:
            sheet.cell(row=row_index, column=worksheet_ensure_column(sheet, header)).value = value
        elif key == "views" and result.get("metricSources", {}).get("views") == "view_query_unknown":
            # Explicit exact-null/conflict invalidates stale display-derived views.
            sheet.cell(row=row_index, column=worksheet_ensure_column(sheet, header)).value = None
    worksheet_ensure_column(sheet, "REPOST")
    status = result_status(result)
    sheet.cell(row=row_index, column=worksheet_ensure_column(sheet, THREADS_SCAN_STATUS_HEADER, hidden=True)).value = status
    if channel or any(value is not None for value in result["metrics"].values()):
        sheet.cell(row=row_index, column=worksheet_ensure_column(sheet, LAST_UPDATE_COLUMN)).value = format_display_datetime()
    return status


def collect_threads_rows(workbook, sheet_name="", selected_partners=None, invalid_links=None):
    """Collect scannable post rows; Threads-host rows that are not post links go to invalid_links."""
    selected = {clean_text(value).casefold() for value in selected_partners or [] if clean_text(value)}
    names = [sheet_name] if sheet_name else workbook_data_sheet_names(workbook)
    rows = []
    for name in names:
        if name not in workbook.sheetnames:
            continue
        sheet = workbook[name]
        link_column = worksheet_find_link_column_index(sheet)
        if not link_column:
            continue
        partner_columns = worksheet_partner_column_indexes(sheet)
        for index in range(2, sheet.max_row + 1):
            url = normalize_threads_url(sheet.cell(row=index, column=link_column).value)
            partners = worksheet_row_partners(sheet, index, partner_columns) if partner_columns else []
            if selected and not selected.intersection(value.casefold() for value in partners):
                continue
            if not is_threads_link(url):
                parsed = urlparse(url)
                host = (parsed.hostname or "").casefold()
                if invalid_links is not None and host in THREADS_HOSTS:
                    kind = "profile" if re.fullmatch(r"/@[^/]+/?", parsed.path) else "malformed"
                    # Host + path only: never echo query strings or userinfo into logs.
                    invalid_links.append({"sheet_name": name, "row": index, "kind": kind, "display": f"{host}{parsed.path}"})
                continue
            rows.append({"sheet_name": name, "row": index, "url": url, "partners": partners})
    return rows


def append_threads_diagnostic(base_dir, entry):
    directory = os.path.join(base_dir, "data")
    os.makedirs(directory, exist_ok=True)
    path = os.path.join(directory, "threads_scan_history.json")
    history = []
    try:
        with open(path, encoding="utf-8") as stream:
            payload = json.load(stream)
        if isinstance(payload, dict) and isinstance(payload.get("history"), list):
            history = payload["history"]
    except (OSError, ValueError):
        pass
    write_json_atomic(path, {"history": [entry, *history[:49]]})


async def run_threads_scraper(file_path, websocket_manager=None, worker_count=3, selected_partners=None,
                              sheet_name="", mode="hybrid", base_dir="", use_proxy=False, proxy_text="", session_cookies=None, proxy_configs=None):
    started = time.monotonic()
    diagnostic = {
        "startedAt": format_display_datetime(), "platform": "threads", "fileLabel": os.path.basename(file_path),
        "scanSheet": sheet_name, "mode": mode, "authSelected": session_cookies is not None, "authState": "unchecked" if session_cookies is not None else "anonymous",
        "proxyEnabled": bool(use_proxy), "proxyCount": 0, "proxyTypes": [],
        "workers": worker_count, "totalUrls": 0, "totalRows": 0, "processedUrls": 0, "processedRows": 0,
        "success": 0, "partial": 0, "error": 0, "errorKinds": {}, "phase": "failed",
    }
    try:
        await _run_threads_scraper(file_path, websocket_manager, worker_count, selected_partners,
                                   sheet_name, mode, base_dir, use_proxy, proxy_text, diagnostics=diagnostic,
                                   session_cookies=session_cookies, proxy_configs=proxy_configs)
        diagnostic["phase"] = "completed"
    except asyncio.CancelledError:
        diagnostic["phase"] = "cancelled"
        raise
    except Exception as error:
        diagnostic["failureType"] = type(error).__name__
        raise
    finally:
        diagnostic["finishedAt"] = format_display_datetime()
        diagnostic["durationSeconds"] = round(time.monotonic() - started, 2)
        try:
            await finish_pending_task(asyncio.create_task(asyncio.to_thread(
                append_threads_diagnostic, base_dir or os.path.dirname(os.path.abspath(file_path)), diagnostic)))
        except Exception:
            # Diagnostic I/O must not replace a scan/save failure or a cancellation.
            if websocket_manager:
                try:
                    await websocket_manager.broadcast_log("Không lưu được chẩn đoán phiên Threads; dữ liệu quét không thay đổi.")
                except Exception:
                    pass


async def _preflight_session(session_cookies, proxies, diagnostics):
    """Verify the cookie on the one route the whole run will use; never fall back to direct."""
    check = await threads_session.verify_cookies(session_cookies, proxies[0] if proxies else None)
    if diagnostics is not None:
        diagnostics["authState"] = check["state"]
    if check["state"] != "valid":
        reported = {"expired", "invalid", "checkpoint", "proxy_unsupported"}
        raise threads_session.SessionError(check["state"] if check["state"] in reported else "unknown")


async def _request_once(url, config, session_cookies):
    options = {"proxy_config": config} if config else {}
    if session_cookies is not None:
        options["session_cookies"] = session_cookies
    return await finish_pending_task(asyncio.create_task(asyncio.to_thread(fetch_threads_http, url, **options)))


async def _request_check(url, config, session_cookies):
    """HTTP with at most one retry for a transport error or missing core metrics."""
    request_result = await _request_once(url, config, session_cookies)
    if request_result.get("error_kind") not in TERMINAL_ERRORS:
        if (missing_core_metrics(request_result) and not request_result.get("error")) or request_result.get("error_kind") == "transport":
            await asyncio.sleep(0.5)
            retried = await _request_once(url, config, session_cookies)
            request_result = merge_results(request_result, retried)
            if retried.get("error_kind") in TERMINAL_ERRORS:
                request_result["error_kind"] = retried["error_kind"]
    return request_result


def needs_browser_fallback(request_result):
    """Hybrid opens Chromium only for a missing core metric; never for shares alone."""
    return request_result.get("error_kind") not in TERMINAL_ERRORS and missing_core_metrics(request_result)


async def _browser_check(url, config, session_cookies, get_browser):
    """One Chromium load on the URL's own route (never another proxy or direct)."""
    options = {"proxy_config": config} if config else {}
    if session_cookies is not None:
        options["session_cookies"] = session_cookies
    try:
        active_browser = await get_browser()
    except Exception:
        return failure("Không khởi chạy được Chromium cho Threads", "browser_start")
    return await fetch_threads_browser(active_browser, url, **options)


class _RouteSlots:
    """Page loads in flight for one scan (HTTP and Chromium share the limit).

    Item tasks queue here in sheet order. A Hybrid browser fallback is served
    before queued HTTP work so it never waits behind the rest of the sheet; the
    caller caps fallbacks below the slot count, so HTTP keeps draining.
    """

    def __init__(self, size):
        self._free = size
        self._waiters = {True: deque(), False: deque()}

    @contextlib.asynccontextmanager
    async def hold(self, *, fallback=False):
        if self._free and not any(self._waiters.values()):
            self._free -= 1
        else:
            future = asyncio.get_running_loop().create_future()
            queue = self._waiters[fallback]
            queue.append(future)
            try:
                await future
            except asyncio.CancelledError:
                if future in queue:
                    queue.remove(future)
                elif not future.cancelled():
                    self._release()  # Granted, then cancelled before it ran.
                raise
        try:
            yield
        finally:
            self._release()

    def _release(self):
        self._free += 1
        for queue in (self._waiters[True], self._waiters[False]):
            while self._free and queue:
                future = queue.popleft()
                if not future.done():
                    self._free -= 1
                    future.set_result(None)


async def _broadcast_result(websocket_manager, item, result, status, source, processed, total):
    diagnostic = f"; {result['diagnostic']}" if result.get("diagnostic") else ""
    missing = missing_metric_labels(result)
    completeness = f"; nguồn chưa trả: {', '.join(missing)}" if status == "Success" and missing else ""
    await websocket_manager.broadcast_log(f"Threads {processed}/{total} [{source}] {status}{completeness}{diagnostic}: {item['url']}")
    primary = item["rows"][0]
    await websocket_manager.broadcast_data({
        "id": processed, "url": item["url"], "sheetName": primary["sheet_name"],
        "channelName": result["channel"], "views": result["metrics"]["views"],
        "likes": result["metrics"]["likes"], "comments": result["metrics"]["comments"],
        "saves": result["metrics"]["reposts"], "shares": result["metrics"]["shares"],
        "status": status, "missingMetrics": missing, "metricSources": result.get("metricSources", {}), "worker": source, "singlePartner": len(primary["partners"]) == 1,
    })


async def _rebuild_threads_summary(workbook, sheet_name, selected_partners, websocket_manager):
    """Refresh the sheet's Threads summary tab (never the TikTok one); failure is only a warning."""
    try:
        count = rebuild_summary_sheet(
            workbook,
            summary_update_time=format_display_datetime(),
            selected_partners=selected_partners,
            data_sheet_name=sheet_name,
            platform="threads",
        )
    except Exception as error:
        if websocket_manager:
            await websocket_manager.broadcast_log(f"CẢNH BÁO: Không cập nhật được sheet Tổng kết ({error})")
        return
    if websocket_manager:
        title = summary_sheet_title_for_data_sheet(sheet_name, "threads")
        await websocket_manager.broadcast_log(f"Đã cập nhật {title} ({count} đối tác).")


async def _run_threads_scraper(file_path, websocket_manager=None, worker_count=3, selected_partners=None,
                               sheet_name="", mode="hybrid", base_dir="", use_proxy=False, proxy_text="", diagnostics=None,
                               session_cookies=None, proxy_configs=None):
    proxies = (list(proxy_configs) if proxy_configs is not None else resolve_threads_proxies(base_dir or os.path.dirname(os.path.abspath(file_path)), proxy_text, mode)) if use_proxy else []
    if use_proxy and not proxies:
        raise ValueError("Bật proxy nhưng chưa có proxy hợp lệ.")
    if session_cookies is not None:
        await _preflight_session(session_cookies, proxies, diagnostics)
    workbook = openpyxl.load_workbook(file_path)
    invalid_links = []
    rows = collect_threads_rows(workbook, sheet_name=sheet_name, selected_partners=selected_partners, invalid_links=invalid_links)
    skipped_profiles = sum(1 for invalid in invalid_links if invalid["kind"] == "profile")
    buckets = {}
    for row in rows:
        key = post_identity(row["url"]) or ("share", row["url"].casefold())
        buckets.setdefault(key, {"url": row["url"], "rows": []})["rows"].append(row)
    items = list(buckets.values())
    total = len(items)
    workers = max(1, min(int(worker_count or 3), 5 if mode != "request" else 10))
    if diagnostics is not None:
        diagnostics.update(totalUrls=total, totalRows=len(rows), skippedProfiles=skipped_profiles,
                           skippedInvalidLinks=len(invalid_links) - skipped_profiles, workers=workers,
                           proxyCount=len(proxies), proxyTypes=sorted({config["type"] for config in proxies}))
    if websocket_manager:
        duplicates = [
            {"id": index, "url": item["url"], "locations": [
                {"sheetName": row["sheet_name"], "row": row["row"]} for row in item["rows"]
            ]}
            for index, item in enumerate(items, 1) if len(item["rows"]) > 1
        ]
        await websocket_manager.broadcast_duplicates({"items": duplicates, "duplicateRowCount": len(rows) - total})
        skipped = f"; bỏ qua {len(invalid_links)} link Threads không hợp lệ" if invalid_links else ""
        await websocket_manager.broadcast_log(f"Threads: {len(rows)} dòng, {total} URL{skipped}; chế độ {mode}, {workers} luồng; {len(proxies)} proxy.")
        for invalid in invalid_links:
            if invalid["kind"] == "profile":
                await websocket_manager.broadcast_log(f"Bỏ qua {invalid['sheet_name']} dòng {invalid['row']}: link hồ sơ Threads, cần link /post/<mã bài>.")
            else:
                await websocket_manager.broadcast_log(
                    f"Bỏ qua {invalid['sheet_name']} dòng {invalid['row']}: link Threads sai định dạng ({invalid['display']}), "
                    "cần dạng threads.com/@tên/post/<mã bài> hoặc threads.com/share/<mã>; sửa link rồi quét lại.")
    if not total:
        if websocket_manager:
            await websocket_manager.broadcast_log("Không tìm thấy link Threads trong sheet hoặc đối tác đã chọn.")
            await websocket_manager.broadcast_status({"total": 0, "processed": 0, "success": 0, "hidden": 0, "error": 0, "done": True})
        workbook.close()
        return

    slots = _RouteSlots(workers)
    # Hybrid fallbacks never take every slot while HTTP work is queued.
    fallback_slots = asyncio.Semaphore(max(1, min(BROWSER_FALLBACK_LIMIT, workers - 1)))
    processed = success = hidden = errors = 0
    dirty = False
    completed = False
    started = time.monotonic()

    async with playwright_session(async_playwright) as playwright:
        browser = None
        browser_lock = asyncio.Lock()

        async def get_browser():
            nonlocal browser
            async with browser_lock:
                if browser is None:
                    # Windows Chromium needs launch-time proxy support for context proxies.
                    browser = await playwright.chromium.launch(headless=True, proxy={"server": "per-context"} if proxies else None)
            return browser

        async def save_pending():
            nonlocal dirty
            await asyncio.to_thread(save_workbook_atomic, workbook, file_path)
            dirty = False

        async def check(index, item):
            # A cookie session keeps one route for the whole run; anonymous runs round-robin.
            # HTTP, its retry and any browser fallback of a URL use the same route.
            config = proxies[0 if session_cookies is not None else index % len(proxies)] if proxies else None
            if mode == "browser":
                async with slots.hold():
                    return item, await _browser_check(item["url"], config, session_cookies, get_browser), "Browser"
            async with slots.hold():
                request_result = await _request_check(item["url"], config, session_cookies)
            if mode == "request" or not needs_browser_fallback(request_result):
                return item, request_result, "Request"
            # The HTTP slot is released first: a slow page never holds it.
            async with fallback_slots, slots.hold(fallback=True):
                browser_result = await _browser_check(item["url"], config, session_cookies, get_browser)
            return item, merge_results(request_result, browser_result), "Hybrid"

        tasks = [asyncio.create_task(check(index, item)) for index, item in enumerate(items)]
        closing = None
        try:
            for task in asyncio.as_completed(tasks):
                item, result, source = await task
                if session_cookies is not None and result.get("error_kind") in {"auth_invalid", "auth_checkpoint", "auth_unknown"}:
                    if diagnostics is not None:
                        diagnostics["authState"] = result["error_kind"].removeprefix("auth_")
                    raise threads_session.SessionError(result["error_kind"].removeprefix("auth_"))
                processed += 1
                status = result_status(result)
                if status == "Success":
                    success += 1
                elif status.startswith("Error:"):
                    errors += 1
                else:
                    hidden += 1
                for row in item["rows"]:
                    write_threads_result(workbook[row["sheet_name"]], row["row"], result)
                    dirty = True
                if diagnostics is not None:
                    diagnostics.update(processedUrls=processed, success=success, partial=hidden, error=errors)
                    diagnostics["processedRows"] += len(item["rows"])
                    if result.get("error"):
                        kind = result.get("error_kind", "unknown")
                        diagnostics["errorKinds"][kind] = diagnostics["errorKinds"].get(kind, 0) + 1
                if websocket_manager:
                    await _broadcast_result(websocket_manager, item, result, status, source, processed, total)
                    await websocket_manager.broadcast_status({
                        "total": total, "processed": processed, "success": success,
                        "hidden": hidden, "error": errors, "workers": workers,
                        "mode": "partner" if selected_partners else "full", "done": False,
                        "rate": round(processed / max(time.monotonic() - started, 0.1) * 60, 1),
                    })
                if processed % 5 == 0:
                    await finish_pending_task(asyncio.create_task(save_pending()))
            # Every page load is done: Chromium closes while the workbook is finalized.
            closing = asyncio.ensure_future(close_browser_bounded(browser))
            if sheet_name:
                await _rebuild_threads_summary(workbook, sheet_name, selected_partners, websocket_manager)
                dirty = True
            completed = True
        finally:
            async def cleanup():
                for task in tasks:
                    if not task.done():
                        task.cancel()
                try:
                    await asyncio.gather(*tasks, return_exceptions=True)
                    if dirty:
                        await save_pending()
                finally:
                    try:
                        # Bounded: a slow Chromium exit is force-killed, and
                        # playwright_session waits only briefly for the driver.
                        await (closing if closing is not None else close_browser_bounded(browser))
                    finally:
                        workbook.close()

            await finish_pending_task(asyncio.create_task(cleanup()))
    if completed and websocket_manager:
        await websocket_manager.broadcast_status({
            "total": total, "processed": processed, "success": success,
            "hidden": hidden, "error": errors, "workers": workers,
            "mode": "partner" if selected_partners else "full", "done": True,
        })
