"""The per-sheet, per-platform "Tổng kết" tab written into the workbook after a scan."""

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from riviu.workbook.columns import (
    set_cell_literal,
    worksheet_find_column_index,
    worksheet_find_last_update_column_index,
    worksheet_find_link_column_index,
    worksheet_partner_column_indexes,
)
from riviu.workbook.constants import LAST_UPDATE_COLUMN
from riviu.workbook.dates import display_datetime_sort_key
from riviu.workbook.naming import summary_sheet_title_for_data_sheet, workbook_data_sheet_names
from riviu.workbook.partners import normalize_selected_partner_keys, partner_dedup_key, worksheet_row_partners
from riviu.workbook.platforms import PLATFORMS, get_platform
from riviu.workbook.text import clean_text, metric_total


def read_existing_summary_updates(worksheet):
    updates = {}
    if not worksheet or (worksheet.max_row or 0) < 2:
        return updates
    partner_column = worksheet_find_column_index(worksheet, ["ĐỐI TÁC"])
    update_column = worksheet_find_last_update_column_index(worksheet)
    if not partner_column or not update_column:
        return updates
    for row_index in range(2, (worksheet.max_row or 0) + 1):
        partner = clean_text(worksheet.cell(row=row_index, column=partner_column).value)
        update_value = clean_text(worksheet.cell(row=row_index, column=update_column).value)
        if partner and update_value:
            updates[partner_dedup_key(partner)] = update_value
    return updates


def build_partner_summary_rows(
    workbook,
    summary_update_time="",
    selected_partner=None,
    selected_partners=None,
    previous_updates=None,
    data_sheet_name=None,
    platform="tiktok",
):
    # Only this platform's links count, so a mixed sheet yields one summary per platform.
    spec = get_platform(platform)
    selected_keys = normalize_selected_partner_keys(selected_partner, selected_partners)
    previous_updates = previous_updates or {}
    summary = {}
    # Rows repeat a few update stamps, and each row compares two; parse each stamp once per rebuild.
    sort_keys = {}

    def update_sort_key(value):
        if not isinstance(value, str):
            return display_datetime_sort_key(value)
        key = sort_keys.get(value)
        if key is None:
            key = sort_keys[value] = display_datetime_sort_key(value)
        return key

    data_sheets = workbook_data_sheet_names(workbook)
    requested_sheet = clean_text(data_sheet_name)
    if requested_sheet:
        data_sheets = [requested_sheet] if requested_sheet in workbook.sheetnames else []
    for sheet_name in data_sheets:
        worksheet = workbook[sheet_name]
        link_column = worksheet_find_link_column_index(worksheet)
        if not link_column:
            continue

        partner_columns = worksheet_partner_column_indexes(worksheet)
        if not partner_columns:
            continue

        metric_columns = {
            header: worksheet_find_column_index(worksheet, [header])
            for header in spec.metric_columns
        }
        last_update_column = worksheet_find_last_update_column_index(worksheet)

        for row_index in range(2, (worksheet.max_row or 0) + 1):
            link = clean_text(worksheet.cell(row=row_index, column=link_column).value)
            if not spec.is_link(link):
                continue

            partners = worksheet_row_partners(worksheet, row_index, partner_columns)
            if not partners:
                continue

            for partner in partners:
                # Same grouping as the partner picker, so "Shop A"/"shop a" is one summary row.
                key = partner_dedup_key(partner)
                bucket = summary.setdefault(
                    key,
                    {
                        "ĐỐI TÁC": partner,
                        "TỔNG LINK": 0,
                        # Blank until a known value arrives: a partner whose links are all unknown is not 0.
                        **{f"TỔNG {metric}": "" for metric in spec.metric_columns},
                        LAST_UPDATE_COLUMN: previous_updates.get(key, ""),
                    },
                )
                bucket["TỔNG LINK"] += 1
                if last_update_column:
                    update_value = clean_text(worksheet.cell(row=row_index, column=last_update_column).value)
                    if update_value and update_sort_key(update_value) > update_sort_key(bucket[LAST_UPDATE_COLUMN]):
                        bucket[LAST_UPDATE_COLUMN] = update_value
                for metric, column_index in metric_columns.items():
                    if column_index:
                        value = worksheet.cell(row=row_index, column=column_index).value
                        bucket[f"TỔNG {metric}"] = metric_total([bucket[f"TỔNG {metric}"], value])

    result = []
    if summary_update_time:
        for row in summary.values():
            if not selected_keys or row["ĐỐI TÁC"].casefold() in selected_keys:
                row[LAST_UPDATE_COLUMN] = summary_update_time

    for index, row in enumerate(sorted(summary.values(), key=lambda value: value["ĐỐI TÁC"].casefold()), start=1):
        row["Stt"] = index
        result.append(row)
    return result


def rebuild_summary_sheet(
    workbook,
    summary_update_time="",
    selected_partner=None,
    selected_partners=None,
    data_sheet_name=None,
    platform="tiktok",
):
    spec = get_platform(platform)
    source_sheet = clean_text(data_sheet_name)
    if not source_sheet or source_sheet not in workbook.sheetnames:
        return 0

    summary_sheet = summary_sheet_title_for_data_sheet(source_sheet, spec.key)
    existing_worksheet = workbook[summary_sheet] if summary_sheet in workbook.sheetnames else None
    previous_updates = read_existing_summary_updates(existing_worksheet)
    # Sit right after the data sheet, behind the summaries of platforms listed earlier in the
    # registry, so rebuilding one platform's tab never reorders the other's.
    anchor = source_sheet
    for key in PLATFORMS:
        if key == spec.key:
            break
        sibling = summary_sheet_title_for_data_sheet(source_sheet, key)
        if sibling in workbook.sheetnames:
            anchor = sibling
    insert_index = workbook.sheetnames.index(anchor) + 1
    if summary_sheet in workbook.sheetnames:
        worksheet = workbook[summary_sheet]
        worksheet.delete_rows(1, max(worksheet.max_row or 1, 1))
        if worksheet.title != summary_sheet:
            worksheet.title = summary_sheet
        current_index = workbook.sheetnames.index(summary_sheet)
        if current_index != insert_index:
            workbook.move_sheet(worksheet, offset=insert_index - current_index)
    else:
        worksheet = workbook.create_sheet(summary_sheet, insert_index)

    rows = build_partner_summary_rows(
        workbook,
        summary_update_time=summary_update_time,
        selected_partner=selected_partner,
        selected_partners=selected_partners,
        previous_updates=previous_updates,
        data_sheet_name=source_sheet,
        platform=spec.key,
    )
    summary_columns = spec.summary_columns
    header_fill = PatternFill("solid", fgColor="0B5ED7")
    for column_index, header in enumerate(summary_columns, start=1):
        cell = worksheet.cell(row=1, column=column_index, value=header)
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")

    for row_index, row in enumerate(rows, start=2):
        for column_index, header in enumerate(summary_columns, start=1):
            value = row.get(header, "")
            cell = set_cell_literal(worksheet.cell(row=row_index, column=column_index), None if value == "" else value)
            if header == "ĐỐI TÁC":
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            elif header == LAST_UPDATE_COLUMN:
                cell.alignment = Alignment(horizontal="center")
            else:
                cell.number_format = "#,##0"
                cell.alignment = Alignment(horizontal="right")

    last_row = max(len(rows) + 1, 1)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:{get_column_letter(len(summary_columns))}{last_row}"
    widths = [8, 38, 12, 16, 12, 16, 16, 14, 20]
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width

    return len(rows)
