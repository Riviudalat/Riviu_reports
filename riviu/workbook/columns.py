"""Column lookup shared by pandas frames and openpyxl worksheets.

Scanners, previews and reports must find the same column for the same header, so both kinds of
sheet go through the same header-list helpers here.
"""

from copy import copy

from openpyxl.utils import column_index_from_string, get_column_letter

from riviu.workbook.constants import (
    COLUMN_ALIASES,
    LAST_UPDATE_COLUMN,
    LEGACY_SCAN_STATUS_HEADERS,
    SCAN_STATUS_ALIASES,
    SCAN_STATUS_HEADER,
    TTBD_INTERNAL_HEADERS,
)
from riviu.workbook.links import is_post_url
from riviu.workbook.text import normalize_header, normalize_key


LINK_COLUMN_SAMPLE_ROWS = 200
# "Link kênh"/"Profile URL" point at an account, never at the post a scan or report needs.
NON_POST_LINK_HEADER_MARKERS = ("kenh", "channel", "profile", "tai khoan")


def set_cell_literal(cell, value):
    """Write untrusted input literally, without altering its displayed text.

    Generated formulas must use normal cell assignment instead of this helper.
    """
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"
    return cell


def find_header_position(headers, aliases):
    """0-based position of the first header matching an alias (case/diacritic-insensitive)."""
    alias_set = {normalize_key(alias) for alias in aliases}
    for position, header in enumerate(headers):
        if normalize_key(header) in alias_set:
            return position
    return None


def partner_column_positions(headers):
    """0-based positions of the "Đối tác" column and the blank/"Unnamed" columns right after it."""
    start = None
    for position, header in enumerate(headers):
        if normalize_key(header).startswith("doi tac"):
            start = position
            break
    if start is None:
        return []
    positions = [start]
    for position in range(start + 1, len(headers)):
        header = headers[position]
        raw_key = normalize_header(header)
        if normalize_key(header).startswith("doi tac") or not raw_key or raw_key.startswith("unnamed:"):
            positions.append(position)
        else:
            break
    return positions


def pick_link_column(headers, sample_values):
    """Index of the post-link column among headers matched by substring ('link'/'url').

    `sample_values(index)` yields cell values of that column. Account/profile link columns are
    skipped unless they are the only ones holding post URLs; a column with post URLs wins.
    """
    internal_keys = {normalize_key(header) for header in (*TTBD_INTERNAL_HEADERS, *LEGACY_SCAN_STATUS_HEADERS)}
    post_candidates = []
    account_candidates = []
    for index, header in enumerate(headers):
        key = normalize_key(header)
        if key in internal_keys or not ("link" in key or "url" in key):
            continue
        if any(marker in key for marker in NON_POST_LINK_HEADER_MARKERS):
            account_candidates.append(index)
        else:
            post_candidates.append(index)
    for index in post_candidates + account_candidates:
        if any(is_post_url(value) for value in sample_values(index)):
            return index
    return post_candidates[0] if post_candidates else None


# pandas frames: columns are returned by name.

def find_column_name(frame, aliases):
    columns = list(frame.columns)
    position = find_header_position(columns, aliases)
    return None if position is None else columns[position]


def find_link_column_name(frame):
    column = find_column_name(frame, COLUMN_ALIASES["link"])
    if column:
        return column
    columns = list(frame.columns)
    index = pick_link_column(
        columns,
        lambda position: frame.iloc[:LINK_COLUMN_SAMPLE_ROWS, position].tolist(),
    )
    return None if index is None else columns[index]


def dataframe_partner_columns(frame):
    columns = list(frame.columns)
    return [columns[position] for position in partner_column_positions(columns)]


# openpyxl worksheets: columns are returned as 1-based indexes.
#
# worksheet.max_column walks every cell of the sheet (about 0.1 s for a 1,900-row
# month sheet). A caller that runs several lookups on one sheet reads the header
# row once with worksheet_headers() and passes that list as `headers`;
# worksheet_ensure_column appends any header it adds to that list, so it stays
# equal to a fresh worksheet_headers() as long as only these helpers write row 1.

def worksheet_headers(worksheet):
    max_column = worksheet.max_column or 0
    if max_column < 1:
        return []
    return [worksheet.cell(row=1, column=index).value for index in range(1, max_column + 1)]


def worksheet_find_column_index(worksheet, aliases, headers=None):
    position = find_header_position(worksheet_headers(worksheet) if headers is None else headers, aliases)
    return None if position is None else position + 1


def worksheet_ensure_column(worksheet, header, aliases=None, *, hidden=False, headers=None):
    """Find a column the same way previews/reports do, appending it only when truly absent."""
    index = worksheet_find_column_index(worksheet, aliases or COLUMN_ALIASES.get(header) or (header,), headers=headers)
    if index is None:
        if headers is None:
            index = (worksheet.max_column or 0) + 1
        else:
            index = len(headers) + 1
            headers.append(header)
        worksheet.cell(row=1, column=index).value = header
    if hidden:
        worksheet.column_dimensions[get_column_letter(index)].hidden = True
    return index


def worksheet_show_column(worksheet, index):
    """Unhide one column without unhiding its neighbours.

    Excel saves adjacent columns with the same width/visibility as one <col min max> range, and
    openpyxl keeps that range under its first letter, so the range is split around `index` first.
    """
    dimensions = worksheet.column_dimensions
    for key, dimension in list(dimensions.items()):
        low = dimension.min or column_index_from_string(key)
        high = dimension.max or low
        if low == high or not low <= index <= high:
            continue
        del dimensions[key]
        for start, end in ((low, index - 1), (index, index), (index + 1, high)):
            if start > end:
                continue
            piece = copy(dimension)
            piece.index = get_column_letter(start)
            piece.min = start
            piece.max = end
            dimensions[piece.index] = piece
    letter = get_column_letter(index)
    if letter in dimensions:
        dimensions[letter].hidden = False


def worksheet_ensure_scan_status_column(worksheet, headers=None):
    """Index of the visible "TRẠNG THÁI QUÉT" column, migrating a legacy hidden one in place.

    Older scans kept the status in a hidden "__TTBD_SCAN_STATUS"/"__THREADS_SCAN_STATUS" column.
    That column is renamed and unhidden where it stands, so its values are kept and no duplicate
    status column is appended.
    """
    index = worksheet_ensure_column(worksheet, SCAN_STATUS_HEADER, SCAN_STATUS_ALIASES, headers=headers)
    cell = worksheet.cell(row=1, column=index)
    if cell.value != SCAN_STATUS_HEADER:
        cell.value = SCAN_STATUS_HEADER
        if headers is not None and index <= len(headers):
            headers[index - 1] = SCAN_STATUS_HEADER
    worksheet_show_column(worksheet, index)
    return index


def worksheet_find_link_column_index(worksheet, headers=None):
    if headers is None:
        headers = worksheet_headers(worksheet)
    column_index = worksheet_find_column_index(worksheet, COLUMN_ALIASES["link"], headers=headers)
    if column_index:
        return column_index

    last_row = min(worksheet.max_row or 0, LINK_COLUMN_SAMPLE_ROWS + 1)
    index = pick_link_column(
        headers,
        lambda position: (worksheet.cell(row=row, column=position + 1).value for row in range(2, last_row + 1)),
    )
    return None if index is None else index + 1


def worksheet_find_last_update_column_index(worksheet):
    return worksheet_find_column_index(worksheet, COLUMN_ALIASES[LAST_UPDATE_COLUMN])


def worksheet_partner_column_indexes(worksheet, headers=None):
    headers = worksheet_headers(worksheet) if headers is None else headers
    return [position + 1 for position in partner_column_positions(headers)]
