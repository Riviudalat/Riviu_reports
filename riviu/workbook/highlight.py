"""Row fills on the scanned data sheet: single-partner links (orange) and active videos (blue)."""

from openpyxl.styles import PatternFill

from riviu.workbook.columns import (
    worksheet_find_column_index,
    worksheet_find_link_column_index,
    worksheet_headers,
    worksheet_partner_column_indexes,
)
from riviu.workbook.constants import (
    LEGACY_PHOTO_LINK_FILL_COLOR,
    SCAN_STATUS_ALIASES,
    SINGLE_LINK_FILL_COLOR,
    TTBD_RESOLVED_URL_HEADER,
    TTBD_SOURCE_URL_HEADER,
    VIDEO_LINK_FILL_COLOR,
)
from riviu.workbook.links import is_scrapable_tiktok_url, should_highlight_video_link
from riviu.workbook.partners import worksheet_row_partners
from riviu.workbook.text import clean_text


def _cell_has_fill(cell, colors):
    """True when the cell has a solid fill of `colors` (one color or several; the fill is read once)."""
    if isinstance(colors, str):
        colors = (colors,)
    fill = cell.fill
    if fill is None or fill.fill_type != "solid":
        return False
    rgb = str(fill.fgColor.rgb or "").upper()
    return any(rgb.endswith(color) for color in colors)


def _video_highlight_columns(worksheet, headers):
    """Column indexes read by should_highlight_video_link, keyed by its keyword names."""
    def find(header):
        return worksheet_find_column_index(worksheet, [header], headers=headers)

    return {
        "likes": (find("TIM"), 0),
        "shares": (find("CHIA SẺ"), 0),
        "scan_status": (worksheet_find_column_index(worksheet, SCAN_STATUS_ALIASES, headers=headers), ""),
        "resolved_url": (find(TTBD_RESOLVED_URL_HEADER), ""),
        "resolved_source_url": (find(TTBD_SOURCE_URL_HEADER), ""),
    }


def _is_active_video_row(worksheet, row_index, link, video_columns):
    values = {
        name: worksheet.cell(row=row_index, column=column).value if column else default
        for name, (column, default) in video_columns.items()
    }
    return should_highlight_video_link(link, **values)


_CLEAR_FILL = PatternFill(fill_type=None)


def _repaint_row(worksheet, row_index, max_column, highlighted, fill, stale_colors):
    for column_index in range(1, max_column + 1):
        cell = worksheet.cell(row=row_index, column=column_index)
        if highlighted:
            cell.fill = fill
        elif _cell_has_fill(cell, stale_colors):
            cell.fill = _CLEAR_FILL


def _link_sheet(workbook, data_sheet_name):
    """(worksheet, link column index, header row) of the named sheet, or None without the sheet or a link column.

    The header row is read once: each read walks every cell for max_column.
    """
    source_sheet = clean_text(data_sheet_name)
    if not source_sheet or source_sheet not in workbook.sheetnames:
        return None
    worksheet = workbook[source_sheet]
    headers = worksheet_headers(worksheet)
    link_column = worksheet_find_link_column_index(worksheet, headers=headers)
    if not link_column:
        return None
    return worksheet, link_column, headers


def highlight_single_partner_link_rows(workbook, data_sheet_name):
    """Bôi cam link chỉ có 1 đối tác, trừ video đủ điều kiện bôi xanh."""
    sheet = _link_sheet(workbook, data_sheet_name)
    if sheet is None:
        return 0
    worksheet, link_column, headers = sheet

    partner_columns = worksheet_partner_column_indexes(worksheet, headers=headers)
    if not partner_columns:
        return 0
    video_columns = _video_highlight_columns(worksheet, headers)

    max_row = worksheet.max_row or 0
    max_column = len(headers)
    if max_row < 2 or max_column < 1:
        return 0

    highlight_fill = PatternFill("solid", fgColor=SINGLE_LINK_FILL_COLOR)
    highlighted_count = 0
    for row_index in range(2, max_row + 1):
        link = clean_text(worksheet.cell(row=row_index, column=link_column).value)
        should_highlight = False
        if is_scrapable_tiktok_url(link):
            partners = worksheet_row_partners(worksheet, row_index, partner_columns)
            is_active_video = _is_active_video_row(worksheet, row_index, link, video_columns)
            should_highlight = len(partners) == 1 and not is_active_video
        if should_highlight:
            highlighted_count += 1
        _repaint_row(worksheet, row_index, max_column, should_highlight, highlight_fill, (SINGLE_LINK_FILL_COLOR,))

    return highlighted_count


def highlight_video_link_rows(workbook, data_sheet_name):
    """Bôi xanh video có TIM hoặc CHIA SẺ lớn hơn 0."""
    sheet = _link_sheet(workbook, data_sheet_name)
    if sheet is None:
        return 0
    worksheet, link_column, headers = sheet
    video_columns = _video_highlight_columns(worksheet, headers)

    max_row = worksheet.max_row or 0
    max_column = len(headers)
    if max_row < 2 or max_column < 1:
        return 0

    video_fill = PatternFill("solid", fgColor=VIDEO_LINK_FILL_COLOR)
    video_count = 0
    for row_index in range(2, max_row + 1):
        link = clean_text(worksheet.cell(row=row_index, column=link_column).value)
        is_video = bool(link and _is_active_video_row(worksheet, row_index, link, video_columns))
        if is_video:
            video_count += 1
        _repaint_row(
            worksheet, row_index, max_column, is_video, video_fill,
            (VIDEO_LINK_FILL_COLOR, LEGACY_PHOTO_LINK_FILL_COLOR),
        )

    return video_count
