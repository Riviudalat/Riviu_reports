"""Sheet preview and summary dashboard payloads for the UI."""

import re

import pandas as pd
from openpyxl.utils import get_column_letter

from riviu.workbook.channels import display_channel_name_from_file
from riviu.workbook.columns import dataframe_partner_columns, find_column_name, find_link_column_name
from riviu.workbook.constants import (
    COLUMN_ALIASES,
    LAST_UPDATE_COLUMN,
    TTBD_INTERNAL_HEADERS,
    TTBD_RESOLVED_URL_HEADER,
    TTBD_SCAN_STATUS_HEADER,
    TTBD_SOURCE_URL_HEADER,
)
from riviu.workbook.dates import preview_datetime_text
from riviu.workbook.links import is_tiktok_link, should_highlight_video_link
from riviu.workbook.naming import (
    data_sheet_name_for_summary_title,
    is_summary_sheet_name,
    resolve_summary_sheet,
    summary_sheet_title_for_data_sheet,
)
from riviu.workbook.partners import extract_row_partners
from riviu.workbook.platforms import PLATFORMS, get_platform
from riviu.workbook.report_rows import fill_missing_dates_from_previous
from riviu.workbook.snapshot import open_workbook_snapshot
from riviu.workbook.text import clean_text, is_total_label, metric_number, metric_total, normalize_header, normalize_key


def clean_preview_value(value):
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    text = clean_text(value)
    if text.endswith(" 00:00:00"):
        return text[:10]
    return text


def _total_label_positions(frame, link_column):
    return [index for index, value in enumerate(frame[link_column].tolist()) if is_total_label(value)]


def fill_preview_total_row(frame, link_column, metric_columns, *, platform=None):
    if not link_column or link_column not in frame.columns:
        return frame

    total_positions = _total_label_positions(frame, link_column)
    if not total_positions:
        return frame

    total_position = total_positions[-1]
    if total_position <= 0:
        return frame

    source_frame = frame.iloc[:total_position]
    spec = PLATFORMS.get(platform or "")
    if spec:
        source_frame = frame[frame[link_column].map(spec.is_link)]
    for column in metric_columns:
        if not column or column not in frame.columns:
            continue
        total_value = metric_total(source_frame[column])
        # pandas 3 string columns reject assigning numeric totals. Display rows
        # are intentionally mixed, while the original workbook stays untouched.
        frame[column] = frame[column].astype(object)
        frame.iat[total_position, frame.columns.get_loc(column)] = total_value
    return frame


def evaluate_sequence_formula(formula, row, column_letter, known):
    """Value of a numbering formula such as =ROW()-1 or =A2+1 at `row`; None if not supported.

    `known` maps sheet rows of the same column to their numbers.
    """
    text = re.sub(r"\s+", "", str(formula or "").lstrip("=")).upper()
    match = re.fullmatch(
        r"(?:ROW\(\)|\$?" + re.escape(column_letter.upper()) + r"\$?(\d+))(?:([+-])(\d+))?",
        text,
    )
    if not match:
        return None
    base = row if match.group(1) is None else known.get(int(match.group(1)))
    if base is None:
        return None
    offset = int(match.group(3) or 0)
    return base - offset if match.group(2) == "-" else base + offset


def _sequence_number(value):
    if isinstance(value, (int, float)) and not pd.isna(value) and float(value).is_integer():
        return int(value)
    text = clean_text(value)
    return int(text) if re.fullmatch(r"\d+", text) else None


def fill_formula_sequence_numbers(frame, snapshot, sheet_name):
    """Show STT numbers that come from formulas without a cached value.

    Workbooks saved by openpyxl (every scan) keep formulas but drop their cached values, so a
    "=A2+1" or "=ROW()-1" STT column reads as blank. Cached values stay as they are.
    """
    column = find_column_name(frame, ["STT"])
    if column is None:
        return frame
    values = frame[column].tolist()
    if all(clean_text(value) for value in values):
        return frame
    position = frame.columns.get_loc(column)
    column_letter = get_column_letter(position + 1)
    formulas = snapshot.column_formulas(sheet_name, column_letter)
    if not formulas:
        return frame
    known = {}
    filled = {}
    for index, value in enumerate(values):
        row = index + 2  # pandas takes the header from sheet row 1 and keeps blank rows
        if clean_text(value):
            number = _sequence_number(value)
            if number is not None:
                known[row] = number
            continue
        number = evaluate_sequence_formula(formulas.get(row), row, column_letter, known) if row in formulas else None
        if number is not None:
            known[row] = number
            filled[index] = number
    if filled:
        frame[column] = frame[column].astype(object)
        for index, number in filled.items():
            frame.iat[index, position] = number
    return frame


def preview_column_labels(frame, raw_headers):
    """Display label per frame column, replacing pandas placeholder headers.

    pandas names a blank header "Unnamed: N" and a repeated one "X.1". Blank-header columns
    without data are dropped (label None); with data they become "Cột <letter>". Repeats keep
    a stable "X (2)" label because preview rows are keyed by column name.
    """
    labels = {}
    used = set()
    for position, column in enumerate(frame.columns):
        raw = raw_headers[position] if position < len(raw_headers) else None
        raw_blank = raw is None or (not isinstance(raw, str) and pd.isna(raw)) or not str(raw).strip()
        text = str(column)
        label = column
        if raw_blank and normalize_header(text).startswith("unnamed:"):
            if all(not clean_text(value) for value in frame.iloc[:, position].tolist()):
                labels[column] = None
                continue
            label = f"Cột {get_column_letter(position + 1)}"
        elif not raw_blank and text != str(raw):
            match = re.fullmatch(re.escape(str(raw)) + r"\.(\d+)", text)
            if match:
                label = f"{raw} ({int(match.group(1)) + 1})"
        while label in used:
            label = f"{label} ({len(used) + 1})"
        used.add(label)
        labels[column] = label
    return labels


def read_sheet_preview(file_path, sheet_name=None, limit=None, *, platform=None):
    with open_workbook_snapshot(file_path) as snapshot:
        sheets = snapshot.sheet_names
        current_sheet = sheet_name if sheet_name in sheets else (sheets[0] if sheets else "")
        if not current_sheet:
            return {"sheets": [], "currentSheet": "", "columns": [], "data": [], "message": "Workbook không có sheet nào."}
        preview = snapshot.cached(
            ("preview", current_sheet, platform or "", limit),
            lambda: build_sheet_preview(snapshot, current_sheet, limit, platform=platform),
        )
        hidden_sheets = snapshot.hidden_sheet_names
    return {
        "sheets": sheets,
        "hiddenSheets": hidden_sheets,
        "currentSheet": current_sheet,
        "columns": list(preview["columns"]),
        "data": [dict(row) for row in preview["data"]],
        "totalRows": preview["totalRows"],
        "shownRows": preview["shownRows"],
    }


def build_sheet_preview(snapshot, current_sheet, limit=None, *, platform=None):
    frame = snapshot.frame(current_sheet)
    raw_headers = snapshot.header_row(current_sheet)
    column_labels = preview_column_labels(frame, raw_headers)
    frame = fill_formula_sequence_numbers(frame, snapshot, current_sheet)
    link_column = find_link_column_name(frame)
    channel_column = find_column_name(frame, COLUMN_ALIASES["TÊN KÊNH"])
    date_column = find_column_name(frame, COLUMN_ALIASES["date"])
    partner_columns = dataframe_partner_columns(frame)
    views_column = find_column_name(frame, ["LƯỢT XEM"])
    likes_column = find_column_name(frame, ["TIM"])
    comments_column = find_column_name(frame, ["BÌNH LUẬN"])
    saves_column = find_column_name(frame, ["LƯỢT LƯU"])
    reposts_column = find_column_name(frame, ["REPOST"])
    shares_column = find_column_name(frame, ["CHIA SẺ"])
    scan_status_column = find_column_name(frame, [TTBD_SCAN_STATUS_HEADER])
    resolved_url_column = find_column_name(frame, [TTBD_RESOLVED_URL_HEADER])
    source_url_column = find_column_name(frame, [TTBD_SOURCE_URL_HEADER])
    public_columns = [
        column
        for column in frame.columns
        if clean_text(column) not in TTBD_INTERNAL_HEADERS and column_labels.get(column) is not None
    ]
    metric_columns = [
        views_column,
        likes_column,
        comments_column,
        saves_column,
        reposts_column,
        shares_column,
    ]
    frame = fill_missing_dates_from_previous(frame, date_column, link_column)
    for column in frame.select_dtypes(include=["datetime"]).columns:
        frame[column] = frame[column].map(lambda value: None if pd.isna(value) else preview_datetime_text(value))
    spec = PLATFORMS.get(platform or "")
    if spec and link_column and not is_summary_sheet_name(current_sheet):
        frame = frame[frame[link_column].map(lambda value: spec.is_link(value) or is_total_label(value))].copy()
        frame.reset_index(drop=True, inplace=True)
    frame = fill_preview_total_row(frame, link_column, metric_columns, platform=platform)

    preview_frame = frame
    if limit and link_column and len(frame.index) > limit:
        total_positions = _total_label_positions(frame, link_column)
        if total_positions and total_positions[-1] >= limit:
            preview_frame = pd.concat(
                [frame.head(max(limit - 1, 0)), frame.iloc[[total_positions[-1]]]],
                ignore_index=True,
            )
        else:
            preview_frame = frame.head(limit)
    elif limit:
        preview_frame = frame.head(limit)

    data = []
    for record in preview_frame.to_dict(orient="records"):
        preview_row = {}
        for column in public_columns:
            preview_row[column_labels[column]] = clean_preview_value(record.get(column, ""))
        if link_column and channel_column and column_labels.get(channel_column) is not None:
            channel_label = column_labels[channel_column]
            preview_row[channel_label] = display_channel_name_from_file(
                record.get(link_column, ""),
                preview_row.get(channel_label, ""),
            )
        link_value = record.get(link_column, "") if link_column else ""
        if link_column and is_tiktok_link(link_value):
            if partner_columns:
                row_partners = extract_row_partners(record, partner_columns)
                if len(row_partners) == 1:
                    preview_row["_singlePartner"] = True
            if should_highlight_video_link(
                link_value,
                likes=record.get(likes_column, "") if likes_column else 0,
                shares=record.get(shares_column, "") if shares_column else 0,
                resolved_url=record.get(resolved_url_column, "") if resolved_url_column else "",
                resolved_source_url=record.get(source_url_column, "") if source_url_column else "",
                scan_status=record.get(scan_status_column, "") if scan_status_column else "",
            ):
                preview_row["_videoLink"] = True
        data.append(preview_row)

    # Cached and shared between requests: read_sheet_preview hands out copies.
    return {
        "columns": tuple(column_labels[column] for column in public_columns),
        "data": tuple(data),
        "totalRows": int(len(frame.index)),
        "shownRows": int(len(preview_frame.index)),
    }


SUMMARY_TOTAL_KEYS = {
    "LƯỢT XEM": "views",
    "TIM": "likes",
    "BÌNH LUẬN": "comments",
    "LƯỢT LƯU": "saves",
    "REPOST": "reposts",
    "CHIA SẺ": "shares",
}


def read_summary_dashboard(file_path, data_sheet_name=None, *, platform="tiktok"):
    """Read the platform's summary tab for a data sheet.

    data_sheet_name may also be a summary tab title (a tab the user opened); that tab is read
    as-is, with the platform that owns it, so a Threads tab opened from the TikTok view still
    shows Threads totals.
    """
    spec = get_platform(platform)
    with open_workbook_snapshot(file_path) as snapshot:
        sheet_names = snapshot.sheet_names
        requested_data_sheet = clean_text(data_sheet_name)
        if requested_data_sheet and requested_data_sheet in sheet_names and is_summary_sheet_name(requested_data_sheet):
            summary_sheet = requested_data_sheet
            requested_data_sheet, owner = resolve_summary_sheet(sheet_names, summary_sheet)
            spec = get_platform(owner)
        elif requested_data_sheet:
            summary_sheet = summary_sheet_title_for_data_sheet(requested_data_sheet, spec.key)
            if summary_sheet not in sheet_names:
                summary_sheet = ""
        else:
            summary_sheet = next(
                (
                    sheet for sheet in sheet_names
                    if is_summary_sheet_name(sheet) and resolve_summary_sheet(sheet_names, sheet)[1] == spec.key
                ),
                "",
            )
        if not summary_sheet:
            label = requested_data_sheet or "sheet này"
            return {
                "sheet": "",
                "dataSheet": requested_data_sheet,
                "columns": list(spec.summary_columns),
                "rows": [],
                "totals": {},
                "message": f"Workbook chưa có tổng kết cho {label}. Quét sheet đó để tạo.",
            }

        resolved_data_sheet = requested_data_sheet or data_sheet_name_for_summary_title(
            sheet_names, summary_sheet
        )
        frame = snapshot.frame(summary_sheet).fillna("")

    columns = frame.columns.tolist()
    if LAST_UPDATE_COLUMN not in columns:
        columns.append(LAST_UPDATE_COLUMN)
    rows = []
    partner_column = find_column_name(frame, ["ĐỐI TÁC", "Đối tác"])
    total_link_column = find_column_name(frame, ["TỔNG LINK", "Tổng link"])
    summary_metric_map = {
        SUMMARY_TOTAL_KEYS[metric]: find_column_name(frame, [f"TỔNG {metric}", metric])
        for metric in spec.metric_columns
    }

    numeric_summary_keys = {
        normalize_key(item)
        for item in spec.summary_columns
        if normalize_key(item) not in {"doi tac", normalize_key(LAST_UPDATE_COLUMN)}
    }
    def summary_value(column, value):
        if normalize_key(column) not in numeric_summary_keys:
            return clean_text(value)
        # A blank metric total means nothing was known for that partner; it must not become 0.
        return metric_total([value])

    totals_row = None
    for record in frame.to_dict(orient="records"):
        partner_name = clean_text(record.get(partner_column or "", ""))
        if not partner_name:
            continue
        row = {column: summary_value(column, record.get(column, "")) for column in columns}
        row.setdefault(LAST_UPDATE_COLUMN, "")
        if is_total_label(partner_name):
            totals_row = row
            continue
        rows.append(row)

    source_rows = [totals_row] if totals_row else rows
    totals = {
        "partners": len(rows),
        "links": sum(metric_number(row.get(total_link_column, 0)) for row in source_rows) if total_link_column else 0,
    }
    for key, column in summary_metric_map.items():
        totals[key] = metric_total(row.get(column, "") for row in source_rows) if column else 0

    return {
        "sheet": summary_sheet,
        "dataSheet": resolved_data_sheet,
        "columns": columns,
        "rows": rows,
        "totals": totals,
    }
