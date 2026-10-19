import base64
import hashlib
import hmac
import json
import os
import secrets
import time
import urllib.request
from pathlib import Path
from urllib.parse import parse_qs, urlsplit
from google.auth.transport.requests import Request
from google.oauth2 import service_account
from google.oauth2.credentials import Credentials
from google_auth_oauthlib.flow import Flow, InstalledAppFlow, WSGITimeoutError
from googleapiclient.discovery import build
from googleapiclient.errors import HttpError

from riviu import server_mode

from riviu.workbook_utils import (
    format_google_sheet_datetime,
    get_platform,
    google_to_excel_sheet_titles,
    month_label_for_sheet_name,
    store_google_sheet_title_map,
    set_cell_literal,
    write_json_atomic,
)

CLIENT_SECRET_FILENAME = "google_oauth_client.json"
TOKEN_FILENAME = "google_oauth_token.json"
LOGIN_TIMEOUT_SECONDS = 300
PASTE_BACK_TTL_SECONDS = 600


class GoogleLoginRequired(RuntimeError):
    """Raised instead of opening a browser login from a background push or sync."""

    def __init__(self, message="Cần nạp tài khoản dịch vụ Google hoặc đăng nhập Google lại trong phần Nguồn dữ liệu."):
        super().__init__(message)


class ServiceAccountError(ValueError):
    """A service-account key that cannot be used; the message never contains the key."""


class ServiceAccountAccessError(RuntimeError):
    """Google answered 403/404 under the service account: the sheet is not shared with it."""

    def __init__(self, email, status):
        self.email = email
        self.status = status
        super().__init__(
            f"Tài khoản dịch vụ {email} chưa có quyền với Google Sheet này (Google trả lỗi {status}). "
            f"Chia sẻ Google Sheet cho email {email} với quyền Người chỉnh sửa rồi thử lại."
        )


def fetch_account_email(access_token):
    token = str(access_token or "").strip()
    if not token:
        return ""
    request = urllib.request.Request(
        "https://openidconnect.googleapis.com/v1/userinfo",
        headers={"Authorization": f"Bearer {token}"},
    )
    try:
        with urllib.request.urlopen(request, timeout=10) as response:
            payload = json.loads(response.read().decode("utf-8"))
    except Exception:
        return ""
    return str(payload.get("email") or "").strip()


SCOPES = [
    "https://www.googleapis.com/auth/spreadsheets",
    "https://www.googleapis.com/auth/userinfo.email",
    "openid",
]


def token_has_required_scopes(token_data):
    token_scopes = set(token_data.get("scopes") or [])
    return set(SCOPES).issubset(token_scopes)


def rgb(red, green, blue):
    return {"red": red / 255, "green": green / 255, "blue": blue / 255}


PUSH_COLUMN_WIDTHS = {
    "Stt": 46, "Ngày": 94, "Link": 510, "Tên Kênh": 150, "LƯỢT XEM": 120, "TIM": 68,
    "BÌNH LUẬN": 150, "LƯỢT LƯU": 145, "REPOST": 145, "CHIA SẺ": 112,
    "Cập nhật lần cuối": 145,
}
PARTNER_COLUMN_WIDTH = 210


def column_widths(headers):
    return [
        PUSH_COLUMN_WIDTHS.get(header, PARTNER_COLUMN_WIDTH if str(header).startswith("Đối tác") else 120)
        for header in headers
    ]


def format_result_sheet(service, spreadsheet_id, sheet_id, headers, row_count, metric_count=5):
    column_count = len(headers)
    if not row_count or not column_count:
        return

    frozen_rows = 1
    widths = column_widths(headers)
    metric_end = 4 + metric_count
    requests = [
        {
            "updateSheetProperties": {
                "properties": {
                    "sheetId": sheet_id,
                    "gridProperties": {"frozenRowCount": frozen_rows},
                },
                "fields": "gridProperties.frozenRowCount",
            }
        },
        {
            "repeatCell": {
                "range": {"sheetId": sheet_id, "startRowIndex": 0, "endRowIndex": 1, "startColumnIndex": 0, "endColumnIndex": column_count},
                "cell": {
                    "userEnteredFormat": {
                        "backgroundColor": rgb(208, 226, 243),
                        "horizontalAlignment": "CENTER",
                        "verticalAlignment": "MIDDLE",
                        "wrapStrategy": "WRAP",
                        "textFormat": {
                            "fontFamily": "Times New Roman",
                            "fontSize": 15,
                            "bold": True,
                            "foregroundColor": rgb(0, 0, 0),
                        },
                    }
                },
                "fields": "userEnteredFormat(backgroundColor,horizontalAlignment,verticalAlignment,wrapStrategy,textFormat)",
            }
        },
        {
            "repeatCell": {
                "range": {"sheetId": sheet_id, "startRowIndex": 1, "endRowIndex": row_count, "startColumnIndex": 0, "endColumnIndex": column_count},
                "cell": {
                    "userEnteredFormat": {
                        "verticalAlignment": "MIDDLE",
                        "wrapStrategy": "CLIP",
                        "textFormat": {"fontFamily": "Arial", "fontSize": 10},
                    }
                },
                "fields": "userEnteredFormat(verticalAlignment,wrapStrategy,textFormat)",
            }
        },
        {
            "repeatCell": {
                "range": {"sheetId": sheet_id, "startRowIndex": 1, "endRowIndex": row_count, "startColumnIndex": 0, "endColumnIndex": 2},
                "cell": {"userEnteredFormat": {"horizontalAlignment": "CENTER"}},
                "fields": "userEnteredFormat.horizontalAlignment",
            }
        },
        {
            "repeatCell": {
                "range": {"sheetId": sheet_id, "startRowIndex": 1, "endRowIndex": row_count, "startColumnIndex": 4, "endColumnIndex": metric_end},
                "cell": {
                    "userEnteredFormat": {
                        "horizontalAlignment": "RIGHT",
                        "numberFormat": {"type": "NUMBER", "pattern": "#,##0"},
                    }
                },
                "fields": "userEnteredFormat(horizontalAlignment,numberFormat)",
            }
        },
        {
            "repeatCell": {
                "range": {"sheetId": sheet_id, "startRowIndex": 1, "endRowIndex": row_count, "startColumnIndex": metric_end, "endColumnIndex": column_count},
                "cell": {"userEnteredFormat": {"horizontalAlignment": "LEFT"}},
                "fields": "userEnteredFormat.horizontalAlignment",
            }
        },
        {
            "updateBorders": {
                "range": {"sheetId": sheet_id, "startRowIndex": 0, "endRowIndex": row_count, "startColumnIndex": 0, "endColumnIndex": column_count},
                "top": {"style": "SOLID", "width": 1, "color": rgb(0, 0, 0)},
                "bottom": {"style": "SOLID", "width": 1, "color": rgb(0, 0, 0)},
                "left": {"style": "SOLID", "width": 1, "color": rgb(0, 0, 0)},
                "right": {"style": "SOLID", "width": 1, "color": rgb(0, 0, 0)},
                "innerHorizontal": {"style": "SOLID", "width": 1, "color": rgb(0, 0, 0)},
                "innerVertical": {"style": "SOLID", "width": 1, "color": rgb(0, 0, 0)},
            }
        },
    ]

    if row_count > 2:
        requests.extend([
            {
                "repeatCell": {
                    "range": {
                        "sheetId": sheet_id,
                        "startRowIndex": row_count - 1,
                        "endRowIndex": row_count,
                        "startColumnIndex": 0,
                        "endColumnIndex": column_count,
                    },
                    "cell": {
                        "userEnteredFormat": {
                            "backgroundColor": rgb(232, 240, 254),
                            "textFormat": {"fontFamily": "Arial", "fontSize": 10, "bold": True},
                            "verticalAlignment": "MIDDLE",
                        }
                    },
                    "fields": "userEnteredFormat(backgroundColor,textFormat,verticalAlignment)",
                }
            },
            {
                "repeatCell": {
                    "range": {
                        "sheetId": sheet_id,
                        "startRowIndex": row_count - 1,
                        "endRowIndex": row_count,
                        "startColumnIndex": 4,
                        "endColumnIndex": metric_end,
                    },
                    "cell": {
                        "userEnteredFormat": {
                            "horizontalAlignment": "RIGHT",
                            "numberFormat": {"type": "NUMBER", "pattern": "#,##0"},
                        }
                    },
                    "fields": "userEnteredFormat(horizontalAlignment,numberFormat)",
                }
            },
        ])

    for index, width in enumerate(widths):
        requests.append({
            "updateDimensionProperties": {
                "range": {"sheetId": sheet_id, "dimension": "COLUMNS", "startIndex": index, "endIndex": index + 1},
                "properties": {"pixelSize": width},
                "fields": "pixelSize",
            }
        })

    service.spreadsheets().batchUpdate(
        spreadsheetId=spreadsheet_id,
        body={"requests": requests},
    ).execute()


def apply_link_formatting(service, spreadsheet_id, sheet_id, rows):
    if not rows or len(rows) < 2:
        return

    requests = []
    for row_index, row in enumerate(rows[1:], start=1):
        url = str(row[2] or "").strip() if len(row) > 2 else ""
        if not url.lower().startswith("http"):
            continue
        requests.append({
            "repeatCell": {
                "range": {
                    "sheetId": sheet_id,
                    "startRowIndex": row_index,
                    "endRowIndex": row_index + 1,
                    "startColumnIndex": 2,
                    "endColumnIndex": 3,
                },
                "cell": {
                    "userEnteredFormat": {
                        "textFormat": {
                            "foregroundColor": rgb(0, 74, 198),
                            "underline": True,
                            "link": {"uri": url},
                        }
                    }
                },
                "fields": "userEnteredFormat.textFormat",
            }
        })

    if requests:
        service.spreadsheets().batchUpdate(
            spreadsheetId=spreadsheet_id,
            body={"requests": requests},
        ).execute()


def client_secret_candidates(base_dir):
    return [
        os.path.join(base_dir, CLIENT_SECRET_FILENAME),
        os.path.join(base_dir, "data", CLIENT_SECRET_FILENAME),
    ]


def client_secret_path(base_dir):
    for path in client_secret_candidates(base_dir):
        if os.path.exists(path):
            return path
    return client_secret_candidates(base_dir)[0]


def token_path(base_dir):
    return os.path.join(base_dir, "data", TOKEN_FILENAME)


# ---- Service account ----------------------------------------------------------------
# A Google Cloud service account (x@project.iam.gserviceaccount.com) reads and writes every
# sheet shared with its email as Editor, so a shared server does not depend on one person's
# Google login. Once its key is uploaded it is used for sync, download and push in
# preference to the OAuth token; without it the OAuth behaviour is unchanged.
# Only the Sheets API is called, so the key gets the spreadsheets scope and nothing else.
SERVICE_ACCOUNT_FILENAME = "google_service_account.json"
SERVICE_ACCOUNT_SCOPES = ["https://www.googleapis.com/auth/spreadsheets"]
SERVICE_ACCOUNT_MAX_BYTES = 64 * 1024
SEALED_FORMAT = "riviu-aesgcm-v1"
_SERVICE_ACCOUNT_FIELDS = ("type", "project_id", "private_key_id", "private_key", "client_email", "client_id", "token_uri")
_DEFAULT_TOKEN_URI = "https://oauth2.googleapis.com/token"


def service_account_path(base_dir):
    return os.path.join(base_dir, "data", SERVICE_ACCOUNT_FILENAME)


def validate_service_account_info(payload):
    """The fields google-auth needs, after checking this really is a service-account key."""
    if not isinstance(payload, dict) or payload.get("type") != "service_account":
        raise ServiceAccountError('File không phải khoá tài khoản dịch vụ Google (cần "type": "service_account").')
    email = payload.get("client_email")
    if (not isinstance(email, str) or len(email) > 320 or email.count("@") != 1
            or any(char.isspace() for char in email) or not email.split("@")[0] or not email.split("@")[1]):
        raise ServiceAccountError("Khoá tài khoản dịch vụ thiếu client_email hợp lệ.")
    key = payload.get("private_key")
    if not isinstance(key, str) or "PRIVATE KEY" not in key or len(key) > 16384:
        raise ServiceAccountError("Khoá tài khoản dịch vụ thiếu private_key.")
    token_uri = payload.get("token_uri", _DEFAULT_TOKEN_URI)
    parsed = urlsplit(token_uri) if isinstance(token_uri, str) else None
    # google-auth posts the signed assertion to token_uri: only ever to Google.
    if parsed is None or parsed.scheme != "https" or not (parsed.hostname or "").endswith(".googleapis.com"):
        raise ServiceAccountError("token_uri của khoá tài khoản dịch vụ phải là địa chỉ googleapis.com.")
    info = {field: payload[field] for field in _SERVICE_ACCOUNT_FIELDS if isinstance(payload.get(field), str)}
    info["token_uri"] = token_uri
    return info


def _sealing_binding(base_dir):
    # Bound to the data root, so a copied file does not decrypt for another root.
    root = os.path.normcase(str(Path(base_dir).resolve())).encode("utf-8")
    return hashlib.sha256(b"Riviu.GoogleServiceAccount.server.v1\0" + root).digest()


def _seal(key, base_dir, info):
    from cryptography.hazmat.primitives.ciphers.aead import AESGCM

    nonce = os.urandom(12)
    sealed = AESGCM(key).encrypt(nonce, json.dumps(info).encode("utf-8"), _sealing_binding(base_dir))
    return {"format": SEALED_FORMAT, "nonce": base64.b64encode(nonce).decode("ascii"),
            "data": base64.b64encode(sealed).decode("ascii")}


def _unseal(key, base_dir, envelope):
    from cryptography.hazmat.primitives.ciphers.aead import AESGCM

    try:
        nonce = base64.b64decode(envelope["nonce"], validate=True)
        sealed = base64.b64decode(envelope["data"], validate=True)
        return json.loads(AESGCM(key).decrypt(nonce, sealed, _sealing_binding(base_dir)).decode("utf-8"))
    except Exception:
        raise ServiceAccountError("Không giải mã được khoá tài khoản dịch vụ đã lưu. Nạp lại file JSON.") from None


def save_service_account(base_dir, payload):
    """Store a validated key in the data dir and return its email. In server mode the file
    is AES-GCM encrypted at rest with RIVIU_SECRET_KEY_FILE, like the cookie stores."""
    info = validate_service_account_info(payload)
    path = service_account_path(base_dir)
    os.makedirs(os.path.dirname(path), exist_ok=True)
    key = server_mode.secret_key()
    write_json_atomic(path, _seal(key, base_dir, info) if key is not None else info)
    return info["client_email"]


def load_service_account_info(base_dir):
    """The stored key, or None when no service account is configured."""
    path = service_account_path(base_dir)
    if not os.path.exists(path):
        return None
    try:
        with open(path, "r", encoding="utf-8") as file_obj:
            stored = json.load(file_obj)
    except (OSError, ValueError):
        raise ServiceAccountError("File tài khoản dịch vụ đã lưu bị hỏng. Nạp lại file JSON.") from None
    if isinstance(stored, dict) and stored.get("format") == SEALED_FORMAT:
        key = server_mode.secret_key()
        if key is None:
            raise ServiceAccountError("Khoá tài khoản dịch vụ được mã hoá cho chế độ server; nạp lại file JSON trên máy này.")
        stored = _unseal(key, base_dir, stored)
    return validate_service_account_info(stored)


def delete_service_account(base_dir):
    path = service_account_path(base_dir)
    if os.path.exists(path):
        os.remove(path)


def service_account_status(base_dir):
    try:
        info = load_service_account_info(base_dir)
    except ServiceAccountError as error:
        return {"configured": True, "email": "", "error": str(error)}
    return {"configured": info is not None, "email": info["client_email"] if info else "", "error": ""}


def service_account_credentials(base_dir):
    info = load_service_account_info(base_dir)
    if info is None:
        return None
    try:
        return service_account.Credentials.from_service_account_info(info, scopes=SERVICE_ACCOUNT_SCOPES)
    except Exception:
        raise ServiceAccountError("Không đọc được private_key của tài khoản dịch vụ. Tải lại file JSON từ Google Cloud.") from None


def google_credentials(base_dir):
    """The credential Sheets calls use: the service account when one is configured, else a
    valid OAuth token, else None. Never opens a browser."""
    creds = service_account_credentials(base_dir)
    if creds is not None:
        return creds
    creds = try_load_credentials(base_dir)
    return creds if creds and creds.valid else None


def service_account_access_error(base_dir, error):
    """For a Google 403/404 while the service account is the active credential: an error
    that names the email to share the sheet with. None otherwise."""
    status = getattr(getattr(error, "resp", None), "status", None)
    try:
        status = int(status)
    except (TypeError, ValueError):
        return None
    if not isinstance(error, HttpError) or status not in (403, 404):
        return None
    try:
        info = load_service_account_info(base_dir)
    except ServiceAccountError:
        return None
    return ServiceAccountAccessError(info["client_email"], status) if info else None


def write_token(token_file, creds):
    write_json_atomic(token_file, json.loads(creds.to_json()))


def try_load_credentials(base_dir):
    """Return stored credentials (refreshed when possible) without ever opening a browser."""
    token_file = token_path(base_dir)
    if not os.path.exists(token_file):
        return None
    try:
        with open(token_file, "r", encoding="utf-8") as file_obj:
            token_data = json.load(file_obj)
        if not token_has_required_scopes(token_data):
            return None
        creds = Credentials.from_authorized_user_file(token_file, SCOPES)
        if creds.valid:
            return creds
        if creds.expired and creds.refresh_token and creds.has_scopes(SCOPES):
            creds.refresh(Request())
            write_token(token_file, creds)
            return creds
        return creds
    except Exception:
        return None


def oauth_status(base_dir):
    token_file = token_path(base_dir)
    account_email = ""
    if os.path.exists(token_file):
        try:
            with open(token_file, "r", encoding="utf-8") as file_obj:
                token_data = json.load(file_obj)
            account_email = str(token_data.get("account") or token_data.get("email") or "").strip()
            if not account_email:
                account_email = fetch_account_email(token_data.get("token", ""))
                if account_email:
                    token_data["account"] = account_email
                    write_json_atomic(token_file, token_data)
        except (OSError, json.JSONDecodeError):
            account_email = ""
    creds = try_load_credentials(base_dir)
    valid = bool(creds and creds.valid)
    service = service_account_status(base_dir)
    usable_service = service["configured"] and not service["error"]
    return {
        "configured": os.path.exists(client_secret_path(base_dir)),
        "authorized": os.path.exists(token_file),
        "valid": valid,
        "accountEmail": account_email,
        "serviceAccount": service,
        # The credential Sheets calls use now; the service account wins over OAuth.
        "activeCredential": "service_account" if usable_service else ("oauth" if valid else "none"),
        "ready": usable_service or valid,
    }


def save_oauth_client(base_dir, payload):
    path = client_secret_path(base_dir)
    os.makedirs(os.path.dirname(path), exist_ok=True)
    write_json_atomic(path, payload)
    return path


def load_credentials(base_dir):
    creds = google_credentials(base_dir)
    if creds is not None:
        return creds
    raise GoogleLoginRequired()


def authorize_google(base_dir, timeout_seconds=LOGIN_TIMEOUT_SECONDS):
    """Interactive browser login; only the explicit login endpoint may call this."""
    creds = try_load_credentials(base_dir)
    if creds and creds.valid:
        return True
    client_file = client_secret_path(base_dir)
    if not os.path.exists(client_file):
        raise FileNotFoundError("Chưa có file OAuth client Google.")
    flow = InstalledAppFlow.from_client_secrets_file(client_file, SCOPES)
    try:
        creds = flow.run_local_server(host="127.0.0.1", port=0, open_browser=True, timeout_seconds=timeout_seconds)
    except WSGITimeoutError as error:
        raise TimeoutError("Hết thời gian chờ đăng nhập Google. Hãy thử lại.") from error
    write_token(token_path(base_dir), creds)
    return True


# ---- Server mode: paste-back login ---------------------------------------------------
# The server cannot open a browser or receive the loopback redirect. With the same
# Desktop-type OAuth client, Google redirects the user's own browser to
# http://localhost:<port>/?state=...&code=...; that page fails to load and the user pastes
# the address back. PKCE binds the code to this server's verifier; the state is checked
# whenever the full address is pasted.

def start_paste_back_login(base_dir):
    """A consent URL plus the pending state to keep server side (never sent to the browser)."""
    client_file = client_secret_path(base_dir)
    if not os.path.exists(client_file):
        raise FileNotFoundError("Chưa có file OAuth client Google.")
    redirect_uri = f"http://localhost:{49152 + secrets.randbelow(16383)}/"
    flow = Flow.from_client_secrets_file(client_file, scopes=SCOPES, redirect_uri=redirect_uri, autogenerate_code_verifier=True)
    auth_url, state = flow.authorization_url(access_type="offline", prompt="consent")
    return {
        "authUrl": auth_url,
        "state": state,
        "codeVerifier": flow.code_verifier,
        "redirectUri": redirect_uri,
        "expiresAt": time.time() + PASTE_BACK_TTL_SECONDS,
    }


def paste_back_code(pending, pasted):
    """The authorization code from a pasted redirect address (state checked) or a bare code."""
    if not pending or time.time() > float(pending.get("expiresAt", 0)):
        raise ValueError("Phiên đăng nhập Google đã hết hạn. Bấm Đăng nhập Google để bắt đầu lại.")
    text = str(pasted or "").strip()
    if not text or len(text) > 4096:
        raise ValueError("Hãy dán địa chỉ localhost mà Google chuyển tới.")
    if "://" not in text and "=" not in text:
        return text
    query = parse_qs(urlsplit(text).query if "://" in text else text.lstrip("?"))
    if query.get("error"):
        raise ValueError(f"Google từ chối đăng nhập: {query['error'][0]}")
    state = (query.get("state") or [""])[0]
    code = (query.get("code") or [""])[0]
    if not state or not hmac.compare_digest(state.encode("utf-8"), str(pending.get("state", "")).encode("utf-8")):
        raise ValueError("Địa chỉ dán vào không khớp phiên đăng nhập này (state sai). Bấm Đăng nhập Google để thử lại.")
    if not code:
        raise ValueError("Địa chỉ dán vào không có mã code của Google.")
    return code


def finish_paste_back_login(base_dir, pending, pasted):
    code = paste_back_code(pending, pasted)
    flow = Flow.from_client_secrets_file(
        client_secret_path(base_dir), scopes=SCOPES, state=pending["state"],
        redirect_uri=pending["redirectUri"], code_verifier=pending["codeVerifier"],
    )
    flow.fetch_token(code=code)
    write_token(token_path(base_dir), flow.credentials)
    return True


def sheets_service(base_dir):
    return build("sheets", "v4", credentials=load_credentials(base_dir), cache_discovery=False)


def download_google_sheet_authenticated(base_dir, spreadsheet_id, destination_path):
    creds = google_credentials(base_dir)
    if creds is None:
        raise RuntimeError("Chưa đăng nhập Google hợp lệ.")
    service = build("sheets", "v4", credentials=creds, cache_discovery=False)
    try:
        return _download_google_sheet(service, spreadsheet_id, destination_path)
    except HttpError as error:
        access_error = service_account_access_error(base_dir, error)
        if access_error is not None:
            raise access_error from None
        raise


def _download_google_sheet(service, spreadsheet_id, destination_path):
    from openpyxl import Workbook

    spreadsheet = service.spreadsheets().get(spreadsheetId=spreadsheet_id).execute()
    workbook = Workbook()
    workbook.remove(workbook.active)
    sheets = spreadsheet.get("sheets", [])
    title_mapping = google_to_excel_sheet_titles(sheet["properties"]["title"] for sheet in sheets)
    store_google_sheet_title_map(workbook, title_mapping)
    if not sheets:
        workbook.create_sheet("Sheet1")
    else:
        for sheet_meta in sheets:
            title = sheet_meta["properties"]["title"]
            escaped = title.replace("'", "''")
            values = service.spreadsheets().values().get(
                spreadsheetId=spreadsheet_id,
                range=f"'{escaped}'!A:ZZ",
                valueRenderOption="UNFORMATTED_VALUE",
                dateTimeRenderOption="FORMATTED_STRING",
            ).execute().get("values", [])
            worksheet = workbook.create_sheet(title=title_mapping[title])
            for row_index, row in enumerate(values, start=1):
                for col_index, value in enumerate(row, start=1):
                    set_cell_literal(worksheet.cell(row=row_index, column=col_index), value)
    os.makedirs(os.path.dirname(destination_path), exist_ok=True)
    workbook.save(destination_path)
    return destination_path


def create_result_sheet_title(source_sheet_name="", *, platform="tiktok"):
    timestamp = format_google_sheet_datetime()
    month_label = month_label_for_sheet_name(source_sheet_name)
    title = f"{month_label} {timestamp}" if month_label else timestamp
    prefix = get_platform(platform).google_sheet_prefix
    return f"{prefix} {title}" if prefix else title


def ensure_unique_sheet_title(existing_titles, desired_title):
    title = desired_title[:100]
    if title not in existing_titles:
        return title
    suffix = 2
    base = title[:95]
    while f"{base}-{suffix}" in existing_titles:
        suffix += 1
    return f"{base}-{suffix}"


def push_rows_to_new_sheet(base_dir, spreadsheet_id, rows, source_sheet_name="", *, platform="tiktok"):
    service = sheets_service(base_dir)
    try:
        return _push_rows(service, spreadsheet_id, rows, source_sheet_name, platform)
    except HttpError as error:
        access_error = service_account_access_error(base_dir, error)
        if access_error is not None:
            raise access_error from None
        raise


def _push_rows(service, spreadsheet_id, rows, source_sheet_name, platform):
    spreadsheet = service.spreadsheets().get(
        spreadsheetId=spreadsheet_id,
        fields="sheets.properties.title",
    ).execute()
    existing_titles = [item["properties"]["title"] for item in spreadsheet.get("sheets", [])]
    title = ensure_unique_sheet_title(existing_titles, create_result_sheet_title(source_sheet_name, platform=platform))

    add_sheet_response = service.spreadsheets().batchUpdate(
        spreadsheetId=spreadsheet_id,
        body={"requests": [{"addSheet": {"properties": {"title": title}}}]},
    ).execute()
    sheet_id = add_sheet_response["replies"][0]["addSheet"]["properties"]["sheetId"]

    service.spreadsheets().values().update(
        spreadsheetId=spreadsheet_id,
        range=f"'{title}'!A1",
        valueInputOption="RAW",
        body={"values": rows},
    ).execute()
    # Only the app-generated final metric totals may be interpreted as formulas.
    # All channel, partner, date, and source text above remains RAW, unchanged.
    if len(rows) > 2 and len(rows[-1]) > 2 and rows[-1][2] == "TỔNG":
        for index in range(4, min(4 + len(get_platform(platform).metric_columns), len(rows[-1]))):
            column = chr(ord("A") + index)
            expected = f"=SUM({column}2:{column}{len(rows) - 1})"
            if rows[-1][index] != expected:
                continue
            service.spreadsheets().values().update(
                spreadsheetId=spreadsheet_id,
                range=f"'{title}'!{column}{len(rows)}",
                valueInputOption="USER_ENTERED",
                body={"values": [[expected]]},
            ).execute()
    metric_count = len(get_platform(platform).metric_columns)
    format_result_sheet(service, spreadsheet_id, sheet_id, rows[0], len(rows), metric_count)
    apply_link_formatting(service, spreadsheet_id, sheet_id, rows)
    return title
