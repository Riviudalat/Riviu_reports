"""Resolve and publish the workbooks a request names, bound to one platform.

An explicit `file_id` must name an existing workbook; only an omitted one falls back to the
legacy selection in `riviu.server_context`.
"""

import os
from urllib.parse import quote

from fastapi.responses import JSONResponse
from openpyxl import load_workbook

from riviu import server_context as context
from riviu.workbook_utils import (
    LEGACY_GOOGLE_SHEET_FILE_ID,
    find_data_sheet_names,
    get_platform,
    google_sheet_source_for_file,
    read_sheet_preview,
    safe_join,
    workbook_file_entries,
)


class SourceRequestError(ValueError):
    def __init__(self, message, status_code=400):
        super().__init__(message)
        self.status_code = status_code

    def response(self):
        return JSONResponse({"error": str(self)}, status_code=self.status_code)


def validate_workbook(path):
    book = load_workbook(path, read_only=True)
    try:
        if not book.worksheets:
            raise ValueError("Workbook không có trang tính")
        # Validate the ZIP/XML payload while reading headers, before publishing it.
        for sheet in book:
            next(sheet.iter_rows(values_only=True), None)
    finally:
        book.close()
    return read_sheet_preview(path)


def publish_new_workbook(staged_path, desired_path):
    """Publish validated bytes without ever overwriting an existing path."""
    stem, extension = os.path.splitext(desired_path)
    candidate = desired_path
    suffix = 2
    while True:
        try:
            os.link(staged_path, candidate)
            return candidate
        except FileExistsError:
            candidate = f"{stem} ({suffix}){extension}"
            suffix += 1


def file_entries():
    return workbook_file_entries(context.EXCEL_DIR)


def resolve_file_path(file_id):
    if not file_id:
        return ""
    return safe_join(context.EXCEL_DIR, file_id.replace("\\", "/"))


def validate_platform(platform="tiktok"):
    try:
        return get_platform(platform).key
    except ValueError as error:
        raise SourceRequestError(str(error)) from error


def resolve_source_file(file_id=None, platform="tiktok", *, allow_empty=False):
    """Capture an exact source; only an omitted ID may use legacy selection."""
    platform = validate_platform(platform)
    if file_id is None:
        file_id = ensure_selected_file()
    elif not isinstance(file_id, str):
        raise SourceRequestError("File không hợp lệ")
    else:
        file_id = file_id.replace("\\", "/")
        if not file_id:
            if allow_empty:
                return "", "", "", platform
            raise SourceRequestError("Chưa chọn workbook hợp lệ")
        if file_id not in {entry["id"] for entry in file_entries()}:
            raise SourceRequestError("File không tồn tại", 404)
    target_path = resolve_file_path(file_id)
    if file_id and (not target_path or not os.path.isfile(target_path)):
        raise SourceRequestError("File không tồn tại", 404)
    return file_id, target_path, file_display_label(file_id), platform


def source_sheet_default(file_id, *, scan=False, explicit=False):
    if not explicit:
        return (context.CURRENT_SCAN_SHEET if scan else "") or context.CURRENT_SELECTED_SHEET or ""
    return default_sheet_for_file(file_id)


def ensure_selected_file():
    entries = file_entries()
    available_ids = {entry["id"] for entry in entries}

    if context.CURRENT_SELECTED_FILE and context.CURRENT_SELECTED_FILE in available_ids:
        return context.CURRENT_SELECTED_FILE

    if not entries:
        context.CURRENT_SELECTED_FILE = ""
        context.CURRENT_SELECTED_SHEET = ""
        context.CURRENT_SCAN_SHEET = ""
        return ""

    google_ids = [
        entry["id"]
        for entry in entries
        if entry.get("source") == "google" and entry["id"] != LEGACY_GOOGLE_SHEET_FILE_ID
    ]
    google_ids = sorted(
        google_ids,
        key=lambda file_id: os.path.getmtime(resolve_file_path(file_id)) if os.path.exists(resolve_file_path(file_id)) else 0,
        reverse=True,
    )
    preferred_ids = google_ids + [LEGACY_GOOGLE_SHEET_FILE_ID, "Report_v1.xlsx", "Report_v1_with_partners.xlsx"]
    context.CURRENT_SELECTED_FILE = next((file_id for file_id in preferred_ids if file_id in available_ids), entries[0]["id"])
    return context.CURRENT_SELECTED_FILE


def select_source(file_id, selected_sheet, scan_sheet):
    """Make a workbook the legacy current source used by callers that omit file_id."""
    context.CURRENT_SELECTED_FILE = file_id
    context.CURRENT_SELECTED_SHEET = selected_sheet
    context.CURRENT_SCAN_SHEET = scan_sheet


def current_google_sheet_source():
    current_id = ensure_selected_file()
    return google_sheet_source_for_file(context.EXCEL_DIR, current_id)


def file_display_label(file_id):
    return next((entry["label"] for entry in file_entries() if entry["id"] == file_id), file_id)


def default_sheet_for_file(file_id):
    if not file_id:
        return ""
    try:
        sheets = find_data_sheet_names(resolve_file_path(file_id))
        return sheets[0] if sheets else ""
    except Exception:
        return ""


def file_id_for_path(path):
    return os.path.relpath(path, context.EXCEL_DIR).replace("\\", "/")


def content_disposition(filename):
    return f"attachment; filename*=UTF-8''{quote(filename)}"
