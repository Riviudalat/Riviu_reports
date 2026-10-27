"""Regenerate the usage-guide screenshots in riviu/web/static/guide/.

    uv run python tools/make_guide_screenshots.py

Runs the real app in server mode in-process (temp data dir, throwaway password and key)
with SYNTHETIC workbooks, and drives it with Playwright Chromium at 1280x800. Scans go
through the real runner and WebSocket events; only the per-link TikTok fetch is replaced
by a fake that returns made-up numbers, and Python sockets are limited to loopback, so
nothing reaches TikTok, Threads or Google. (The browser still loads the UI's web fonts.)
Red numbered callouts match the ①②③ in riviu/web/templates/guide.html; update both
together. PNGs are palette-quantized to stay small.
"""

import base64
import io
import os
import shutil
import socket
import sys
import tempfile
import threading
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "riviu" / "web" / "static" / "guide"
VIEWPORT = {"width": 1280, "height": 800}
PASSWORD = "guide-demo-password"
MAX_BYTES = 200 * 1024
TIKTOK_FILE = "Demo TikTok.xlsx"
THREADS_FILE = "Demo Threads.xlsx"
SHEET = "Tháng 10"
PARTNERS = ["Đối tác Mẫu A", "Đối tác Mẫu B", "Đối tác Mẫu C", "Đối tác Mẫu D"]
LINK_COUNT = 14
FAILING_LINK = 3      # shows up under "Lỗi quét"
NO_STATS_LINK = 6     # shows up under "Không số liệu"
FAKE_FETCH_SECONDS = 1.0
# FAKE credentials only: made-up values that are not keys or sessions of any account.
FAKE_SERVICE_ACCOUNT = {
    "type": "service_account", "project_id": "demo-project", "private_key_id": "demo",
    "private_key": "-----BEGIN PRIVATE KEY-----\nFAKE-KEY-FOR-SCREENSHOTS\n-----END PRIVATE KEY-----\n",
    "client_email": "riviu-reports@demo-project.iam.gserviceaccount.com", "client_id": "000000000000000000000",
    "token_uri": "https://oauth2.googleapis.com/token",
}
FAKE_TIKTOK_COOKIES = [
    {"domain": ".tiktok.com", "name": "sessionid", "value": "fake-session-for-screenshots", "path": "/",
     "secure": True, "httpOnly": True, "hostOnly": False, "sameSite": "no_restriction"},
    {"domain": ".tiktok.com", "name": "tt_csrf_token", "value": "fake-csrf", "path": "/",
     "secure": True, "httpOnly": False, "hostOnly": False, "sameSite": "lax"},
]


def block_python_network():
    """Only loopback connections: the app must not reach TikTok/Google while posing."""
    original = socket.socket.connect

    def connect(self, address):
        host = address[0] if isinstance(address, tuple) else str(address)
        if host not in {"127.0.0.1", "::1", "localhost"}:
            raise OSError("network disabled while generating guide screenshots")
        return original(self, address)

    socket.socket.connect = connect


def tiktok_link(number):
    return f"https://www.tiktok.com/@demo/video/74000000000000000{number:02d}"


def write_workbooks(data_dir):
    from openpyxl import Workbook

    workbook = Workbook()
    previous = workbook.active
    previous.title = "Tháng 9"
    headers = ["STT", "NGÀY AIR", "LINK AIR", "TÊN KÊNH", "LƯỢT XEM", "TIM", "BÌNH LUẬN", "LƯỢT LƯU", "CHIA SẺ", "Đối tác"]
    previous.append(headers)
    for number in range(1, 6):
        previous.append([number, f"0{number}/09/2026", f"https://www.tiktok.com/@demo/video/73900000000000000{number:02d}",
                         f"Kênh Demo {number:02d}", 12000 * number, 900 * number, 40 * number, 30 * number, 12 * number,
                         PARTNERS[number % len(PARTNERS)]])
    current = workbook.create_sheet(SHEET)
    current.append(headers)
    for number in range(1, LINK_COUNT + 1):
        current.append([number, f"{number:02d}/10/2026", tiktok_link(number), f"Kênh Demo {number:02d}",
                        None, None, None, None, None, PARTNERS[number % len(PARTNERS)]])
    workbook.save(data_dir / TIKTOK_FILE)

    threads = Workbook()
    sheet = threads.active
    sheet.title = SHEET
    sheet.append(["Ngày", "Link", "Tên Kênh", "LƯỢT XEM", "TIM", "BÌNH LUẬN", "REPOST", "CHIA SẺ", "Đối tác"])
    for number in range(1, 9):
        sheet.append([f"{number:02d}/10/2026", f"https://www.threads.com/@demo/post/DEMO{number:04d}", f"demo_{number:02d}",
                      None, None, None, None, None, PARTNERS[number % len(PARTNERS)]])
    threads.save(data_dir / THREADS_FILE)


def store_fake_tiktok_cookies(data_dir, key):
    """Saved straight into the encrypted vault: importing would verify against TikTok."""
    from riviu.platforms.threads_session import EncryptedFileStore
    from riviu.platforms.tiktok_session import TIKTOK_STORE, normalize_tiktok_cookies

    EncryptedFileStore(data_dir, key, spec=TIKTOK_STORE).save(normalize_tiktok_cookies(FAKE_TIKTOK_COOKIES))


def install_fake_tiktok_fetch():
    from riviu.platforms.tiktok import statuses, workers
    from riviu.platforms.tiktok.parsing import empty_metrics

    def fake_scrape(url, retries=0, **_):
        time.sleep(FAKE_FETCH_SECONDS)
        number = int(url.rstrip("/")[-2:])
        if number == FAILING_LINK:
            return empty_metrics(), "", "Error: HTTP 403", retries + 1, ""
        if number == NO_STATS_LINK:
            return empty_metrics(), f"Kênh Demo {number:02d}", statuses.STATUS_TIKTOK_NO_STATS, 1, url
        data = {"Views": 18500 * number + 731, "Likes": 1210 * number + 17, "Comments": 37 * number + 4,
                "Saves": 52 * number + 9, "Shares": 21 * number + 3}
        return data, f"Kênh Demo {number:02d}", "Success", 1, url

    workers.scrape_link_with_retries_request = fake_scrape


def start_server(port):
    import uvicorn
    from riviu.app import app

    server = uvicorn.Server(uvicorn.Config(app, host="127.0.0.1", port=port, log_level="warning"))
    thread = threading.Thread(target=server.run, daemon=True)
    thread.start()
    deadline = time.time() + 30
    while not server.started:
        if time.time() > deadline:
            raise RuntimeError("server did not start")
        time.sleep(0.1)
    return server, thread


CALLOUT_JS = """
(boxes) => {
  for (const node of document.querySelectorAll('.guide-shot-callout')) node.remove();
  for (const [index, box] of boxes.entries()) {
    const frame = document.createElement('div');
    frame.className = 'guide-shot-callout';
    Object.assign(frame.style, {position: 'fixed', left: `${box.x - 4}px`, top: `${box.y - 4}px`,
      width: `${box.width + 8}px`, height: `${box.height + 8}px`, border: '3px solid #e11d48',
      borderRadius: '9px', zIndex: 2147483646, pointerEvents: 'none', boxSizing: 'border-box'});
    const badge = document.createElement('div');
    badge.className = 'guide-shot-callout';
    badge.textContent = String(index + 1);
    Object.assign(badge.style, {position: 'fixed', left: `${Math.max(2, box.x - 15)}px`, top: `${Math.max(2, box.y - 15)}px`,
      width: '24px', height: '24px', borderRadius: '50%', background: '#e11d48', color: '#fff',
      font: '700 13px/24px Inter, Arial, sans-serif', textAlign: 'center', zIndex: 2147483647,
      pointerEvents: 'none', boxShadow: '0 2px 6px rgba(0,0,0,.25)'});
    document.body.append(frame, badge);
  }
}
"""


def save_png(raw, path):
    from PIL import Image

    image = Image.open(io.BytesIO(raw)).convert("RGB")
    for colors in (256, 128, 64):
        quantized = image.quantize(colors=colors, method=Image.Quantize.MEDIANCUT, dither=Image.Dither.NONE)
        buffer = io.BytesIO()
        quantized.save(buffer, format="PNG", optimize=True)
        if buffer.tell() <= MAX_BYTES:
            break
    path.write_bytes(buffer.getvalue())
    return buffer.tell()


class Shooter:
    def __init__(self, page):
        self.page = page
        self.sizes = {}

    def shot(self, name, *targets):
        boxes = []
        for target in targets:
            locator = self.page.locator(target) if isinstance(target, str) else target
            box = locator.first.bounding_box()
            if box is None:
                raise RuntimeError(f"{name}: callout target not visible: {target}")
            boxes.append(box)
        self.page.evaluate(CALLOUT_JS, boxes)
        self.page.wait_for_timeout(150)
        size = save_png(self.page.screenshot(), OUTPUT_DIR / f"{name}.png")
        self.page.evaluate(CALLOUT_JS, [])
        self.sizes[name] = size
        print(f"{name}.png  {size / 1024:.0f} KB")


def run(base_url, data_dir):
    from playwright.sync_api import sync_playwright

    with sync_playwright() as playwright:
        browser = playwright.chromium.launch()
        context = browser.new_context(viewport=VIEWPORT, locale="vi-VN", device_scale_factor=1)
        page = context.new_page()
        shooter = Shooter(page)
        settle = lambda ms=600: page.wait_for_timeout(ms)

        # 1. Login (server mode: `/` is the public viewer; the admin page is /admin).
        page.goto(base_url + "/admin")
        page.wait_for_url("**/login?next=/admin")
        page.fill("#password", PASSWORD)
        settle(300)
        shooter.shot("01-dang-nhap", "#password", "button[type=submit]")
        page.click("button[type=submit]")
        page.wait_for_selector("#guideButton")
        settle(1500)

        # 2. Platform bar and the source drawer button.
        page.click("#platformTikTok")
        settle()
        shooter.shot("02-chon-nen-tang", "#platformTikTok", "#sourceDrawerButton")

        # 3. Source drawer: pick the synthetic workbook, then point at link / sync / upload.
        page.click("#sourceDrawerButton")
        settle()
        page.select_option("#excelFileSelect", label=page.locator("#excelFileSelect option", has_text="Demo TikTok").first.inner_text())
        settle(1500)
        page.fill("#googleSheetUrlInput", "https://docs.google.com/spreadsheets/d/DEMO-SHEET-ID/edit")
        shooter.shot("03-nap-nguon", "#googleSheetUrlInput", "#syncSheetBtn", "#uploadArea")
        page.keyboard.press("Escape")
        settle()

        # 4. Scan sheet.
        page.select_option("#scanSheetSelect", SHEET)
        settle(800)
        sheet_tab = page.locator(".sheet-tab", has_text=SHEET).first
        sheet_tab.click()
        settle(1500)
        shooter.shot("04-chon-sheet", "#scanSheetSelect", sheet_tab)

        # 5. Mode, workers and proxy.
        page.select_option("#scrapeModeSelect", "request")
        page.select_option("#workerCountSelect", "1")
        shooter.shot("05-che-do-quet", "#scrapeModeSelect", "#workerCountSelect", "#proxyCard", "#tiktokCookieCard")

        # 5b. Cookie TikTok (fake saved cookies; the switch stays off for the scan below).
        page.click("#tiktokCookieBtn")
        page.wait_for_selector("#tiktokCookieModal.active")
        settle(800)
        shooter.shot("05b-cookie-tiktok", "#tiktokCookieWarning", "#tiktokCookieModal .threads-cookie-tabs", "#tiktokCookieImport")
        page.evaluate("closeTikTokCookieModal()")
        settle()

        # 6-7. A scan through the real runner and WebSocket, with the fake fetch.
        page.click("#startBtn")
        page.click(".workspace-tab[data-workspace-tab=live]")
        page.wait_for_function("document.querySelectorAll('#dataFeed tr').length >= 4", timeout=30000)
        settle(300)
        shooter.shot("06-dang-quet", "#startBtn", ".progress-section", ".workspace-tab[data-workspace-tab=live]")
        page.click(".workspace-tab[data-workspace-tab=errors]")
        page.wait_for_function("!document.querySelector('#failedLinksBody .workspace-empty')", timeout=30000)
        settle(300)
        shooter.shot("07-huy-quet", "#cancelBtn", ".workspace-tab[data-workspace-tab=errors]")
        page.wait_for_function("document.getElementById('cancelBtn').disabled", timeout=90000)
        settle(2500)

        # 8. Results in the scanned sheet.
        page.click(".workspace-tab[data-workspace-tab=sheet]")
        settle(800)
        page.locator(".sheet-tab", has_text=SHEET).first.click()
        settle(1500)
        shooter.shot("08-ket-qua", ".workspace-tab[data-workspace-tab=sheet]", ".workspace-action[data-action-for=sheet][title='Tải sheet']")

        # 9. The "Tổng kết" tab.
        summary_tab = page.locator(".sheet-tab", has_text="Tổng kết").first
        summary_tab.click()
        settle(1500)
        shooter.shot("09-tong-ket", summary_tab)

        # 10. Partner reports.
        shooter.shot("10a-mo-doi-tac", "button[title='Đối tác và báo cáo']")
        page.click("button[title='Đối tác và báo cáo']")
        page.wait_for_selector("#partnerList .partner-checkbox")
        settle(800)
        checkboxes = page.locator("#partnerList .partner-checkbox")
        checkboxes.nth(0).check()
        checkboxes.nth(1).check()
        settle(500)
        shooter.shot("10-xuat-bao-cao", "#partnerList", "#exportReportBtn")
        page.keyboard.press("Escape")
        settle()

        # 11. Google: service-account upload (recommended), OAuth upload and login status.
        page.click("#sourceDrawerButton")
        settle()
        shooter.shot("11-google", "#googleServiceAccountBtn", "#googleOauthBtn", "#googleLoginBtn")
        from riviu import google_sheets_sync

        # A FAKE service-account key (encrypted at rest, as in server mode); nothing calls Google.
        google_sheets_sync.save_service_account(str(data_dir), FAKE_SERVICE_ACCOUNT)
        page.evaluate("refreshGoogleOauthStatus()")
        page.wait_for_selector("#googleServiceAccountBox:not([hidden])")
        settle(500)
        shooter.shot("11b-tai-khoan-dich-vu", "#googleServiceAccountEmail", "#googleServiceAccountCopy", "#pushGoogleBtn")
        # A synthetic Desktop OAuth client enables the login button; the consent URL is built locally.

        google_sheets_sync.save_oauth_client(str(data_dir), {"installed": {
            "client_id": "demo.apps.googleusercontent.com", "client_secret": "demo",
            "auth_uri": "https://accounts.google.com/o/oauth2/auth", "token_uri": "https://oauth2.googleapis.com/token",
            "redirect_uris": ["http://localhost"],
        }})
        page.evaluate("refreshGoogleOauthStatus()")
        page.wait_for_function("!document.getElementById('googleLoginBtn').disabled")
        page.click("#googleLoginBtn")
        page.wait_for_selector("#googlePasteModal.active")
        page.fill("#googlePasteInput", "http://localhost:51234/?state=…&code=…&scope=…")
        settle(400)
        shooter.shot("12-google-dan-link", "#googlePasteAuthLink", "#googlePasteInput", "#googlePasteSubmit")
        page.evaluate("closeGooglePasteModal()")
        page.evaluate("closeCompactDrawers()")
        settle()

        # 12. Threads.
        page.click("#platformThreads")
        settle(800)
        page.click("#sourceDrawerButton")
        settle()
        page.select_option("#excelFileSelect", label=page.locator("#excelFileSelect option", has_text="Demo Threads").first.inner_text())
        settle(1500)
        page.evaluate("closeCompactDrawers()")
        settle(1200)
        shooter.shot("13-threads", "#platformThreads", "#threadsCookieCard")
        page.click("#threadsCookieBtn")
        page.wait_for_selector("#threadsCookieModal.active")
        settle(800)
        shooter.shot("14-cookie-threads", ".threads-cookie-tabs", "#threadsCookieImport")
        page.evaluate("closeThreadsCookieModal()")
        settle()

        # 13. Logout.
        shooter.shot("15-dang-xuat", "#logoutButton")
        browser.close()
        return shooter.sizes


def main():
    workdir = Path(tempfile.mkdtemp(prefix="riviu-guide-"))
    data_dir = workdir / "data-root"
    (data_dir / "data").mkdir(parents=True)
    key_file = workdir / "secret.key"
    secret_key = os.urandom(32)
    key_file.write_text(base64.b64encode(secret_key).decode())
    if os.name == "posix":
        key_file.chmod(0o600)
    with socket.socket() as probe:
        probe.bind(("127.0.0.1", 0))
        port = probe.getsockname()[1]

    from riviu import server_mode

    os.environ.update({
        "RIVIU_SERVER_MODE": "1",
        "RIVIU_DATA_DIR": str(data_dir),
        "RIVIU_PUBLIC_ORIGIN": "https://113-161-254-76.sslip.io",
        "RIVIU_PASSWORD_HASH": server_mode.hash_password(PASSWORD, n=2 ** 14),
        "RIVIU_SECRET_KEY_FILE": str(key_file),
        "RIVIU_PORT": str(port),
    })
    block_python_network()
    write_workbooks(data_dir / "data")
    store_fake_tiktok_cookies(data_dir, secret_key)
    install_fake_tiktok_fetch()
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    server, thread = start_server(port)
    try:
        sizes = run(f"http://localhost:{port}", data_dir)
    finally:
        server.should_exit = True
        thread.join(timeout=15)
        shutil.rmtree(workdir, ignore_errors=True)
    too_big = {name: size for name, size in sizes.items() if size > MAX_BYTES}
    if too_big:
        print(f"Over {MAX_BYTES // 1024} KB: {too_big}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.path.insert(0, str(ROOT))
    raise SystemExit(main())
