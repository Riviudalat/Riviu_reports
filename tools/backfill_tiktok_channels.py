"""Fill empty TikTok channel names in Riviu Data from V1 workbooks (one-off, rerun after re-imports).

The Sheet -> riviu_data load never carried TÊN KÊNH for TikTok, and Riviu Data's importer
ignores it on purpose (bao_cao belongs to Riviu Reports). This copies the best channel name
per link from the given workbooks into bao_cao.tiktok_chi_so, only where ten_kenh is empty,
never over a locked row (khoa_chi_so), and inserts the chi_so row for posts that have none.

Run inside the V2 container (RIVIU_DATABASE_URL_FILE points at the riviu_scanner DSN):
    python tools/backfill_tiktok_channels.py /tmp/wb/*.xlsx            # dry run: counts only
    python tools/backfill_tiktok_channels.py --apply /tmp/wb/*.xlsx
Workbooks earlier on the command line win when two carry the same link with equal quality.
"""

from __future__ import annotations

import argparse
import asyncio
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import load_workbook  # noqa: E402

from riviu.platforms.tiktok.channels import channel_name_quality  # noqa: E402
from riviu.store import config as store_config  # noqa: E402
from riviu.store.postgres import connect  # noqa: E402
from riviu.workbook.columns import worksheet_find_column_index, worksheet_find_link_column_index, worksheet_headers  # noqa: E402
from riviu.workbook.constants import COLUMN_ALIASES  # noqa: E402
from riviu.workbook.links import normalize_tiktok_url  # noqa: E402
from riviu.workbook.text import clean_text  # noqa: E402

UPSERT_SQL = (
    "INSERT INTO bao_cao.tiktok_chi_so AS c (id, ten_kenh) "
    "SELECT u.id, u.ten_kenh FROM unnest($1::bigint[], $2::text[], $3::text[]) AS u(id, link_seen, ten_kenh) "
    "WHERE EXISTS (SELECT 1 FROM public.v_tiktok v WHERE v.id = u.id AND v.link_air = u.link_seen) "
    "ON CONFLICT (id) DO UPDATE SET ten_kenh = EXCLUDED.ten_kenh "
    "WHERE NOT c.khoa_chi_so AND coalesce(c.ten_kenh, '') = '' "
    "RETURNING c.id"
)


def workbook_channels(paths):
    """{normalized link: best channel name} over every sheet of the given workbooks."""
    best = {}
    for path in paths:
        workbook = load_workbook(path, data_only=True)
        try:
            for worksheet in workbook.worksheets:
                headers = worksheet_headers(worksheet)
                link_index = worksheet_find_link_column_index(worksheet, headers=headers)
                channel_index = worksheet_find_column_index(worksheet, COLUMN_ALIASES["TÊN KÊNH"], headers=headers)
                if not link_index or not channel_index:
                    continue
                for row in worksheet.iter_rows(min_row=2, values_only=True):
                    raw_link = row[link_index - 1] if len(row) >= link_index else None
                    link = normalize_tiktok_url(clean_text(raw_link))
                    channel = clean_text(row[channel_index - 1] if len(row) >= channel_index else "")
                    quality = channel_name_quality(channel, link)
                    if link and quality > best.get(link, ("", 0))[1]:
                        best[link] = (channel, quality)
        finally:
            workbook.close()
    return {link: channel for link, (channel, _) in best.items()}


async def backfill(paths, apply):
    channels = workbook_channels(paths)
    print(f"workbook links with a channel name: {len(channels)}")
    conn = await connect(store_config.load_config().database_url)
    try:
        records = await conn.fetch(
            "SELECT v.id, v.link_air FROM public.v_tiktok v "
            "LEFT JOIN bao_cao.tiktok_chi_so c ON c.id = v.id "
            "WHERE coalesce(c.ten_kenh, '') = '' AND NOT coalesce(c.khoa_chi_so, false)"
        )
        ids, links, names = [], [], []
        for record in records:
            name = channels.get(normalize_tiktok_url(clean_text(record["link_air"])))
            if name:
                ids.append(record["id"])
                links.append(record["link_air"])
                names.append(name)
        print(f"posts with an empty channel: {len(records)}; matched in workbooks: {len(ids)}")
        if not apply:
            print("dry run: nothing written (pass --apply)")
            return
        async with conn.transaction():
            written = await conn.fetch(UPSERT_SQL, ids, links, names)
        print(f"rows written: {len(written)}")
    finally:
        await conn.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--apply", action="store_true", help="write to the database (default: dry run)")
    parser.add_argument("workbooks", nargs="+", type=Path)
    args = parser.parse_args()
    asyncio.run(backfill(args.workbooks, args.apply))


if __name__ == "__main__":
    main()
