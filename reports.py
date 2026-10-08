"""Partner report workbooks, export payloads and Google push rows.

Pure workbook/report building: no FastAPI, scan state, locks or WebSocket code.
`app.py` owns the routes that call into this module; this module must not import `app`.
"""
import io
import os
import re
import shutil
import sys
import tempfile
import zipfile
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import coordinate_from_string
from openpyxl.utils.units import pixels_to_EMU

from workbook_utils import (
    LAST_UPDATE_COLUMN,
    SINGLE_LINK_FILL_COLOR,
    VIDEO_LINK_FILL_COLOR,
    build_workbook_rows,
    clean_text,
    find_data_sheet_names,
    format_display_datetime,
    format_filename_datetime,
    format_metric,
    get_platform,
    is_exportable_report_row,
    is_failed_channel_name,
    list_workbook_partners,
    metric_total,
    partner_dedup_key,
    set_cell_literal,
    should_highlight_video_link,
)


# Same resolution as app.application_resource_dir(): bundled resources when frozen, else the repository.
RESOURCE_DIR = os.path.abspath(getattr(sys, "_MEIPASS", os.path.dirname(os.path.abspath(__file__))))
LOGO_PATH = os.path.join(RESOURCE_DIR, "logo.png")


def safe_report_name(name):
    filename = re.sub(r'[\\/:*?"<>|]+', "-", clean_text(name))
    filename = re.sub(r"\s+", " ", filename).strip(" .")
    return filename[:90] if filename else "doi_tac"


def spreadsheet_text(value):
    return clean_text(value)


def spreadsheet_date_text(value):
    if pd.isna(value):
        return ""
    if hasattr(value, "to_pydatetime"):
        return spreadsheet_text(value.to_pydatetime().strftime("%d/%m/%Y"))
    if isinstance(value, datetime):
        return spreadsheet_text(value.strftime("%d/%m/%Y"))
    parsed = parse_report_date(value)
    if not pd.isna(parsed):
        return spreadsheet_text(parsed.to_pydatetime().strftime("%d/%m/%Y"))
    text = clean_text(value)
    if text.endswith(" 00:00:00"):
        text = text[:10]
    return spreadsheet_text(text)


def parse_report_date(value):
    """Parse report dates as Vietnamese day/month/year values."""
    if pd.isna(value):
        return pd.NaT
    if hasattr(value, "to_pydatetime") or isinstance(value, datetime):
        return value
    text = clean_text(value)
    if not text:
        return pd.NaT
    # Keep ISO timestamps unambiguous; all other slash/dash dates are D/M/Y.
    if re.match(r"^\d{4}[-/]\d{1,2}[-/]\d{1,2}", text):
        return pd.to_datetime(text, errors="coerce")
    return pd.to_datetime(text, errors="coerce", dayfirst=True)


def spreadsheet_hyperlink(value):
    return clean_text(value)


def excel_column_name(index):
    name = ""
    while index > 0:
        index, remainder = divmod(index - 1, 26)
        name = chr(65 + remainder) + name
    return name


def column_width_to_pixels(width):
    if not width:
        width = 8.43
    return int(width * 7 + 5)


def row_height_to_pixels(height):
    if not height:
        height = 15
    return int(height * 96 / 72)


def add_centered_image_to_cell(ws, cell_address, image_path, max_height_px=36):
    if not image_path or not os.path.exists(image_path):
        return False
    try:
        img = ExcelImage(image_path)
    except Exception:
        return False

    col_letter, row_number = coordinate_from_string(cell_address)
    col_idx = ord(col_letter.upper()) - ord("A")
    row_idx = row_number - 1

    if img.height > max_height_px:
        scale = max_height_px / img.height
        img.height = max_height_px
        img.width = int(img.width * scale)

    cell_width_px = column_width_to_pixels(ws.column_dimensions[col_letter].width)
    max_width_px = max(cell_width_px - 4, 24)
    if img.width > max_width_px:
        scale = max_width_px / img.width
        img.width = max_width_px
        img.height = int(img.height * scale)

    cell_height_px = row_height_to_pixels(ws.row_dimensions[row_number].height)
    col_off = pixels_to_EMU(max((cell_width_px - img.width) / 2, 0))
    row_off = pixels_to_EMU(max((cell_height_px - img.height) / 2, 0))
    img.anchor = OneCellAnchor(
        _from=AnchorMarker(col=col_idx, colOff=col_off, row=row_idx, rowOff=row_off),
        ext=XDRPositiveSize2D(pixels_to_EMU(img.width), pixels_to_EMU(img.height)),
    )
    ws.add_image(img)
    return True


REPORT_COLUMN_WIDTHS = {
    "NGÀY AIR": 14, "TÊN KÊNH": 24, "LINK AIR": 72, "LƯỢT XEM": 14, "TIM": 12,
    "BÌNH LUẬN": 14, "LƯỢT LƯU": 14, "REPOST": 14, "CHIA SẺ": 12,
}
REPORT_TEXT_COLUMNS = {"LINK AIR", "TÊN KÊNH"}
REPORT_ACCENT = "FF6B00"
REPORT_ACCENT_TEXT = "9A3412"
REPORT_TOTAL_FILL = "FFF7ED"


def report_cell_value(header, value):
    if header == "NGÀY AIR":
        if pd.isna(value):
            return ""
        if hasattr(value, "to_pydatetime"):
            return value.to_pydatetime()
        parsed_date = parse_report_date(value)
        return parsed_date.to_pydatetime() if not pd.isna(parsed_date) else clean_text(value)
    if header in REPORT_TEXT_COLUMNS:
        return clean_text(value)
    return format_metric(value)


def report_row_fill(spec, source_row):
    partners = source_row.get("partners")
    partners = partners if isinstance(partners, (list, tuple)) else []
    if spec.highlights_video_links and should_highlight_video_link(
        clean_text(source_row.get("LINK AIR", "")),
        likes=source_row.get("TIM", 0),
        shares=source_row.get("CHIA SẺ", 0),
        resolved_url=source_row.get("_resolvedUrl", ""),
        resolved_source_url=source_row.get("_resolvedSourceUrl", ""),
        scan_status=source_row.get("_scanStatus", ""),
    ):
        return PatternFill("solid", fgColor=VIDEO_LINK_FILL_COLOR)
    if len(partners) == 1:
        return PatternFill("solid", fgColor=SINGLE_LINK_FILL_COLOR)
    return None


def write_report_banner(ws, spec, partner, link_count, columns):
    ws.row_dimensions[1].height = 42
    ws.row_dimensions[2].height = 26
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    add_centered_image_to_cell(ws, "A1", LOGO_PATH)

    ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=len(columns))
    title_cell = ws["B1"]
    title_cell.value = f"BÁO CÁO {spec.label.upper()} - ĐỐI TÁC: {partner}"
    title_cell.fill = PatternFill("solid", fgColor=REPORT_ACCENT)
    title_cell.font = Font(color="FFFFFF", bold=True, size=14)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")

    # Row 2: metadata across all but the last column, platform icon in the last one.
    last_column = get_column_letter(len(columns))
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(columns) - 1)
    ws.column_dimensions[last_column].width = REPORT_COLUMN_WIDTHS.get(columns[-1], 12)
    platform_icon = os.path.join(RESOURCE_DIR, "static", "platform-icons", f"{spec.key}.png")
    add_centered_image_to_cell(ws, f"{last_column}2", platform_icon, max_height_px=22)
    ws["A2"] = f"Tổng link: {link_count} • Ngày cập nhật: {format_display_datetime()}"
    ws["A2"].font = Font(color=REPORT_ACCENT_TEXT, italic=True, bold=True)
    ws["A2"].alignment = Alignment(horizontal="center")


def write_report_total(ws, spec, frame, columns, total_row, border):
    total_fill = PatternFill("solid", fgColor=REPORT_TOTAL_FILL)
    ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=3)
    for col_index in range(1, len(columns) + 1):
        cell = ws.cell(row=total_row, column=col_index)
        cell.fill = total_fill
        cell.border = border
    total_label = ws.cell(row=total_row, column=1, value="TỔNG")
    total_label.font = Font(bold=True, color=REPORT_ACCENT_TEXT)
    total_label.alignment = Alignment(horizontal="center", vertical="center")
    for header in spec.metric_columns:
        if header not in frame.columns:
            continue
        total_value = metric_total(frame[header])
        if total_value == "":
            continue
        cell = ws.cell(row=total_row, column=columns.index(header) + 1, value=total_value)
        cell.font = Font(bold=True, color=REPORT_ACCENT_TEXT)
        cell.number_format = "#,##0"
        cell.alignment = Alignment(horizontal="right")


def build_partner_report(partner, rows, *, apply_min_views=True, min_views=100, platform="tiktok"):
    spec = get_platform(platform)
    columns = list(spec.report_columns)
    valid_rows = [
        row for row in rows
        if is_exportable_report_row(row, apply_min_views=apply_min_views, min_views=min_views)
    ]
    frame = pd.DataFrame(valid_rows)
    wb = Workbook()
    ws = wb.active
    ws.title = "Báo cáo"
    table_header_row = 3
    data_start_row = table_header_row + 1
    thin = Side(style="thin", color="D8DEE9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    write_report_banner(ws, spec, partner, len(frame), columns)
    for col_index, header in enumerate(columns, start=1):
        cell = ws.cell(row=table_header_row, column=col_index, value=header)
        cell.fill = PatternFill("solid", fgColor=REPORT_ACCENT)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    for row_index, (_, source_row) in enumerate(frame.iterrows(), start=data_start_row):
        for col_index, header in enumerate(columns, start=1):
            value = report_cell_value(header, source_row.get(header, ""))
            cell = set_cell_literal(ws.cell(row=row_index, column=col_index), value)
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=(header in {"LINK AIR", "TÊN KÊNH"}))
            if header == "LINK AIR" and isinstance(value, str) and value.startswith("http"):
                cell.hyperlink = value
                cell.style = "Hyperlink"
            elif header == "NGÀY AIR":
                cell.number_format = "dd/mm/yyyy"
                cell.alignment = Alignment(horizontal="center", vertical="top")
            elif header not in REPORT_TEXT_COLUMNS:
                cell.number_format = "#,##0"
                cell.alignment = Alignment(horizontal="right", vertical="top")
        row_fill = report_row_fill(spec, source_row)
        if row_fill is not None:
            for col_index in range(1, len(columns) + 1):
                ws.cell(row=row_index, column=col_index).fill = row_fill

    total_row = len(frame) + data_start_row if len(frame) else None
    if total_row:
        write_report_total(ws, spec, frame, columns, total_row, border)

    ws.freeze_panes = f"A{data_start_row}"
    filter_last_row = (total_row - 1) if total_row else max(data_start_row, table_header_row)
    ws.auto_filter.ref = f"A{table_header_row}:{get_column_letter(len(columns))}{filter_last_row}"
    for index, header in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(index)].width = REPORT_COLUMN_WIDTHS.get(header, 14)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def build_google_push_rows(rows, *, platform="tiktok"):
    spec = get_platform(platform)
    normalized_rows = []
    max_partner_columns = 1
    for row in rows:
        partner_names = row.get("partners", [])
        if not isinstance(partner_names, list):
            partner_names = []
        partner_names = [clean_text(name) for name in partner_names if clean_text(name)]
        max_partner_columns = max(max_partner_columns, len(partner_names))
        normalized_rows.append((row, partner_names))

    headers = ["Stt", "Ngày", "Link", "Tên Kênh", *spec.metric_columns]
    headers.extend(["Đối tác" if index == 0 else f"Đối tác {index + 1}" for index in range(max_partner_columns)])
    headers.append(LAST_UPDATE_COLUMN)

    def metric_cell(row, metric):
        # Unknown metrics stay blank in Google too; only confirmed numbers are written.
        value = format_metric(row.get(metric, ""))
        return "" if value is None else value

    values = [headers]
    for index, (row, partner_names) in enumerate(normalized_rows, start=1):
        padded_partners = partner_names + [""] * (max_partner_columns - len(partner_names))
        channel_name = clean_text(row.get("TÊN KÊNH", ""))
        if is_failed_channel_name(channel_name):
            channel_name = "Lỗi"
        values.append([
            index,
            spreadsheet_date_text(row.get("NGÀY AIR", "")),
            spreadsheet_hyperlink(row.get("LINK AIR", "")),
            channel_name,
            *[metric_cell(row, metric) for metric in spec.metric_columns],
            *padded_partners,
            clean_text(row.get(LAST_UPDATE_COLUMN, "")),
        ])
    if len(values) > 1:
        data_last_row = len(values)
        total_row = ["", "", "TỔNG", ""]
        for offset in range(len(spec.metric_columns)):
            column_index = 5 + offset
            column_name = excel_column_name(column_index)
            metric_values = [row[column_index - 1] for row in values[1:]]
            # Same rule as the Excel report: sum the known values, blank only when none is known.
            known = metric_total(metric_values) != ""
            total_row.append(f"=SUM({column_name}2:{column_name}{data_last_row})" if known else "")
        total_row.extend([""] * (len(headers) - len(total_row)))
        values.append(total_row)
    return values


def build_export_payload(target_path, selected_partners, apply_min_views, min_views, requested_sheet_name="", platform="tiktok"):
    # One immutable copy feeds discovery and every partner report in this export.
    # Atomic scan saves may replace the source while this function is running.
    with tempfile.TemporaryDirectory(prefix="riviu-export-") as directory:
        snapshot = os.path.join(directory, os.path.basename(target_path))
        shutil.copyfile(target_path, snapshot)
        return _build_export_payload(snapshot, selected_partners, apply_min_views, min_views, requested_sheet_name, platform)


def _build_export_payload(target_path, selected_partners, apply_min_views, min_views, requested_sheet_name="", platform="tiktok"):
    spec = get_platform(platform)
    data_sheets = find_data_sheet_names(target_path)
    report_sheet = clean_text(requested_sheet_name) or (data_sheets[0] if data_sheets else "")
    if report_sheet and report_sheet not in data_sheets:
        raise ValueError(f"Sheet {report_sheet} không tồn tại trong file.")
    available_partners = {
        partner_dedup_key(name): name
        for name in list_workbook_partners(target_path, sheet_name=report_sheet, platform=spec.key)
    }
    partners = list(dict.fromkeys(
        available_partners[partner_dedup_key(name)]
        for name in selected_partners
        if clean_text(name) and partner_dedup_key(name) in available_partners
    ))
    if not partners:
        raise ValueError("Không tìm thấy đối tác đã chọn trong file")
    export_timestamp = format_filename_datetime()
    sheet_tag = safe_report_name(report_sheet) if report_sheet else ""
    name_tags = [tag for tag in (spec.file_tag, sheet_tag, export_timestamp) if tag]

    def report_bytes(partner):
        rows = build_workbook_rows(target_path, selected_partner=partner, sheet_name=report_sheet, platform=spec.key)
        return build_partner_report(partner, rows, apply_min_views=apply_min_views, min_views=min_views, platform=spec.key)

    if len(partners) == 1:
        partner = partners[0]
        return {
            "content": report_bytes(partner),
            "media_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "filename": f"{' '.join([safe_report_name(partner), *name_tags])}.xlsx",
        }

    archive = io.BytesIO()
    used_names = set()
    with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED) as zip_file:
        for partner in partners:
            base_name = " ".join([safe_report_name(partner), *name_tags])
            filename = f"{base_name}.xlsx"
            counter = 2
            while filename in used_names:
                filename = f"{base_name}_{counter}.xlsx"
                counter += 1
            used_names.add(filename)
            zip_file.writestr(filename, report_bytes(partner))
    return {
        "content": archive.getvalue(),
        "media_type": "application/zip",
        "filename": f"{' '.join(['bao_cao_doi_tac', *name_tags])}.zip",
    }
