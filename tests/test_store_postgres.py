"""Postgres backend (RIVIU_STORAGE=postgres) with the database faked: config, write rules,
workbook parity, runners on a store session and the /internal/scan contract.

The same contracts against a real Postgres 17 built from Riviu Data's migrations are in
tests/test_store_postgres_integration.py (opt-in).
"""

import asyncio
import datetime as dt
import hashlib

import openpyxl
import pytest
from fastapi.testclient import TestClient

from riviu import internal_scan, scan_state
from riviu.platforms import threads
from riviu.platforms import tiktok as scraper
from riviu.platforms.tiktok.workbook_io import collect_rows
from riviu.store import config as store_config
from riviu.store import postgres, rules
from riviu.workbook.partners import (
    doi_tac_from_partners,
    partners_from_doi_tac,
    row_partners_from_values,
)
from riviu.workbook.platforms import get_platform
from riviu.workbook.report_rows import build_workbook_rows
from riviu.workbook.text import format_metric
from tests.pg_fakes import Events, FakeConn, NoPlaywright

TOKEN = "t" * 40
NOW = dt.datetime(2026, 10, 5, 1, 30, tzinfo=dt.UTC)
TIKTOK_URL = "https://www.tiktok.com/@demo/video/7400000000000000001"
THREADS_URL = "https://www.threads.com/@miri_viu/post/DdRIHGWCURV"


# ---- configuration -----------------------------------------------------------------

def test_storage_defaults_to_workbook_and_postgres_needs_server_mode_and_scanner_role(tmp_path):
    assert store_config.load_config({}).backend == "workbook"
    with pytest.raises(store_config.StoreConfigError, match="server"):
        store_config.load_config({"RIVIU_STORAGE": "postgres", "RIVIU_DATABASE_URL": "postgresql://riviu_scanner@h/db"})
    server = {"RIVIU_STORAGE": "postgres", "RIVIU_SERVER_MODE": "1"}
    with pytest.raises(store_config.StoreConfigError, match="riviu_scanner"):
        store_config.load_config({**server, "RIVIU_DATABASE_URL": "postgresql://riviudata:x@h/db"})
    with pytest.raises(store_config.StoreConfigError, match="dual"):
        store_config.load_config({**server, "RIVIU_STORAGE": "dual"})
    dsn_file = tmp_path / "dsn"
    dsn_file.write_text("postgresql://riviu_scanner:pw@db:5432/riviu_data\n", encoding="utf-8")
    tokens = tmp_path / "tokens"
    digest = hashlib.sha256(b"u" * 40).hexdigest()
    tokens.write_text(f"# timer\n{TOKEN}\nsha256:{digest} report:read\n", encoding="utf-8")
    config = store_config.load_config({**server, "RIVIU_DATABASE_URL_FILE": str(dsn_file), "RIVIU_SCAN_TOKEN_FILE": str(tokens)})
    assert config.postgres and config.database_url.endswith("/riviu_data")
    assert [sorted(token.scopes) for token in config.scan_tokens] == [["scan:run"], ["report:read"]]
    assert "pw" not in repr(config)
    tokens.write_text("short\n", encoding="utf-8")
    with pytest.raises(store_config.StoreConfigError, match="32"):
        store_config.load_config({**server, "RIVIU_DATABASE_URL_FILE": str(dsn_file), "RIVIU_SCAN_TOKEN_FILE": str(tokens)})


def test_month_is_a_year_qualified_range():
    assert rules.parse_thang("2025-12") == (dt.date(2025, 12, 1), dt.date(2026, 1, 1))
    assert rules.parse_thang("2026-01") == (dt.date(2026, 1, 1), dt.date(2026, 2, 1))
    for bad in ("Tháng 10", "2026-13", "10/2026", ""):
        with pytest.raises(ValueError):
            rules.parse_thang(bad)
    # 2026-09-30 17:30 UTC is already October in Vietnam.
    assert rules.current_thang(dt.datetime(2026, 9, 30, 17, 30, tzinfo=dt.UTC)) == "2026-10"


# ---- write rules: NULL is unknown, 0 is a confirmed zero ----------------------------

METRICS_ALL = ("luot_xem", "tim", "binh_luan", "luot_luu", "chia_se")


@pytest.mark.parametrize(("status", "data", "expected_metrics", "touches_time"), [
    ("Success", {"Views": "0", "Likes": None, "Comments": "", "Saves": "1.234", "Shares": 7},
     {"luot_xem": 0, "tim": None, "binh_luan": None, "luot_luu": 1234, "chia_se": 7}, True),
    ("Ẩn số liệu: TikTok không trả số liệu", {"Views": "5"}, dict.fromkeys(METRICS_ALL), False),
    ("Error: HTTP 403", {"Views": "5"}, {}, False),
])
def test_tiktok_write_columns_follow_the_workbook_rules(status, data, expected_metrics, touches_time):
    columns = rules.tiktok_update_columns("Demo Shop", TIKTOK_URL, data, "Demo Shop", status, now=NOW)
    assert {key: columns[key] for key in METRICS_ALL if key in columns} == expected_metrics
    assert columns["trang_thai_quet"] == status
    assert ("cap_nhat_lan_cuoi" in columns) is touches_time
    assert ("link_nguon_phan_giai" in columns) is (status == "Success")
    # An unchanged channel is never rewritten (keeps a staff edit made during the scan).
    assert "ten_kenh" not in columns


def test_tiktok_failure_never_downgrades_a_usable_channel_and_success_upgrades_a_handle():
    failed = rules.tiktok_update_columns("@demo", TIKTOK_URL, {}, "Lỗi", "Error: HTTP 403", now=NOW)
    assert "ten_kenh" not in failed
    upgraded = rules.tiktok_update_columns("@demo", TIKTOK_URL, {"Views": "1"}, "Demo Shop", "Success", now=NOW)
    assert upgraded["ten_kenh"] == "Demo Shop"


@pytest.mark.parametrize(("result", "expected"), [
    # A missing value keeps the old one; a confirmed 0 is written as 0.
    ({"channel": "miri", "metrics": {"views": 0, "likes": 4, "comments": None, "reposts": None, "shares": None}, "error": ""},
     {"ten_kenh": "miri", "luot_xem": 0, "tim": 4, "trang_thai_quet": "Success"}),
    # The explicit view_query_unknown invalidation clears stale views.
    ({"channel": "", "metrics": {"views": None, "likes": 2, "comments": None, "reposts": None, "shares": None}, "error": "",
      "metricSources": {"views": "view_query_unknown"}},
     {"luot_xem": None, "tim": 2, "trang_thai_quet": "Success"}),
    # A failure writes only its status and keeps every number and the last-update time.
    ({"channel": "", "metrics": {"views": None, "likes": None, "comments": None, "reposts": None, "shares": None},
      "error": "Bài Threads bị giới hạn theo vị trí hiện tại"},
     {"trang_thai_quet": "Error: Bài Threads bị giới hạn theo vị trí hiện tại"}),
])
def test_threads_write_columns_follow_the_workbook_rules(result, expected):
    columns = rules.threads_update_columns(result, now=NOW)
    touched = {key: value for key, value in columns.items() if key != "cap_nhat_lan_cuoi"}
    assert touched == expected
    assert ("cap_nhat_lan_cuoi" in columns) is (result["error"] == "")


# ---- parity with the workbook path ---------------------------------------------------

PARTNER_CELLS = [
    ("Cafe A", None),
    ("DANH SÁCH\n- Cafe A\n2. Tiệm Bánh B", "Cafe\u200b A"),
    ("Cafe A, Quán C", None),  # one workbook partner with a comma: cannot be stored (doi_tac_from_partners refuses)
    ("  Bún   Bò\u00a0Huế ", "Cafe e\u0301"),  # NFD é and odd spaces
    (None, None),
]


def test_partner_lists_survive_the_doi_tac_round_trip():
    for cells in PARTNER_CELLS:
        partners = row_partners_from_values(cells)
        if any("," in name for name in partners):
            with pytest.raises(ValueError):
                doi_tac_from_partners(partners)
            continue
        assert partners_from_doi_tac(doi_tac_from_partners(partners)) == partners


def _tiktok_workbook(tmp_path, rows):
    path = tmp_path / "tiktok.xlsx"
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Data"
    sheet.append(["NGÀY AIR", "LINK AIR", "TÊN KÊNH", "LƯỢT XEM", "TIM", "BÌNH LUẬN", "LƯỢT LƯU", "CHIA SẺ",
                  "Cập nhật lần cuối", "Đối tác", "Đối tác 2", "__TTBD_SCAN_STATUS", "__TTBD_RESOLVED_URL", "__TTBD_SOURCE_URL"])
    for row in rows:
        sheet.append(row)
    book.save(path)
    book.close()
    return path


TIKTOK_ROWS = [
    [dt.datetime(2026, 10, 1), " www.tiktok.com/@demo/video/7400000000000000001 ", "Demo", 0, None, 5, 1, 2,
     "01/10/2026-08:30", "Cafe A", None, "Success", "https://www.tiktok.com/@demo/video/7400000000000000001",
     "https://www.tiktok.com/@demo/video/7400000000000000001"],
    [dt.datetime(2026, 10, 2), "https://vt.tiktok.com/ZSabc123/", "", None, None, None, None, None, None,
     "DANH SÁCH\n- Cafe A\n2. Tiệm Bánh B", "Cafe\u200b A", "", "https://www.tiktok.com/@x/video/7400000000000000009",
     "https://vt.tiktok.com/ZSabc123/"],
    [dt.datetime(2026, 10, 3), "https://www.tiktok.com/@demo", "Demo", 1, 1, 1, 1, 1, None, "Quán C", None, None, None, None],
    [dt.datetime(2026, 10, 4), "javascript:alert(1)", "x", None, None, None, None, None, None, "Quán C", None, None, None, None],
]


def _db_record(row_id, row):
    ngay, link, channel, views, likes, comments, saves, shares, updated, p1, p2, status, resolved, source = row
    return {
        "id": row_id, "ngay_air": ngay.date(), "link_air": link, "ten_kenh": channel or None,
        "doi_tac": doi_tac_from_partners(row_partners_from_values([p1, p2])),
        "luot_xem": views, "tim": likes, "binh_luan": comments, "luot_luu": saves, "chia_se": shares,
        "cap_nhat_lan_cuoi": dt.datetime.strptime(updated, "%d/%m/%Y-%H:%M").replace(tzinfo=rules.LOCAL_TZ) if updated else None,
        "trang_thai_quet": status or None, "link_phan_giai": resolved, "link_nguon_phan_giai": source,
    }


def test_postgres_rows_match_the_workbook_rows(tmp_path):
    path = _tiktok_workbook(tmp_path, TIKTOK_ROWS)
    records = [_db_record(index, row) for index, row in enumerate(TIKTOK_ROWS, start=1)]
    book = openpyxl.load_workbook(path)
    for selected in (None, ["cafe a"]):
        workbook_scan = collect_rows(book, selected_partners=selected, sheet_name="Data")
        store_scan = postgres.tiktok_scan_rows(records, "2026-10", {name.casefold() for name in selected or []})
        keys = ("url", "partners", "expected_media_id")
        assert [{key: row[key] for key in keys} for row in store_scan] == [{key: row[key] for key in keys} for row in workbook_scan]
    book.close()

    def comparable(row):
        return {key: (format_metric(value) if key in rules.ALL_REPORT_METRICS else value)
                for key, value in row.items() if key != "sheet_name"}

    workbook_report = build_workbook_rows(path, sheet_name="Data", platform="tiktok")
    store_report = [rules.report_row(get_platform("tiktok"), record, "2026-10") for record in records]
    assert [comparable(row) for row in store_report if row] == [comparable(row) for row in workbook_report]


def test_threads_rows_match_the_workbook_rows(tmp_path):
    cells = [
        (THREADS_URL + "?xmt=1", "Cafe A"),
        ("threads.net/@miri_viu/post/AbC_123/", "Cafe A\nTiệm B"),
        ("https://www.threads.com/@miri_viu", "Cafe A"),  # profile: invalid
        ("https://www.threads.com/@miri_viu/p/x", "Tiệm B"),  # malformed
        ("https://www.tiktok.com/@demo/video/1", "Cafe A"),  # other platform: skipped
    ]
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Data"
    sheet.append(["Link", "Đối tác"])
    for link, partner in cells:
        sheet.append([link, partner])
    records = [
        {"id": index, "link_air": link, "doi_tac": doi_tac_from_partners(row_partners_from_values([partner])), "ten_kenh": None}
        for index, (link, partner) in enumerate(cells, start=1)
    ]
    for selected in (None, ["tiệm b"]):
        invalid = []
        workbook_rows = threads.collect_threads_rows(book, "Data", selected, invalid)
        store_rows, store_invalid = postgres.threads_scan_rows(records, "2026-10", threads.selected_threads_partner_keys(selected))
        assert [(row["url"], row["partners"]) for row in store_rows] == [(row["url"], row["partners"]) for row in workbook_rows]
        assert [(item["kind"], item["display"]) for item in store_invalid] == [(item["kind"], item["display"]) for item in invalid]


# ---- runners on a store session (store faked at the connection) ----------------------

def _session(platform, records):
    store = postgres.PostgresStore(platform, "postgresql://riviu_scanner@fake/db", connector=None)
    conn = FakeConn(records)
    if platform == "tiktok":
        rows, invalid = postgres.tiktok_scan_rows(records, "2026-10"), []
    else:
        rows, invalid = postgres.threads_scan_rows(records, "2026-10")
    return postgres.PostgresScanSession(store, conn, "2026-10", rows, invalid_links=invalid, run_id=7, events=Events()), conn


def test_tiktok_scan_writes_through_the_session_with_guarded_updates(monkeypatch, tmp_path):
    records = [
        {"id": 1, "link_air": TIKTOK_URL, "doi_tac": "Cafe A", "ten_kenh": "Demo", "link_phan_giai": None, "link_nguon_phan_giai": None},
        {"id": 2, "link_air": TIKTOK_URL + "?x=1", "doi_tac": "Cafe B", "ten_kenh": None, "link_phan_giai": None, "link_nguon_phan_giai": None},
    ]
    session, conn = _session("tiktok", records)
    monkeypatch.setattr(scraper.runner, "async_playwright", NoPlaywright)
    monkeypatch.setattr(scraper.workers, "_run_request_scrape", lambda *_a, **_kw: (
        {"Views": "999", "Likes": "0", "Comments": None, "Saves": "2", "Shares": "3"}, "Demo", "Success", 1, TIKTOK_URL,
    ))
    monkeypatch.setattr(scraper.history, "append_scrape_history", lambda *_a: pytest.fail("no JSON history for a store scan"))
    asyncio.run(scraper.run_scraper("2026-10", Events(), worker_count=1, retries=0, base_dir=tmp_path, store_session=session))

    assert session.written == 2 and session.totals["success"] == 2
    written = {}
    for sql, args in conn.updates:
        assert "t.link_air = u.link_seen AND t.da_xoa_luc IS NULL AND NOT t.khoa_chi_so" in sql
        columns = [part.split(" = ")[0] for part in sql.split("SET ", 1)[1].split(" FROM", 1)[0].split(", ")]
        for index, row_id in enumerate(args[0]):
            written[row_id] = {"link_seen": args[1][index], **{column: args[2 + n][index] for n, column in enumerate(columns)}}
    assert written[1]["link_seen"] == TIKTOK_URL and written[2]["link_seen"] == TIKTOK_URL + "?x=1"
    for row in written.values():
        assert (row["luot_xem"], row["tim"], row["binh_luan"]) == (999, 0, None)
    # Row 1 already holds that channel: only row 2 names ten_kenh.
    assert "ten_kenh" not in written[1] and written[2]["ten_kenh"] == "Demo"
    assert any("set_config('app.user'" in sql and args == ("app:riviu-reports-scanner",) for sql, args in conn.calls)
    assert any("scan_run_item" in sql for sql, _ in conn.calls)
    assert not list(tmp_path.glob("*.json"))


def test_threads_scan_writes_through_the_session_and_keeps_the_diagnostic(monkeypatch, tmp_path):
    session, conn = _session("threads", [{"id": 5, "link_air": THREADS_URL, "doi_tac": "Cafe A", "ten_kenh": None}])
    monkeypatch.setattr(threads, "async_playwright", NoPlaywright)
    monkeypatch.setattr(threads, "fetch_threads_http", lambda _url: {
        "channel": "miri_viu", "metrics": {"views": 372, "likes": 0, "comments": 1, "reposts": 2, "shares": 3}, "error": "",
    })
    monkeypatch.setattr(threads, "append_threads_diagnostic", lambda *_a: pytest.fail("no JSON diagnostic for a store scan"))
    asyncio.run(threads.run_threads_scraper("2026-10", Events(), sheet_name="2026-10", mode="request", base_dir=str(tmp_path),
                                            store_session=session))
    assert session.written == 1
    assert session.diagnostic["phase"] == "completed" and session.diagnostic["success"] == 1
    [(sql, args)] = conn.updates
    assert sql.startswith("UPDATE public.threads") and "repost = u.repost" in sql


def test_a_row_edited_during_the_scan_is_skipped_and_logged():
    records = [{"id": 1, "link_air": THREADS_URL, "doi_tac": None, "ten_kenh": None}]
    session, conn = _session("threads", records)
    conn.reject_ids = {1}
    session.write_threads_result(1, {"channel": "x", "metrics": {"views": 1}, "error": ""})
    assert asyncio.run(session.flush()) == (True, None)
    assert session.written == 0 and session.skipped == [{"id": 1, "reason": "link đã đổi"}]
    assert any("đã bị sửa trong lúc quét" in line for line in session.events.logs)
    item_sql, item_args = next((sql, args) for sql, args in conn.calls if "scan_run_item" in sql)
    assert item_args[2] == [1] and item_args[3] == ["Bỏ qua: link đã đổi"]


def test_a_failed_flush_keeps_the_writes_for_the_next_one():
    session, conn = _session("threads", [{"id": 1, "link_air": THREADS_URL, "doi_tac": None, "ten_kenh": None}])
    session.write_threads_result(1, {"channel": "x", "metrics": {"views": 1}, "error": ""})

    async def broken(*_args):
        raise ConnectionError("db down")

    conn.fetch, original = broken, conn.fetch
    assert asyncio.run(session.flush()) == (False, "other")
    assert session.pending_count == 1
    conn.fetch = original
    assert asyncio.run(session.flush()) == (True, None) and session.written == 1


def test_only_scanner_columns_can_be_buffered():
    session, _conn = _session("threads", [{"id": 1, "link_air": THREADS_URL, "doi_tac": None, "ten_kenh": None}])
    with pytest.raises(postgres.StoreError):
        session._buffer(1, {"link_air": "x"}, "Success")


# ---- POST /internal/scan --------------------------------------------------------------

class FakeStore:
    def __init__(self, conn):
        self.conn = conn
        self.platform = "threads"

    async def connect(self):
        return self.conn

    async def open_scan(self, thang, *, run_id=None, events=None):
        session, _ = _session("threads", [{"id": 1, "link_air": THREADS_URL, "doi_tac": None, "ten_kenh": None}])
        session.run_id, session.events, session.conn = run_id, events, self.conn
        return session


@pytest.fixture
def internal_api(monkeypatch, tmp_path):
    from riviu import app as backend

    monkeypatch.setattr(store_config, "ACTIVE", store_config.StoreConfig(
        backend="postgres", database_url="postgresql://riviu_scanner@fake/db",
        scan_tokens=(store_config.ScanToken(store_config.token_digest(TOKEN), frozenset({"scan:run"})),
                     store_config.ScanToken(store_config.token_digest("r" * 40), frozenset({"report:read"}))),
    ))
    monkeypatch.setattr(scan_state, "manager", scan_state.ConnectionManager())
    monkeypatch.setattr(scan_state, "SCRAPE_TASK", None)
    monkeypatch.setattr(backend.server_context, "EXCEL_DIR", str(tmp_path))
    state = {"conn": FakeConn(), "scans": []}
    monkeypatch.setattr(internal_scan, "make_store", lambda _platform: FakeStore(state["conn"]))

    async def fake_scan(session, events):
        state["scans"].append(session.label)
        session.write_threads_result(1, {"channel": "x", "metrics": {"views": 0}, "error": ""})
        await session.flush()
        session.totals = {"success": 1}

    monkeypatch.setattr(internal_scan, "SCANNERS", {"tiktok": fake_scan, "threads": fake_scan})
    with TestClient(backend.app, base_url="http://riviu-reports:8000") as client:
        yield client, state


def _post(client, body, token=TOKEN, **headers):
    if token:
        headers["Authorization"] = f"Bearer {token}"
    return client.post("/internal/scan", json=body, headers=headers)


def test_internal_scan_auth(internal_api):
    client, state = internal_api
    body = {"platform": "threads", "thang": "2026-10", "trigger": "timer"}
    assert _post(client, body, token=None).status_code == 401
    assert _post(client, body, token="x" * 40).status_code == 401
    assert _post(client, body, token="r" * 40).status_code == 403
    assert _post(client, body, Origin="http://riviu-reports:8000").status_code == 403
    assert state["scans"] == []


def test_internal_scan_is_absent_in_workbook_mode(internal_api, monkeypatch):
    client, _ = internal_api
    monkeypatch.setattr(store_config, "ACTIVE", store_config.WORKBOOK)
    assert _post(client, {"platform": "threads"}).status_code == 404


@pytest.mark.parametrize("body", [
    {"platform": "facebook"}, {"platform": "threads", "thang": "Tháng 10"},
    {"platform": "threads", "trigger": "cron"}, {"platform": "threads", "scan_run_id": True}, {"platform": "threads", "x": 1},
])
def test_internal_scan_rejects_invalid_bodies(internal_api, body):
    client, state = internal_api
    assert _post(client, body).status_code == 422
    assert state["scans"] == []


def test_timer_scan_runs_to_the_end_and_manual_scan_returns_at_once(internal_api):
    client, state = internal_api
    response = _post(client, {"platform": "threads", "thang": "2026-10", "trigger": "timer"})
    assert response.status_code == 200
    summary = response.json()
    assert summary["phase"] == "hoan_tat" and summary["scan_run_id"] == 41 and summary["totals"]["written"] == 1
    final = [args for sql, args in state["conn"].calls if sql.startswith("UPDATE public.scan_run SET totals")]
    assert final and final[-1][2] == "hoan_tat" and final[-1][5] is True
    assert any("pg_advisory_unlock" in sql for sql, _ in state["conn"].calls) and state["conn"].closed

    state["conn"] = FakeConn()
    response = _post(client, {"platform": "threads", "thang": "2026-10", "started_by": "user:an", "scan_run_id": 9})
    assert response.status_code == 202 and response.json()["scan_run_id"] == 9
    for _ in range(100):
        if state["conn"].closed:
            break
        asyncio.run(asyncio.sleep(0.02))
    assert state["scans"] == ["2026-10", "2026-10"]
    reuse = next(args for sql, args in state["conn"].calls if sql.startswith("UPDATE public.scan_run SET phase"))
    assert reuse[:2] == (9, "threads")


def test_internal_scan_refuses_while_a_scan_runs(internal_api):
    client, state = internal_api
    state["conn"] = FakeConn(lock_ok=False)
    response = _post(client, {"platform": "threads", "thang": "2026-10", "trigger": "timer"})
    assert response.status_code == 409 and state["conn"].closed
    state["conn"] = FakeConn()
    scan_state.manager.running = True  # a scan started from the UI in this process
    assert _post(client, {"platform": "threads", "trigger": "timer"}).status_code == 409
    assert state["scans"] == []
