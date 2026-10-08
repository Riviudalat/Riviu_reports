import json
from datetime import datetime
from types import SimpleNamespace
from unittest.mock import Mock

import openpyxl
import pytest

from riviu import google_sheets_sync
from riviu.platforms.tiktok import build_result_sheet
from riviu.workbook_utils import (
    build_workbook_rows,
    clean_text,
    find_data_sheet_names,
    format_excel_sheet_datetime,
    google_sheet_original_title,
    read_sheet_preview,
    result_sheet_display_name,
)


def test_result_tab_is_valid_excel_title_and_reopens(tmp_path):
    book = openpyxl.Workbook()
    title = result_sheet_display_name("25-09-2026-10:30")
    book.create_sheet(title)
    build_result_sheet(book, [], "")
    assert format_excel_sheet_datetime(datetime(2026, 9, 25, 10, 30)) == "25-09-2026-10-30"
    path = tmp_path / "results.xlsx"
    book.save(path)
    reopened = openpyxl.load_workbook(path)
    assert title in reopened.sheetnames
    assert all(":" not in name for name in reopened.sheetnames)
    reopened.close()


def _mock_sheets_api(monkeypatch, titles):
    service = Mock()
    sheets = service.spreadsheets.return_value
    sheets.get.return_value.execute.return_value = {
        "sheets": [{"properties": {"title": title}} for title in titles]
    }

    def values_get(**kwargs):
        original = kwargs["range"][1:-6].replace("''", "'")
        request = Mock()
        request.execute.return_value = {"values": [["Original title", "LINK AIR"], [original, "https://www.threads.com/@demo/post/123"]]}
        return request

    sheets.values.return_value.get.side_effect = values_get
    monkeypatch.setattr(google_sheets_sync, "try_load_credentials", lambda *_: SimpleNamespace(valid=True))
    monkeypatch.setattr(google_sheets_sync, "build", lambda *_, **__: service)


def test_google_import_sanitizes_unique_titles_and_preserves_original_mapping(tmp_path, monkeypatch):
    titles = [
        "T9 25-09-2026-10:30",
        "Campaign/A:*?[]\\",
        "Campaign-A-----",
        "A" * 40,
        "A" * 39 + "B",
        "Report Seeding Threads T9 25-09-2026-10:30",
        "Report Seeding Threads T9 25-09-2026-10:31",
    ]
    observed_maps = []
    for index, order in enumerate((titles, list(reversed(titles)))):
        _mock_sheets_api(monkeypatch, order)
        path = tmp_path / f"import-{index}.xlsx"
        google_sheets_sync.download_google_sheet_authenticated(tmp_path, "spreadsheet", str(path))
        book = openpyxl.load_workbook(path)
        mapping = json.loads(book.custom_doc_props["Riviu.GoogleSheetTitles"].value)
        observed_maps.append(mapping)
        assert set(mapping) == set(titles)
        assert len({name.casefold() for name in book.sheetnames}) == len(titles)
        assert all(len(name) <= 31 and not any(char in name for char in ":\\/?*[]") for name in book.sheetnames)
        for original, local in mapping.items():
            assert book[local]["A2"].value == original
            assert google_sheet_original_title(path, local) == original
        book.close()
        data_sheets = find_data_sheet_names(path)
        assert mapping[titles[0]] not in data_sheets
        assert not any(name.startswith("Report Seeding Threads") for name in data_sheets)
        assert len(data_sheets) == 4
    assert observed_maps[0] == observed_maps[1]


@pytest.mark.parametrize("column", ["TIM", "REPOST"])
def test_preview_sums_formatted_metric_cells_with_string_dtype(tmp_path, column):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["LINK AIR", column])
    sheet.append(["https://www.threads.com/@demo/post/1", "1,234"])
    sheet.append(["https://www.threads.com/@demo/post/2", "1.234.567"])
    sheet.append(["TỔNG", "stale"])
    path = tmp_path / "formatted.xlsx"
    book.save(path)
    preview = read_sheet_preview(path)
    assert preview["data"][-1][column] == "1235801"
    # The display calculation must not rewrite the user's source workbook.
    original = openpyxl.load_workbook(path)
    assert original.active["B4"].value == "stale"
    original.close()


def test_preview_keeps_unknown_repost_total_blank(tmp_path):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["LINK AIR", "REPOST"])
    sheet.append(["https://www.threads.com/@demo/post/1", None])
    sheet.append(["TỔNG", None])
    path = tmp_path / "unknown-repost.xlsx"
    book.save(path)
    assert read_sheet_preview(path)["data"][-1]["REPOST"] == ""


@pytest.mark.parametrize("platform", ["tiktok", "threads"])
def test_platform_dates_fill_merged_groups_but_stop_at_group_boundaries(tmp_path, platform):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Data"
    sheet.append(["NGÀY AIR", "LINK AIR"])
    sheet.append(["25/09/2026", "https://www.threads.com/@demo/post/1"])
    sheet.append([None, "https://www.threads.net/@demo/post/2"])
    sheet.merge_cells("A2:A3")
    sheet.append([None, "TỔNG"])
    sheet.append([None, "https://www.threads.com/share/three/"])
    sheet.append(["26/09/2026", "https://www.threads.com/@demo/post/4"])
    sheet.append([None, None])
    sheet.append([None, "https://www.threads.com/@demo/post/5"])
    sheet.append(["27/09/2026", "https://www.threads.com/@demo/post/6"])
    sheet.append([None, "Nhóm mới"])
    sheet.append([None, "https://www.threads.com/@demo/post/7"])
    if platform == "tiktok":
        for index, row in enumerate(sheet.iter_rows(min_row=2), start=1):
            if "threads." in (row[1].value or ""):
                row[1].value = f"https://www.tiktok.com/@demo/video/{index}"
    path = tmp_path / "group-dates.xlsx"
    book.save(path)
    rows = build_workbook_rows(path, platform=platform)
    assert [clean_text(row["NGÀY AIR"]) for row in rows] == [
        "25/09/2026", "25/09/2026", "", "26/09/2026", "", "27/09/2026", ""
    ]
