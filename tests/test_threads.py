import io
import asyncio
import json
import os

from riviu.platforms import threads as threads_scraper

import openpyxl
import pytest
from playwright.async_api import async_playwright

from riviu.reports import build_export_payload, build_google_push_rows, build_partner_report
from riviu.google_sheets_sync import create_result_sheet_title

from riviu.platforms.threads import (
    parse_action_count,
    parse_threads_http,
    fetch_threads_browser,
    fetch_threads_http,
    result_status,
    run_threads_scraper,
    same_threads_post,
    write_threads_result,
)
from riviu.workbook_utils import (
    build_workbook_rows,
    is_internal_workbook_filename,
    is_threads_link,
    list_workbook_partners_with_link_counts,
    rebuild_summary_sheet,
    worksheet_find_column_index,
)


URL = "https://www.threads.com/@miri_viu/post/DdRIHGWCURV?xmt=sample"
SHARE_URL = "https://www.threads.com/share/GlQTw_txU/"


def test_threads_url_requires_real_post_host_and_identity():
    assert is_threads_link(URL)
    assert is_threads_link("https://www.threads.net/@miri_viu/post/DdRIHGWCURV")
    assert is_threads_link(SHARE_URL)
    assert not is_threads_link("https://threads.com.evil.test/@miri_viu/post/DdRIHGWCURV")
    assert not is_threads_link("https://www.tiktok.com/@miri_viu/video/1")
    assert same_threads_post(URL, "https://www.threads.com/@miri_viu/post/DdRIHGWCURV")
    assert not same_threads_post(URL, "https://www.threads.com/?error=invalid_post")
    assert same_threads_post(SHARE_URL, "https://www.threads.com/@trangxinh718/post/DdQ8PEekxig")
    assert not same_threads_post(SHARE_URL, "https://www.threads.com/?error=invalid_post")


def test_http_parser_keeps_missing_counts_unknown():
    html = '<title>Example</title><script type="application/json">{"require":[["BarcelonaLoggedOutExpansionGating",[],{"view_counts":372},7623]]}</script>'
    result = parse_threads_http(html, URL, URL)
    assert result["channel"] == "miri_viu"
    assert result["metrics"] == {"views": 372, "likes": None, "comments": None, "reposts": None, "shares": None}
    assert parse_threads_http(html, URL, "https://www.threads.com/?error=invalid_post")["error"]
    share = parse_threads_http(html, SHARE_URL, "https://www.threads.com/@trangxinh718/post/DdQ8PEekxig")
    assert share["channel"] == "trangxinh718"


def target_media(**overrides):
    media = {
        "code": "DdRIHGWCURV", "user": {"username": "miri_viu"},
        "like_count": 0, "like_and_view_counts_disabled": False,
        "text_post_app_info": {"direct_reply_count": 0, "repost_count": 0, "reshare_count": None, "quote_count": 99},
    }
    media.update(overrides)
    return media


def target_payload(*media):
    return {"require": [["RelayPrefetchedStreamCache", "next", [], [
        "adp_BarcelonaPostPageTargetQueryRelayPreloader_fixture",
        {"__bbox": {"result": {"data": {"media": item}}}},
    ]] for item in media]}


def json_script(document):
    return '<script type="application/json">' + json.dumps(document) + '</script>'


def test_target_json_preserves_explicit_zeros_and_null_shares():
    content = json_script(target_payload(target_media()))
    result = parse_threads_http(content, URL, URL)
    assert result["metrics"] == {"views": None, "likes": 0, "comments": 0, "reposts": 0, "shares": None}
    assert result_status(result) == "Success"
    assert threads_scraper.missing_metric_labels(result) == ["LƯỢT XEM", "CHIA SẺ"]


@pytest.mark.parametrize("value, expected", [(0, 0), (12, 12), (None, None), (True, None), (-1, None), ("12", None), (1.5, None)])
def test_target_json_validates_count_types(value, expected):
    media = target_media(like_count=value, text_post_app_info={
        "direct_reply_count": value, "repost_count": value, "reshare_count": value,
    })
    metrics = parse_threads_http(json_script(target_payload(media)), URL, URL)["metrics"]
    assert all(metrics[key] == expected for key in ("likes", "comments", "reposts", "shares"))


@pytest.mark.parametrize("flag", [True, None, "false", 0])
def test_target_json_does_not_expose_hidden_or_unverified_likes(flag):
    media = target_media(like_count=12, like_and_view_counts_disabled=flag)
    metrics = parse_threads_http(json_script(target_payload(media)), URL, URL)["metrics"]
    assert metrics["likes"] is None
    assert metrics["comments"] == 0


@pytest.mark.parametrize("media", [target_media(code="Other"), target_media(user={"username": "other"}), target_media(user=None)])
def test_target_json_rejects_wrong_identity(media):
    assert all(value is None for value in parse_threads_http(json_script(target_payload(media)), URL, URL)["metrics"].values())


def test_target_json_rejects_decoys_outside_typed_target_query():
    media = target_media(like_count=999)
    content = json_script({"caption": json.dumps(target_payload(media)), "recommended_post": media})
    content += json_script({"require": [["RelayPrefetchedStreamCache", "next", [], [
        "adp_BarcelonaLoggedOutRelatedPostsQueryRelayPreloader_fixture",
        {"__bbox": {"result": {"data": {"media": media}}}},
    ]]]})
    assert all(value is None for value in parse_threads_http(content, URL, URL)["metrics"].values())


@pytest.mark.parametrize("second", [target_media(like_count=7), target_media(like_count=None), target_media(like_and_view_counts_disabled=True)])
def test_target_json_conflicting_values_remain_unknown(second):
    metrics = parse_threads_http(json_script(target_payload(target_media(), second)), URL, URL)["metrics"]
    assert metrics["likes"] is None
    assert metrics["reposts"] == 0


def test_partial_hybrid_keeps_browser_diagnostic_without_losing_views():
    first = {"channel": "miri_viu", "metrics": {"views": 135}, "error": ""}
    second = {"channel": "", "metrics": threads_scraper.empty_metrics(), "error": "Browser: timeout"}
    merged = threads_scraper.merge_results(first, second)
    assert merged["metrics"]["views"] == 135
    assert merged["diagnostic"] == "Browser: timeout"
    assert result_status(merged) == "Success"


def test_target_json_missing_fields_and_malformed_scripts_are_unknown():
    media = target_media(text_post_app_info=None)
    media.pop("like_count")
    content = '<script type="application/json">{broken</script>' + json_script(target_payload(media))
    assert all(value is None for value in parse_threads_http(content, URL, URL)["metrics"].values())


def test_hybrid_saves_explicit_zero_counts_and_partial_status(tmp_path, monkeypatch):
    path = tmp_path / "zero_counts.xlsx"
    book = openpyxl.Workbook()
    book.active.append(["Link", "LƯỢT XEM", "TIM", "BÌNH LUẬN", "REPOST", "CHIA SẺ"])
    book.active.append([URL])
    book.save(path)
    book.close()
    monkeypatch.setattr(threads_scraper, "fetch_threads_http", lambda _url: {
        "channel": "miri_viu", "metrics": {**threads_scraper.empty_metrics(), "views": 135}, "error": "",
    })

    async def rendered(_browser, _url):
        return parse_threads_http(json_script(target_payload(target_media())), URL, URL)

    monkeypatch.setattr(threads_scraper, "fetch_threads_browser", rendered)
    asyncio.run(run_threads_scraper(path, mode="hybrid"))
    saved = openpyxl.load_workbook(path)
    try:
        assert [saved.active.cell(2, col).value for col in range(2, 7)] == [135, 0, 0, 0, None]
        col = worksheet_find_column_index(saved.active, [threads_scraper.THREADS_SCAN_STATUS_HEADER])
        assert saved.active.cell(2, col).value == "Success"
    finally:
        saved.close()


@pytest.mark.parametrize("metrics", [
    {"views": 0}, {"likes": 0}, {"views": 314, "likes": 2, "comments": 6, "reposts": 0},
    {"likes": 0, "comments": 0, "reposts": 0, "shares": 0},
])
def test_confirmed_metrics_including_zero_are_success(metrics):
    result = {"metrics": {**threads_scraper.empty_metrics(), **metrics}, "error": ""}
    assert result_status(result) == "Success"


def test_no_metric_evidence_is_not_success():
    assert result_status({"metrics": threads_scraper.empty_metrics(), "error": ""}).startswith("Ẩn số liệu:")
    assert result_status({"metrics": {"views": 0}, "error": "Không đọc được bài"}).startswith("Error:")


def test_action_count_does_not_convert_blank_to_zero():
    assert parse_action_count("4") == 4
    assert parse_action_count("1.2K") == 1200
    assert parse_action_count("") is None


@pytest.mark.parametrize("text,count", [("32.3K", 32300), ("4.1M", 4100000), ("1.15K", 1150), ("2.7K", 2700)])
def test_abbreviated_counts_are_exact(text, count):
    # Binary floats truncated 32.3K to 32299 and 4.1M to 4099999.
    assert parse_action_count(text) == count


NFD_COMMENTS = "BI\u0300NH LUA\u0323\u0302N"  # BÌNH LUẬN decomposed (NFD), as some exports write it


def test_threads_result_reuses_existing_header_spellings():
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    # Without diacritics, and NFD-decomposed: the preview reads both as the same column.
    sheet.append(["Link", "Tên kênh", "Luot xem", "TIM", NFD_COMMENTS, "REPOST", "CHIA SẺ"])
    sheet.append([URL])
    result = {"channel": "miri_viu", "error": "",
              "metrics": {"views": 15, "likes": 2, "comments": 3, "reposts": 0, "shares": None}}
    write_threads_result(sheet, 2, result)
    headers = [cell.value for cell in sheet[1]]
    assert headers[:7] == ["Link", "Tên kênh", "Luot xem", "TIM", NFD_COMMENTS, "REPOST", "CHIA SẺ"]
    assert "LƯỢT XEM" not in headers and "BÌNH LUẬN" not in headers and "Tên Kênh" not in headers
    assert [sheet.cell(2, column).value for column in range(2, 7)] == ["miri_viu", 15, 2, 3, 0]


def test_threads_save_temp_file_is_hidden_from_workbook_lists(tmp_path):
    target = tmp_path / "threads.xlsx"
    written = []

    class Workbook:
        def save(self, path):
            written.append(os.path.basename(path))
            # An interrupted save leaves this name behind; it must never look like a workbook.
            assert is_internal_workbook_filename(path) and not str(path).endswith(".xlsx")
            with open(path, "wb") as stream:
                stream.write(b"saved")

    threads_scraper.save_workbook_atomic(Workbook(), target)
    assert target.read_bytes() == b"saved" and len(written) == 1
    assert os.listdir(tmp_path) == ["threads.xlsx"]
    assert result_status({"metrics": {"views": 372, "likes": 4, "comments": 1, "reposts": None, "shares": 2}, "error": ""}) == "Success"


def test_workbook_filters_platform_and_preserves_missing_metrics(tmp_path):
    path = tmp_path / "mixed.xlsx"
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Data"
    sheet.append(["Ngày", "Link", "Tên Kênh", "LƯỢT XEM", "TIM", "BÌNH LUẬN", "LƯỢT LƯU", "CHIA SẺ", "Đối tác"])
    sheet.append(["14/9/2026", URL, "", 100, 2, "", "", "", "Cafe A"])
    sheet.append(["14/9/2026", "https://www.tiktok.com/@demo/video/123", "", 300, 5, 1, 2, 3, "Cafe A"])
    book.save(path)

    threads_rows = build_workbook_rows(path, platform="threads")
    assert len(threads_rows) == 1
    assert threads_rows[0]["REPOST"] == ""
    assert len(build_workbook_rows(path)) == 1

    write_threads_result(sheet, 2, {"channel": "miri_viu", "metrics": {"views": 372, "likes": None, "comments": 1, "reposts": None, "shares": 2}, "error": ""})
    assert sheet["D2"].value == 372
    assert sheet["E2"].value == 2
    assert sheet["F2"].value == 1
    assert sheet["G2"].value == ""
    assert sheet["H2"].value == 2
    assert sheet["C2"].value == "miri_viu"
    assert sheet["G3"].value == 2
    assert any(cell.value == "REPOST" for cell in sheet[1])


def test_threads_exports_use_repost_and_keep_unknown_blank():
    rows = [{
        "NGÀY AIR": "14/9/2026", "TÊN KÊNH": "miri_viu", "LINK AIR": URL,
        "LƯỢT XEM": 372, "TIM": 4, "BÌNH LUẬN": 1, "REPOST": "",
        "CHIA SẺ": 2, "partners": ["Cafe A"],
    }]
    pushed = build_google_push_rows(rows, platform="threads")
    assert pushed[0][7] == "REPOST"
    assert pushed[1][7] == ""
    assert "Trạng thái" not in pushed[0]
    assert pushed[2][7] == ""

    report = openpyxl.load_workbook(io.BytesIO(build_partner_report("Cafe A", rows, platform="threads")))
    sheet = report.active
    assert sheet["G3"].value == "REPOST"
    assert sheet["G4"].value in (None, "")
    assert sheet["I3"].value is None
    assert sheet["G5"].value is None


def test_threads_google_result_tab_has_distinct_name():
    assert create_result_sheet_title("Tháng 9", platform="threads").startswith("Report Seeding Threads ")


def test_threads_partner_export_excludes_tiktok(tmp_path):
    path = tmp_path / "mixed.xlsx"
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Data"
    sheet.append(["Ngày", "Link", "Tên Kênh", "LƯỢT XEM", "TIM", "BÌNH LUẬN", "LƯỢT LƯU", "CHIA SẺ", "Đối tác"])
    sheet.append(["14/9/2026", URL, "miri_viu", 180, 4, 1, "", 2, "Cafe A"])
    sheet.append(["14/9/2026", "https://www.tiktok.com/@demo/video/123", "TikTok", 500, 5, 2, 3, 4, "Cafe A"])
    sheet.append(["14/9/2026", "https://www.tiktok.com/@other/video/234", "TikTok 2", 300, 2, 1, 0, 0, "TikTok Only"])
    book.save(path)

    partners = list_workbook_partners_with_link_counts(path, sheet_name="Data", platform="threads", min_views=100)
    assert [partner["name"] for partner in partners] == ["Cafe A"]
    assert partners[0]["linkCount"] == 1
    payload = build_export_payload(path, ["Cafe A"], True, 100, "Data", "threads")
    report = openpyxl.load_workbook(io.BytesIO(payload["content"]))
    assert "Threads" in payload["filename"]
    assert report.active["C4"].value == URL
    assert report.active["C5"].value is None
    assert report.active["G3"].value == "REPOST"
    assert report.active["I3"].value is None


def test_threads_runner_saves_before_done_and_ignores_tiktok(tmp_path, monkeypatch):
    path = tmp_path / "mixed.xlsx"
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Data"
    sheet.append(["Link", "Tên Kênh", "LƯỢT XEM", "TIM", "Đối tác"])
    sheet.append([URL, "", "", "", "Cafe A"])
    sheet.append(["https://www.tiktok.com/@demo/video/123", "TikTok", 12, 5, "Cafe A"])
    rebuild_summary_sheet(book, data_sheet_name="Data")
    book.save(path)

    monkeypatch.setattr("riviu.platforms.threads.fetch_threads_http", lambda _url: {
        "channel": "miri_viu", "metrics": {"views": 372, "likes": 4, "comments": 1, "reposts": None, "shares": 2}, "error": "",
    })

    class Manager:
        async def broadcast_log(self, _message):
            pass

        async def broadcast_duplicates(self, _data):
            pass

        async def broadcast_data(self, _data):
            pass

        async def broadcast_status(self, data):
            if data.get("done"):
                saved = openpyxl.load_workbook(path)
                assert saved.active["C2"].value == 372
                assert saved.active["C3"].value == 12
                # The scan rebuilds its own Threads summary and leaves the TikTok one as it was.
                assert saved.sheetnames == ["Data", "Tổng kết data", "Tổng kết Threads data"]
                tiktok, threads = saved["Tổng kết data"], saved["Tổng kết Threads data"]
                assert [cell.value for cell in tiktok[2]][1:4] == ["Cafe A", 1, 12]
                assert [cell.value for cell in threads[1]][6] == "TỔNG REPOST"
                assert [cell.value for cell in threads[2]][1:8] == ["Cafe A", 1, 372, 4, 1, None, 2]
                assert threads.cell(row=2, column=9).value
                saved.close()

    asyncio.run(run_threads_scraper(path, Manager(), sheet_name="Data", mode="request"))


@pytest.mark.skipif(not os.environ.get("THREADS_LIVE_URL"), reason="live Threads check is opt-in")
def test_live_threads_http_and_browser_agree_on_post(tmp_path):
    url = os.environ["THREADS_LIVE_URL"]
    request = fetch_threads_http(url)

    async def browser_check():
        async with async_playwright() as playwright:
            browser = await playwright.chromium.launch(headless=True)
            original_new_page = browser.new_page
            source = {}

            async def capture_page():
                page = await original_new_page()
                close = page.close

                async def capture_close():
                    content = await page.content()
                    source.update(parse_threads_http(content, url, page.url))
                    _, source["has_target"] = threads_scraper._embedded_post_metrics(content, threads_scraper.post_identity(page.url))
                    await close()

                page.close = capture_close
                return page

            browser.new_page = capture_page
            try:
                return await fetch_threads_browser(browser, url), source
            finally:
                await browser.close()

    rendered, source = asyncio.run(browser_check())
    assert not request["error"], request
    assert not rendered["error"], rendered
    assert rendered["channel"] == request["channel"]
    assert rendered["metrics"]["views"] is not None
    assert rendered["metrics"]["likes"] is not None
    assert rendered["metrics"]["comments"] is not None
    # Verify against the same rendered target payload, including explicit null shares.
    if source["has_target"]:
        for key in threads_scraper.ACTION_TITLES:
            assert rendered["metrics"][key] == source["metrics"][key]
    else:
        assert rendered["metrics"]["shares"] is not None

    path = tmp_path / "live_threads.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["Link", "Tên Kênh", "LƯỢT XEM", "TIM", "BÌNH LUẬN", "LƯỢT LƯU", "CHIA SẺ", "Đối tác"])
    sheet.append([url, "", "", "", "", "", "", "Cafe A"])
    workbook.save(path)
    asyncio.run(run_threads_scraper(path, sheet_name="Data", mode="hybrid"))
    saved = openpyxl.load_workbook(path)
    assert saved.active["B2"].value == rendered["channel"]
    assert isinstance(saved.active["C2"].value, int)
    assert saved.active["F2"].value is None
    for key, header in threads_scraper.METRICS.items():
        column = worksheet_find_column_index(saved.active, [header])
        actual = saved.active.cell(2, column).value if column else None
        assert actual == rendered["metrics"][key]
    saved.close()


@pytest.mark.skipif(not os.environ.get("THREADS_SHARE_URL"), reason="live Threads share check is opt-in")
def test_live_threads_share_redirect_in_http_and_browser():
    url = os.environ["THREADS_SHARE_URL"]
    request = fetch_threads_http(url)

    async def browser_check():
        async with async_playwright() as playwright:
            browser = await playwright.chromium.launch(headless=True)
            try:
                return await fetch_threads_browser(browser, url)
            finally:
                await browser.close()

    rendered = asyncio.run(browser_check())
    assert not request["error"], request
    assert not rendered["error"], rendered
    assert request["channel"] == rendered["channel"]
    assert request["metrics"]["views"] is not None


@pytest.mark.skipif(not os.environ.get("RIVIU_TEST_URL"), reason="local UI check is opt-in")
def test_threads_ui_uses_repost_column_not_tiktok_saves():
    from playwright.sync_api import sync_playwright

    from urllib.parse import parse_qs, urlparse

    # Self-contained and read-only: the server's saved sources, workbook list and selection are
    # never read or written. Every source API answers from one synthetic Threads workbook.
    file_id, label = "data/threads-ui-fixture.xlsx", "Threads UI fixture"
    empty_source = {"fileId": "", "displaySheet": "", "scanSheet": "", "pushSheet": "", "url": ""}
    sources = {"tiktok": dict(empty_source),
               "threads": {**empty_source, "fileId": file_id, "displaySheet": "Data", "scanSheet": "Data", "pushSheet": "Data"}}
    preview = {
        "sheets": ["Data"], "currentSheet": "Data",
        "columns": ["Link", "Tên Kênh", "LƯỢT XEM", "TIM", "BÌNH LUẬN", "LƯỢT LƯU", "REPOST", "CHIA SẺ"],
        "data": [{"Link": URL, "Tên Kênh": "miri_viu", "LƯỢT XEM": 372, "TIM": 4,
                  "BÌNH LUẬN": 1, "LƯỢT LƯU": 99, "REPOST": 3, "CHIA SẺ": 2}],
    }

    def query(route):
        params = parse_qs(urlparse(route.request.url).query)
        return params.get("platform", ["tiktok"])[0], params.get("file_id", [""])[0]

    def source_preferences(route):
        if route.request.method == "GET":
            route.fulfill(json={"sources": sources})
        else:
            route.fulfill(json={"success": True})

    def list_files(route):
        platform, requested = query(route)
        selected = file_id if platform == "threads" and requested == file_id else ""
        route.fulfill(json={
            "files": [{"id": file_id, "label": label, "source": "google"}],
            "current": selected, "file_id": selected, "currentLabel": label if selected else "",
            "currentSheet": "Data" if selected else "", "sheets": ["Data"] if selected else [],
            "scanSheet": "Data" if selected else "", "googleSheetUrl": "", "platform": platform,
            "googlePushReady": False, "googleOAuthConfigured": False, "googleOAuthAuthorized": False,
        })

    def select_file(route):
        body = route.request.post_data_json or {}
        selected = body.get("file_id") or body.get("filename") or ""
        if selected != file_id:
            route.fulfill(status=400, json={"error": "Không tìm thấy file"})
            return
        route.fulfill(json={"success": True, "selected": file_id, "file_id": file_id, "sheet": "Data",
                            "scanSheet": "Data", "platform": body.get("platform", "threads")})

    def preview_excel(route):
        platform, requested = query(route)
        if platform == "threads" and requested == file_id:
            route.fulfill(json={**preview, "file": file_id, "file_id": file_id, "fileLabel": label,
                                "platform": platform, "summarySource": "Data"})
        else:
            route.fulfill(json={"sheets": [], "currentSheet": "", "columns": [], "data": [], "platform": platform})

    server = urlparse(os.environ["RIVIU_TEST_URL"])
    unmocked_writes = []

    def read_only_guard(route):
        # Registered first, so it only sees requests no mock above claimed.
        request = urlparse(route.request.url)
        if request.netloc == server.netloc and route.request.method not in {"GET", "HEAD"}:
            unmocked_writes.append(f"{route.request.method} {request.path}")
            route.abort()
        else:
            route.continue_()

    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(headless=True)
        try:
            for width, height in [(1440, 900), (390, 844)]:
                page = browser.new_page(viewport={"width": width, "height": height})
                errors = []
                page.on("pageerror", lambda error: errors.append(str(error)))
                page.route("**/*", read_only_guard)
                page.route("**/source-preferences*", source_preferences)
                page.route("**/list-files*", list_files)
                page.route("**/select-file*", select_file)
                page.route("**/preview-excel*", preview_excel)
                try:
                    page.goto(os.environ["RIVIU_TEST_URL"], wait_until="domcontentloaded")
                    page.wait_for_function("websocketSessionReady")
                    page.locator("#platformThreads").click()
                    page.get_by_role("columnheader", name="REPOST").wait_for()
                    headers = page.locator("#previewHeader th").all_inner_texts()
                    cells = page.locator("#previewBody tr td").all_inner_texts()
                    assert "LƯỢT LƯU" not in headers
                    assert cells[headers.index("REPOST")] == "3"
                    assert page.locator("#liveSavedHeader").inner_text() == "REPOST"
                    page.evaluate("""() => {
                        startBtn.disabled = true;
                        cancelBtn.disabled = false;
                        updateProgress({phase: 'running', total: 1, processed: 1, success: 0, hidden: 1, error: 0, done: false});
                    }""")
                    assert page.locator("#startBtn").is_disabled()
                    assert page.locator("#progressStatus").inner_text() == ""
                    page.evaluate("updateProgress({phase: 'completed', total: 1, processed: 1, success: 0, hidden: 1, error: 0, done: true})")
                    assert not page.locator("#startBtn").is_disabled()
                    assert "không trả số liệu" in page.locator("#progressStatus").inner_text()
                    assert not errors, errors
                finally:
                    # Leave the app while the mocks are still served: a source save queued at the
                    # end would otherwise slip past the routes during page.close() and hit the server.
                    page.goto("about:blank")
                    page.close()
            assert not unmocked_writes, unmocked_writes
        finally:
            browser.close()
