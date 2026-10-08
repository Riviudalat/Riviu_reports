"""Source-aware APIs use only temporary workbooks and mocked remote I/O."""
import asyncio
import io
import json
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from starlette.datastructures import UploadFile
from starlette.websockets import WebSocketDisconnect

import app as backend
from workbook_utils import (
    PLATFORMS,
    google_sheet_source_for_file,
    register_google_sheet_source,
    summary_sheet_title_for_data_sheet,
)


def make_source(path, marker):
    book = Workbook()
    sheet = book.active
    sheet.title = f"{marker} Data"
    sheet.append(["Link", "Tên Kênh", "Đối tác", "LƯỢT XEM", "TIM", "REPOST"])
    sheet.append([f"https://www.tiktok.com/@{marker}/video/123", marker, f"{marker} TikTok", 1000, 10, ""])
    sheet.append([f"https://www.threads.com/@{marker}/post/Fixture", marker, f"{marker} Threads", 2000, 20, 2])
    second = book.create_sheet(f"{marker} Other")
    second.append(["Link", "Tên Kênh", "Đối tác", "LƯỢT XEM"])
    second.append([f"https://www.tiktok.com/@{marker}/video/456", marker, f"{marker} Other Partner", 3000])
    # The explicit-source tests query platform=threads, so the fixture holds the Threads summary.
    summary = book.create_sheet(summary_sheet_title_for_data_sheet(sheet.title, "threads"))
    summary.append(PLATFORMS["threads"].summary_columns)
    summary.append([1, f"{marker} Summary", 2, 3000, 30, 0, 0, 0, "fixture"])
    book.save(path)
    book.close()


@pytest.fixture
def source_backend(tmp_path, monkeypatch):
    monkeypatch.setattr(backend, "EXCEL_DIR", str(tmp_path))
    monkeypatch.setattr(backend, "CURRENT_SELECTED_FILE", "A.xlsx")
    monkeypatch.setattr(backend, "CURRENT_SELECTED_SHEET", "A Other")
    monkeypatch.setattr(backend, "CURRENT_SCAN_SHEET", "A Other")
    monkeypatch.setattr(backend, "SCRAPE_TASK", None)
    monkeypatch.setattr(backend, "SCAN_STARTING", False)
    monkeypatch.setattr(backend, "SOURCE_BUSY", False)
    monkeypatch.setattr(backend, "DESKTOP_UPDATE_PENDING", False)
    monkeypatch.setattr(backend, "manager", backend.ConnectionManager())
    monkeypatch.setattr(backend, "oauth_status", lambda _: {"valid": True, "configured": True})
    for marker in ("A", "B"):
        make_source(tmp_path / f"{marker}.xlsx", marker)
        register_google_sheet_source(str(tmp_path), f"{marker}.xlsx", f"https://docs.google.com/spreadsheets/d/ID_{marker}/edit")
    return tmp_path


@pytest.fixture
def client(source_backend):
    with TestClient(backend.app, base_url="http://127.0.0.1:1231") as client:
        assert client.get("/").status_code == 200
        yield client


def selection():
    return backend.CURRENT_SELECTED_FILE, backend.CURRENT_SELECTED_SHEET, backend.CURRENT_SCAN_SHEET


def test_explicit_list_and_empty_source_do_not_rebind(client):
    before = selection()
    state = client.get("/list-files", params={"file_id": "B.xlsx", "platform": "threads"}).json()
    assert state["current"] == state["file_id"] == "B.xlsx"
    assert state["currentLabel"] == "B.xlsx"
    assert state["currentSheet"] == state["scanSheet"] == "B Data"
    assert state["sheets"] == ["B Data", "B Other"]
    assert state["googleSheetUrl"].endswith("/ID_B/edit")
    assert state["platform"] == "threads" and state["googlePushReady"]
    empty = client.get("/list-files", params={"file_id": "", "platform": "threads"}).json()
    assert empty["current"] == empty["currentLabel"] == empty["googleSheetUrl"] == ""
    assert empty["sheets"] == [] and len(empty["files"]) == 2
    assert not empty["googlePushReady"]
    assert selection() == before


def test_explicit_preview_summary_partners_download_and_export_use_b(client, source_backend):
    before = selection()
    query = {"file_id": "B.xlsx", "platform": "threads"}
    preview = client.get("/preview-excel", params=query).json()
    assert preview["file"] == preview["file_id"] == "B.xlsx"
    assert preview["currentSheet"] == "B Data" and preview["platform"] == "threads"
    assert "B Threads" in json.dumps(preview)
    summary = client.get("/summary-dashboard", params=query).json()
    assert summary["file"] == "B.xlsx" and summary["dataSheet"] == "B Data"
    assert "B Summary" in json.dumps(summary)
    partners = client.get("/report-partners", params=query).json()
    assert partners["file"] == "B.xlsx" and partners["platform"] == "threads"
    assert [item["name"] for item in partners["partners"]] == ["B Threads"]
    download = client.get("/download-excel", params=query)
    assert download.status_code == 200
    assert download.content == (source_backend / "B.xlsx").read_bytes()
    exported = client.post("/export-report", json={**query, "partners": ["B Threads"], "sheetName": "B Data"})
    assert exported.status_code == 200
    book = load_workbook(io.BytesIO(exported.content))
    try:
        values = str(list(book.active.values))
        assert "B Threads" in values and "threads.com/@B" in values
        assert "tiktok.com/@B" not in values and "A TikTok" not in values
    finally:
        book.close()
    assert selection() == before


def test_sheet_lists_report_the_workbooks_hidden_sheets(client, source_backend):
    book = Workbook()
    old_month = book.active
    old_month.title = "Tháng 6"
    old_month.sheet_state = "hidden"
    for sheet in (old_month, book.create_sheet("Tháng 8")):
        sheet.append(["Link", "Tên Kênh", "Đối tác", "LƯỢT XEM"])
        sheet.append([f"https://www.tiktok.com/@{sheet.title[-1]}/video/1", "kenh", "Cafe", 100])
    book.save(source_backend / "C.xlsx")
    book.close()
    query = {"file_id": "C.xlsx", "platform": "tiktok"}
    for route in ("/list-files", "/preview-excel", "/report-partners"):
        data = client.get(route, params=query).json()
        assert data["hiddenSheets"] == ["Tháng 6"], route
        assert set(data["sheets"]) == {"Tháng 6", "Tháng 8"}, route
    assert client.get("/list-files", params={"file_id": "B.xlsx", "platform": "tiktok"}).json()["hiddenSheets"] == []


def test_threads_report_apis_honour_min_views_like_tiktok(client, source_backend):
    book = Workbook()
    sheet = book.active
    sheet.title = "Data"
    sheet.append(["Link", "Tên Kênh", "Đối tác", "LƯỢT XEM"])
    sheet.append(["https://www.threads.com/@hi/post/AAA", "hi", "Cafe", 500])
    sheet.append(["https://www.threads.com/@lo/post/BBB", "lo", "Cafe", 50])
    sheet.append(["https://www.threads.com/@zero/post/CCC", "zero", "Cafe", 0])
    book.save(source_backend / "C.xlsx")
    book.close()
    query = {"file_id": "C.xlsx", "platform": "threads", "sheet_name": "Data"}

    def link_count(**params):
        [partner] = client.get("/report-partners", params={**query, **params}).json()["partners"]
        return partner["linkCount"], partner["rawLinkCount"]

    assert link_count() == (1, 3)
    assert link_count(min_views=40) == (2, 3)
    assert link_count(apply_min_views="false") == (3, 3)

    def exported_links(**options):
        response = client.post("/export-report", json={
            "file_id": "C.xlsx", "platform": "threads", "sheetName": "Data", "partners": ["Cafe"], **options,
        })
        assert response.status_code == 200
        report = load_workbook(io.BytesIO(response.content))
        try:
            return [row[2] for row in report.active.iter_rows(min_row=4, values_only=True) if row[2] and "threads.com" in row[2]]
        finally:
            report.close()

    assert exported_links() == ["https://www.threads.com/@hi/post/AAA"]
    assert len(exported_links(minViews=40)) == 2
    assert len(exported_links(applyMinViews=False)) == 3


def test_explicit_push_uses_b_provenance_and_never_registers_override(client, source_backend, monkeypatch):
    calls = []
    monkeypatch.setattr(backend, "push_rows_to_new_sheet", lambda base, sid, rows, **kw: calls.append((sid, rows, kw)) or "Result")
    before = selection()
    query = {"file_id": "B.xlsx", "platform": "threads"}
    pushed = client.post("/push-google-sheet", json=query).json()
    assert pushed["file_id"] == "B.xlsx" and pushed["sourceSheet"] == "B Data"
    assert calls[0][0] == "ID_B" and calls[0][2]["platform"] == "threads"
    assert "B Threads" in str(calls[0][1]) and "A TikTok" not in str(calls[0][1])
    assert client.post("/push-google-sheet", json={**query, "url": "https://docs.google.com/spreadsheets/d/OVERRIDE/edit"}).status_code == 200
    assert calls[1][0] == "OVERRIDE"
    assert google_sheet_source_for_file(str(source_backend), "B.xlsx")["spreadsheetId"] == "ID_B"
    assert selection() == before


@pytest.mark.parametrize("route", ["/preview-excel", "/summary-dashboard", "/report-partners", "/download-excel"])
@pytest.mark.parametrize("file_id", ["", "missing.xlsx", "../A.xlsx", "/A.xlsx"])
def test_explicit_invalid_get_source_never_falls_back(client, route, file_id):
    before = selection()
    response = client.get(route, params={"file_id": file_id})
    assert response.status_code in (400, 404)
    assert "error" in response.json() and selection() == before


@pytest.mark.parametrize("route", ["/push-google-sheet", "/export-report"])
@pytest.mark.parametrize("file_id", ["", "missing.xlsx", "../A.xlsx", 42])
def test_explicit_invalid_body_source_never_falls_back(client, route, file_id):
    before = selection()
    response = client.post(route, json={"file_id": file_id, "partners": ["A TikTok"]})
    assert response.status_code in (400, 404)
    assert selection() == before


def test_explicit_preview_is_readonly_even_for_current_file(client):
    before = selection()
    explicit = client.get("/preview-excel", params={"file_id": "A.xlsx", "sheet_name": "A Data"}).json()
    assert explicit["currentSheet"] == "A Data" and selection() == before
    legacy = client.get("/preview-excel", params={"sheet_name": "A Data"}).json()
    assert legacy["currentSheet"] == "A Data"
    assert backend.CURRENT_SELECTED_SHEET == "A Data"


def test_legacy_no_argument_direct_calls_and_default_activation(source_backend, monkeypatch):
    assert asyncio.run(backend.list_files())["current"] == "A.xlsx"
    assert asyncio.run(backend.preview_excel())["currentSheet"] == "A Other"
    assert asyncio.run(backend.summary_dashboard())["file"] == "A.xlsx"
    assert asyncio.run(backend.report_partners())["file"] == "A.xlsx"
    download = asyncio.run(backend.download_excel())
    assert "A.xlsx" in download.headers["content-disposition"]
    # A snapshot of the bytes, not a stream of the live file: an open download must not block the scan's atomic replace.
    assert download.body == (Path(backend.EXCEL_DIR) / "A.xlsx").read_bytes()
    calls = []
    monkeypatch.setattr(backend, "push_rows_to_new_sheet", lambda base, sid, rows, **kw: calls.append(sid) or "Result")
    assert asyncio.run(backend.push_google_sheet())["success"] and calls == ["ID_A"]
    uploaded = asyncio.run(backend.upload_excel(UploadFile(filename="Legacy.xlsx", file=io.BytesIO((source_backend / "B.xlsx").read_bytes()))))
    assert uploaded["success"] and uploaded["platform"] == "tiktok"
    assert selection() == ("Legacy.xlsx", "B Data", "B Data")


def test_sync_activate_false_publishes_registered_source_without_activation(client, source_backend, monkeypatch):
    monkeypatch.setattr(backend, "fetch_google_spreadsheet_title", lambda _: "Synced")
    monkeypatch.setattr(backend, "sync_download_google_sheet", lambda base, url, sid, destination: make_source(destination, "New"))
    logs = []
    async def log(message, **kwargs): logs.append(message)
    monkeypatch.setattr(backend.manager, "broadcast_log", log)
    before = selection()
    data = {"url": "https://docs.google.com/spreadsheets/d/SYNC/edit", "platform": "threads", "activate": False}
    response = client.post("/sync-google-sheet", json=data).json()
    assert response["success"] and response["platform"] == "threads"
    assert response["currentSheet"] == response["scanSheet"] == "New Data"
    assert response["sheets"] == ["New Data", "New Other", "Tổng kết Threads new data"]
    assert (source_backend / response["file_id"]).is_file()
    assert google_sheet_source_for_file(str(source_backend), response["file_id"])["spreadsheetId"] == "SYNC"
    assert selection() == before and "sheet=New Data" in logs[0]
    activated = client.post("/sync-google-sheet", json={"url": data["url"]}).json()
    assert activated["success"] and selection() == (activated["file"], "New Data", "New Data")


def test_upload_multipart_activate_false_preserves_current(client, source_backend):
    before = selection()
    response = client.post("/upload-excel", data={"platform": "threads", "activate": "false"}, files={"file": ("Uploaded.xlsx", (source_backend / "B.xlsx").read_bytes())}).json()
    assert response["success"] and response["platform"] == "threads" and not response["activate"]
    assert response["file_id"] == "Uploaded.xlsx" and response["sheets"] == ["B Data", "B Other"]
    assert response["sheet"] == response["scanSheet"] == "B Data"
    assert (source_backend / "Uploaded.xlsx").is_file() and selection() == before


def test_select_restores_sheets_and_validates_before_committing(client):
    before = selection()
    rejected = client.post("/select-file", json={"file_id": "B.xlsx", "sheet_name": "B Other", "scan_sheet": "A Other", "platform": "threads"})
    assert rejected.status_code == 400 and selection() == before
    response = client.post("/select-file", json={"file_id": "B.xlsx", "sheet_name": "B Other", "scan_sheet": "B Data", "platform": "threads"}).json()
    assert response["success"] and response["file_id"] == "B.xlsx" and response["platform"] == "threads"
    assert selection() == ("B.xlsx", "B Other", "B Data")
    legacy = client.post("/select-file", json={"filename": "A.xlsx"}).json()
    assert legacy["success"] and selection() == ("A.xlsx", "A Data", "A Data")


@pytest.mark.parametrize("route", ["/list-files", "/preview-excel", "/summary-dashboard", "/report-partners", "/download-excel"])
def test_wrong_platform_is_rejected_before_file_resolution(client, monkeypatch, route):
    def unexpected(): pytest.fail("Invalid platform must not resolve legacy current file")
    monkeypatch.setattr(backend, "ensure_selected_file", unexpected)
    assert client.get(route, params={"platform": "wrong"}).status_code == 400


@pytest.mark.parametrize("route", ["/select-file", "/sync-google-sheet", "/push-google-sheet", "/export-report"])
def test_wrong_body_platform_is_rejected_before_file_resolution(client, monkeypatch, route):
    def unexpected(): pytest.fail("Invalid platform must not inspect files")
    monkeypatch.setattr(backend, "file_entries", unexpected)
    assert client.post(route, json={"platform": "wrong", "file_id": "B.xlsx"}).status_code == 400


def test_wrong_upload_platform_rejected_before_publication(client, source_backend):
    assert client.post("/upload-excel", data={"platform": "wrong"}, files={"file": ("Never.xlsx", b"invalid")}).status_code == 400
    assert not (source_backend / "Never.xlsx").exists()


def test_local_token_and_origin_middleware_still_protect_explicit_sources(source_backend):
    with TestClient(backend.app, base_url="http://127.0.0.1:1231") as client:
        assert client.get("/preview-excel?file_id=B.xlsx").status_code == 403
        client.get("/")
        assert client.get("/preview-excel?file_id=B.xlsx", headers={"Origin": "https://foreign.example"}).status_code == 403
        assert client.get("/list-files?file_id=B.xlsx", headers={"Host": "foreign.example:1231"}).status_code == 403


async def websocket_start(monkeypatch, payload):
    class Socket:
        def __init__(self): self.messages, self.received = [], False
        async def accept(self): pass
        async def send_json(self, message): self.messages.append(message)
        async def receive_text(self):
            if not self.received:
                self.received = True
                return json.dumps({"action": "start", "sheet_name": "A Data", **payload})
            await asyncio.sleep(0)
            raise WebSocketDisconnect()
    monkeypatch.setattr(backend, "trusted_local_request", lambda *args, **kwargs: True)
    monkeypatch.setattr(backend, "valid_local_session", lambda *args: True)
    socket = Socket()
    await backend.websocket_endpoint(socket)
    if backend.SCRAPE_TASK:
        await backend.SCRAPE_TASK
    return socket.messages


@pytest.mark.parametrize("payload", [{"platform": "wrong"}, {"file_id": ""}])
def test_websocket_rejects_platform_and_empty_source_before_legacy_resolution(source_backend, monkeypatch, payload):
    def unexpected(): pytest.fail("Rejected start must not resolve current workbook")
    monkeypatch.setattr(backend, "ensure_selected_file", unexpected)
    messages = asyncio.run(websocket_start(monkeypatch, payload))
    assert any(message.get("data", {}).get("phase") == "failed" for message in messages)
    assert backend.SCRAPE_TASK is None


@pytest.mark.parametrize("use_proxy", [False, True])
def test_session_start_follows_proxy_toggle_without_separate_consent(source_backend, monkeypatch, use_proxy):
    cookies = ({"name": "sessionid", "value": "synthetic-only"},)
    generations, calls = [], []
    class Session:
        def snapshot(self, generation):
            generations.append(generation)
            return cookies
    async def runner(path, options, **kwargs): calls.append(options)
    monkeypatch.setattr(backend, "threads_session_manager", lambda: Session())
    monkeypatch.setattr(backend, "run_scraper_safely", runner)
    asyncio.run(websocket_start(monkeypatch, {
        "platform": "threads", "file_id": "A.xlsx", "use_threads_session": True,
        "threads_session_generation": "fixture-generation", "use_proxy": use_proxy,
        "proxy_text": "first.example:8080\nsecond.example:8080",
    }))
    assert generations == ["fixture-generation"] and len(calls) == 1
    assert calls[0]["threads_cookies"] == cookies
    assert calls[0]["threads_generation"] == "fixture-generation"
    proxies = calls[0]["threads_proxies"]
    assert ([proxy["host"] for proxy in proxies] if use_proxy else proxies) == (["first.example", "second.example"] if use_proxy else [])


def test_session_start_rejects_socks5_auth_proxy_before_dispatch(source_backend, monkeypatch):
    class Session:
        def snapshot(self, generation):
            return ({"name": "sessionid", "value": "synthetic-only"},)
    calls = []
    async def runner(path, options, **kwargs): calls.append(options)
    monkeypatch.setattr(backend, "threads_session_manager", lambda: Session())
    monkeypatch.setattr(backend, "run_scraper_safely", runner)
    messages = asyncio.run(websocket_start(monkeypatch, {
        "platform": "threads", "file_id": "A.xlsx", "use_threads_session": True,
        "threads_session_generation": "fixture-generation", "scrape_mode": "request",
        "use_proxy": True, "proxy_text": "socks5://fixture-user:fixture-password@proxy.example:1080",
    }))
    assert not calls
    assert any("SOCKS5" in message.get("message", "") for message in messages)
    assert "fixture-password" not in json.dumps(messages)


def test_session_generation_still_required(source_backend, monkeypatch):
    calls = []
    async def runner(path, options, **kwargs): calls.append(options)
    monkeypatch.setattr(backend, "run_scraper_safely", runner)
    messages = asyncio.run(websocket_start(monkeypatch, {"platform": "threads", "use_threads_session": True}))
    assert not calls
    assert any("cookie Threads" in message.get("message", "") for message in messages)
