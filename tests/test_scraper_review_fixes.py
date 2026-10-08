"""Regression fixtures for scraper/data-loss and per-request proxy fixes."""

import asyncio
from concurrent.futures import ThreadPoolExecutor
import io
import json
import socket
import threading
import time
import urllib.request

import openpyxl
import pytest
from playwright.async_api import async_playwright

import proxy_utils
import scraper
import threads_scraper


THREADS_URL = "https://www.threads.com/@miri_viu/post/DdRIHGWCURV"


@pytest.mark.parametrize("hydrated, delayed, has_card", [(False, False, True), (True, False, True), (True, True, True), (True, False, False)])
def test_threads_browser_uses_target_actions_and_not_caption_or_recommendations(hydrated, delayed, has_card):
    html = '''<html><body><header>miri_viu</header>
      <article><a href="/@other/post/Other">other</a>
        <div role="button"><svg title="Like"></svg>999</div>
        <div role="button"><svg title="Comment"></svg>888</div></article>
      <main><a href="/@miri_viu/post/DdRIHGWCURV">Thread 377 views</a>
      <article><a href="/@miri_viu/post/DdRIHGWCURV">miri_viu</a>
        <p>999999 views this week!</p>
        <div role="button" aria-label="Views">377 views</div>
        <div class="actions">
          <div role="button"><svg title="Like"></svg>4</div>
          <div role="button"><svg title="Comment"></svg>1</div>
          <div role="button"><svg title="Repost"></svg></div>
          <div role="button"><svg title="Share"></svg>2</div>
        </div>
      </article></main></body></html>'''

    if hydrated:
        html = html.replace('</svg>4', '</svg>').replace('</svg>1', '</svg>').replace('</svg>2', '</svg>')
        document = {"require": [["RelayPrefetchedStreamCache", "next", [], [
            "adp_BarcelonaPostPageTargetQueryRelayPreloader_fixture",
            {"__bbox": {"result": {"data": {"media": {
                "code": "DdRIHGWCURV", "user": {"username": "miri_viu"},
                "like_count": 0, "like_and_view_counts_disabled": False,
                "text_post_app_info": {"direct_reply_count": 0, "repost_count": 0, "reshare_count": None},
            }}}}},
        ]]]}
        if delayed:
            html += '<script>setTimeout(() => { const s = document.createElement("script"); s.type = "application/json"; s.textContent = ' + json.dumps(json.dumps(document)) + '; document.body.append(s); }, 600);</script>'
        else:
            html += '<script type="application/json">' + json.dumps(document) + '</script>'
        if not has_card:
            html = html.replace('/@miri_viu/post/DdRIHGWCURV', '/@other/post/Wrong')

    async def check():
        async with async_playwright() as playwright:
            browser = await playwright.chromium.launch(headless=True)
            # Route the real Chromium navigation; extraction still runs on real DOM.
            original_new_page = browser.new_page

            async def fixture_page():
                page = await original_new_page()
                await page.route("**/*", lambda route: route.fulfill(body=html, content_type="text/html"))
                return page

            browser.new_page = fixture_page
            try:
                return await threads_scraper.fetch_threads_browser(browser, THREADS_URL)
            finally:
                await browser.close()

    result = asyncio.run(check())
    assert result["error"] == ""
    if hydrated:
        assert result["metrics"] == {"views": 377 if has_card else None, "likes": 0, "comments": 0, "reposts": 0, "shares": None}
        if not has_card:
            assert result["diagnostic"] == "Không xác minh được card Threads cần quét"
    else:
        assert result["metrics"] == {"views": 377, "likes": 4, "comments": 1, "reposts": None, "shares": 2}


def test_threads_hybrid_preserves_verified_http_metrics():
    content = '<script type="application/json">' + json.dumps({"require": [
        ["BarcelonaLoggedOutExpansionGating", [], {"view_counts": 377}, 7623],
    ]}) + '</script>'
    http = threads_scraper.parse_threads_http(content, THREADS_URL, THREADS_URL)
    browser = {"channel": "miri_viu", "metrics": {"views": 999999, "likes": 4}, "error": ""}
    merged = threads_scraper.merge_results(http, browser)
    assert merged["metrics"]["views"] == 377
    assert merged["metrics"]["likes"] == 4


@pytest.mark.parametrize("content", [
    '<p>"view_counts":999999</p>',
    '<script type="application/json">{"recommended_post":{"code":"OtherPost","view_counts":999999}}</script>',
    '<script type="application/json">{"caption":"\\\"view_counts\\\":999999"}</script>',
])
def test_threads_http_rejects_views_without_typed_post_provenance(content):
    result = threads_scraper.parse_threads_http(content, THREADS_URL, THREADS_URL)
    assert result["metrics"]["views"] is None


def test_threads_http_ignores_caption_count_alongside_real_post_module():
    content = '<script type="application/json">' + json.dumps({"require": [
        ["BarcelonaLoggedOutExpansionGating", [], {"view_counts": 377}, 7623],
        {"caption": '"view_counts":999999'},
        {"recommended_post": {"code": "OtherPost", "view_counts": 444}},
    ]}) + '</script>'
    assert threads_scraper.parse_threads_http(content, THREADS_URL, THREADS_URL)["metrics"]["views"] == 377


class _NoBrowser:
    async def __aenter__(self):
        return object()

    async def __aexit__(self, *_args):
        return False


def _scan_fixture(tmp_path, monkeypatch):
    path = tmp_path / "report.xlsx"
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Data"
    sheet.append(["Link", "Tên Kênh", "LƯỢT XEM", "TIM", "BÌNH LUẬN", "LƯỢT LƯU", "CHIA SẺ"])
    sheet.append(["https://www.tiktok.com/@demo/video/123", "Demo", 10, 1, 0, 0, 0])
    book.save(path)
    book.close()
    monkeypatch.setattr(scraper, "async_playwright", _NoBrowser)
    monkeypatch.setattr(scraper, "_run_request_scrape", lambda *_a, **_kw: (
        {"Views": "999", "Likes": "10", "Comments": "1", "Saves": "2", "Shares": "3"},
        "Demo", "Success", 1, "https://www.tiktok.com/@demo/video/123",
    ))
    return path


class _CancelAfterData:
    def __init__(self):
        self.statuses = []

    async def broadcast_log(self, *_a, **_kw):
        pass

    async def broadcast_status(self, data):
        self.statuses.append(data)

    async def broadcast_data(self, data):
        assert data["views"] == 999
        raise asyncio.CancelledError()


class _CancelAfterResultLog(_CancelAfterData):
    async def broadcast_log(self, *_a, **kwargs):
        if (kwargs.get("details") or {}).get("kind") == "scrape_ok":
            raise asyncio.CancelledError()


@pytest.mark.parametrize("manager_type", [_CancelAfterData, _CancelAfterResultLog])
def test_tiktok_cancel_saves_results_already_sent_to_ui(tmp_path, monkeypatch, manager_type):
    path = _scan_fixture(tmp_path, monkeypatch)
    book = openpyxl.load_workbook(path)
    book.active.append(["TỔNG", None, "=SUM(C2:C2)"])
    book.save(path)
    book.close()
    manager = manager_type()
    with pytest.raises(asyncio.CancelledError):
        asyncio.run(scraper.run_scraper(path, manager, worker_count=1, retries=0, sheet_name="Data"))
    saved = openpyxl.load_workbook(path)
    assert saved.active["C2"].value == 999
    # The cancel save must not drop the TỔNG row the scan removed in memory.
    assert saved.active["A3"].value == "TỔNG"
    assert saved.active["C3"].value == "=SUM(C2:C2)"
    saved.close()
    assert not any(status.get("done") for status in manager.statuses)


def test_tiktok_cancel_save_failure_surfaces_and_keeps_original(tmp_path, monkeypatch):
    path = _scan_fixture(tmp_path, monkeypatch)
    original = path.read_bytes()
    manager = _CancelAfterData()

    def locked(*_args):
        raise PermissionError("fixture locked destination")

    monkeypatch.setattr(scraper.os, "replace", locked)
    with pytest.raises(RuntimeError, match="Excel"):
        asyncio.run(scraper.run_scraper(path, manager, worker_count=1, retries=0, sheet_name="Data"))
    assert path.read_bytes() == original
    assert not list(tmp_path.glob("*.tmp"))
    assert not any(status.get("done") for status in manager.statuses)


def test_request_scrapes_use_a_per_run_pool_that_is_shut_down(tmp_path, monkeypatch):
    path = _scan_fixture(tmp_path, monkeypatch)
    threads = []
    pools = []

    class TrackedPool(ThreadPoolExecutor):
        def __init__(self, *args, **kwargs):
            super().__init__(*args, **kwargs)
            pools.append(self)

    def scrape(*_args, **_kwargs):
        threads.append(threading.current_thread().name)
        return (
            {"Views": "999", "Likes": "10", "Comments": "1", "Saves": "2", "Shares": "3"},
            "Demo", "Success", 1, "https://www.tiktok.com/@demo/video/123",
        )

    monkeypatch.setattr(scraper, "ThreadPoolExecutor", TrackedPool, raising=False)
    monkeypatch.setattr(scraper, "_run_request_scrape", scrape)
    monkeypatch.setattr(scraper, "append_scrape_history", lambda *_a: None)
    asyncio.run(scraper.run_scraper(path, worker_count=1, retries=0, sheet_name="Data"))

    # Not the loop's shared default executor, which also runs workbook saves.
    assert threads and all(name.startswith("riviu-tiktok") for name in threads)
    assert len(pools) == 1 and pools[0]._max_workers == 1 and pools[0]._shutdown


def test_atomic_save_finishes_before_cancellation_returns(tmp_path, monkeypatch):
    path = _scan_fixture(tmp_path, monkeypatch)
    book = openpyxl.load_workbook(path)
    book.active["C2"].value = 1234
    entered = threading.Event()
    release = threading.Event()
    save = book.save

    def slow_save(destination):
        entered.set()
        assert release.wait(timeout=5)
        save(destination)

    monkeypatch.setattr(book, "save", slow_save)

    async def check():
        task = asyncio.create_task(scraper.save_workbook(book, path))
        await asyncio.to_thread(entered.wait, 5)
        task.cancel()
        await asyncio.sleep(0)
        assert not task.done()
        release.set()
        with pytest.raises(asyncio.CancelledError):
            await task

    try:
        asyncio.run(check())
        saved = openpyxl.load_workbook(path)
        assert saved.active["C2"].value == 1234
        saved.close()
        assert not list(tmp_path.glob("*.tmp"))
    finally:
        release.set()
        book.close()


def test_result_sheet_failure_does_not_skip_summary(tmp_path, monkeypatch):
    path = _scan_fixture(tmp_path, monkeypatch)

    def failed_result(*_a, **_kw):
        raise ValueError("fixture invalid result tab")

    def summary(book, **_kw):
        book.create_sheet("Summary rebuilt").append(["fresh"])
        return 1

    monkeypatch.setattr(scraper, "build_result_sheet", failed_result)
    monkeypatch.setattr(scraper, "rebuild_summary_sheet", summary)
    monkeypatch.setattr(scraper, "append_scrape_history", lambda *_a: None)
    asyncio.run(scraper.run_scraper(path, worker_count=1, retries=0, sheet_name="Data", create_result_sheet=True))
    saved = openpyxl.load_workbook(path)
    assert saved["Summary rebuilt"]["A1"].value == "fresh"
    saved.close()


def test_socks_and_direct_transports_are_isolated_during_concurrent_requests(monkeypatch):
    import http.client
    import socks

    original_socket = socket.socket
    original_default = socks.socksocket.default_proxy
    routes = {}
    barrier = threading.Barrier(3)

    class FakeSocket:
        def setsockopt(self, *_a):
            pass

        def sendall(self, *_a):
            pass

        def makefile(self, *_a):
            return io.BytesIO(b"HTTP/1.1 200 OK\r\nContent-Length: 2\r\n\r\nOK")

        def close(self):
            pass

    def socks_connect(address, **kwargs):
        routes.setdefault(threading.current_thread().name, []).append(kwargs.get("proxy_addr"))
        barrier.wait(timeout=5)
        return FakeSocket()

    def direct_connect(address, *_a, **_kw):
        routes.setdefault(threading.current_thread().name, []).append("direct")
        barrier.wait(timeout=5)
        return FakeSocket()

    # The real urllib opener, request and HTTP connection run. Only sockets are fake.
    monkeypatch.setattr(socks, "create_connection", socks_connect)
    monkeypatch.setattr(socket, "create_connection", direct_connect)
    configs = [proxy_utils.normalize_proxy_config({"type": "socks5", "host": host, "port": 1080})
               for host in ("proxy-a.example", "proxy-b.example")]
    proxy_utils.set_session_proxies([])

    def request_twice(config):
        name = threading.current_thread().name
        for _ in range(2):
            with proxy_utils.urlopen_with_config(urllib.request.Request("http://fixture.invalid/post"), config, timeout=1) as response:
                assert response.read() == b"OK"
        return routes[name]

    try:
        with ThreadPoolExecutor(max_workers=3) as executor:
            futures = [executor.submit(request_twice, config) for config in [*configs, None]]
            assert [future.result() for future in futures] == [
                ["proxy-a.example", "proxy-a.example"], ["proxy-b.example", "proxy-b.example"], ["direct", "direct"],
            ]
        assert socket.socket is original_socket
        assert socks.socksocket.default_proxy == original_default
    finally:
        socket.socket = original_socket
        socks.socksocket.default_proxy = original_default
        proxy_utils.set_session_proxies([])


def _threads_scan_fixture(tmp_path, monkeypatch, count=1):
    path = tmp_path / "threads.xlsx"
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Data"
    sheet.append(["Link", "Tên Kênh", "LƯỢT XEM"])
    for index in range(count):
        sheet.append([f"https://www.threads.com/@demo/post/Post{index}", "Demo", 10])
    book.save(path)
    book.close()
    monkeypatch.setattr(threads_scraper, "async_playwright", _NoBrowser)
    monkeypatch.setattr(threads_scraper, "fetch_threads_http", lambda _url: {
        "channel": "demo", "metrics": {"views": 999, "likes": None, "comments": None, "reposts": None, "shares": None}, "error": "",
    })
    return path


def test_threads_cancellation_joins_autosave_before_cleanup_save(tmp_path, monkeypatch):
    path = _threads_scan_fixture(tmp_path, monkeypatch, count=5)
    first_started = threading.Event()
    release_first = threading.Event()
    overlap_seen = threading.Event()
    save = threads_scraper.save_workbook_atomic
    lock = threading.Lock()
    active = calls = 0

    def tracked_save(book, destination):
        nonlocal active, calls
        with lock:
            active += 1
            calls += 1
            number = calls
            if active > 1:
                overlap_seen.set()
        try:
            if number == 1:
                first_started.set()
                assert release_first.wait(timeout=5)
            save(book, destination)
        finally:
            with lock:
                active -= 1

    monkeypatch.setattr(threads_scraper, "save_workbook_atomic", tracked_save)

    async def check():
        task = asyncio.create_task(threads_scraper.run_threads_scraper(path, sheet_name="Data", mode="request"))
        try:
            assert await asyncio.to_thread(first_started.wait, 5)
            task.cancel()
            # A blocked first writer forces the old implementation's cleanup
            # writer to overlap. The fixed runner waits for the owned writer.
            await asyncio.to_thread(overlap_seen.wait, 0.2)
            assert not overlap_seen.is_set()
            assert not task.done()
        finally:
            release_first.set()
            with pytest.raises(asyncio.CancelledError):
                await task

    asyncio.run(check())
    saved = openpyxl.load_workbook(path)
    assert [saved.active.cell(row, 3).value for row in range(2, 7)] == [999] * 5
    saved.close()


def _needs_browser_fallback(monkeypatch):
    """Request fails retryably, so the link goes to a (fake) browser worker that succeeds."""
    monkeypatch.setattr(scraper, "_run_request_scrape", lambda *_a, **_kw: (
        scraper.empty_metrics(), "", "Error: HTTP 500", 1, "",
    ))

    class Context:
        async def new_page(self):
            return object()

        async def close(self):
            pass

    async def make_context(browser, proxy_configs=None):
        assert browser is not None
        return Context()

    async def browser_scrape(_page, url, _retries, channel_cache=None):
        return {"Views": "999", "Likes": "10", "Comments": "1", "Saves": "2", "Shares": "3"}, "Demo", "Success", 1, url

    monkeypatch.setattr(scraper, "make_browser_context", make_context)
    monkeypatch.setattr(scraper, "scrape_with_retries", browser_scrape)
    # Skip the polite pause between browser links.
    monkeypatch.setattr(scraper.random, "uniform", lambda _low, _high: 0)


class _SlowExitPlaywright:
    """Chromium whose exit outlasts the scan, as seen live on a busy Windows host.

    A graceful browser.close() takes SLOW_EXIT seconds unless stopping the driver
    force-kills Chromium first, and the killed process then keeps the driver stop
    waiting up to SLOW_EXIT more.
    """

    SLOW_EXIT = 5.0

    def __init__(self):
        self.events = []
        self.killed = asyncio.Event()
        self.exited = asyncio.Event()
        fake = self

        class Browser:
            contexts = []

            async def close(self):
                fake.events.append("close")
                try:
                    await asyncio.wait_for(fake.killed.wait(), fake.SLOW_EXIT)
                except asyncio.TimeoutError:
                    return
                raise RuntimeError("fixture: target closed")

        class Chromium:
            async def launch(self, **_kwargs):
                return Browser()

        self.chromium = Chromium()

    def __call__(self):
        return self

    async def __aenter__(self):
        return self

    async def __aexit__(self, *_args):
        # Closing the driver's stdin makes it force-kill every browser.
        self.events.append("stop")
        self.killed.set()
        try:
            await asyncio.wait_for(self.exited.wait(), self.SLOW_EXIT)
        except asyncio.TimeoutError:
            pass


@pytest.mark.parametrize("platform", ["tiktok", "threads"])
def test_finished_scan_does_not_wait_for_slow_chromium_exit(tmp_path, monkeypatch, platform):
    monkeypatch.setattr(scraper, "BROWSER_CLOSE_TIMEOUT", 0.2)
    monkeypatch.setattr(scraper, "PLAYWRIGHT_STOP_TIMEOUT", 0.2)
    fake = _SlowExitPlaywright()
    if platform == "tiktok":
        path = _scan_fixture(tmp_path, monkeypatch)
        # Hybrid starts Chromium only for a fallback, so make this link need one.
        _needs_browser_fallback(monkeypatch)
        monkeypatch.setattr(scraper, "async_playwright", fake)
        scan = lambda: scraper.run_scraper(path, worker_count=1, retries=0, sheet_name="Data",
                                           use_request=True, browser_fallback=True)
    else:
        path = _threads_scan_fixture(tmp_path, monkeypatch)
        monkeypatch.setattr(threads_scraper, "async_playwright", fake)

        async def browser_result(_browser, _url):
            return {"channel": "demo", "metrics": {"views": 999, "likes": 1, "comments": 0, "reposts": 0, "shares": 0}, "error": ""}

        monkeypatch.setattr(threads_scraper, "fetch_threads_browser", browser_result)
        scan = lambda: threads_scraper.run_threads_scraper(path, sheet_name="Data", mode="hybrid")

    async def check():
        started = time.monotonic()
        await scan()
        # Old code awaited both slow exits in full (~2 x SLOW_EXIT) before returning.
        assert time.monotonic() - started < 2.5
        # The kill was requested before the scan reported completion.
        assert fake.events == ["close", "stop"]
        assert scraper.browser_cleanup_pending()
        fake.exited.set()
        for _ in range(100):
            if not scraper.browser_cleanup_pending():
                break
            await asyncio.sleep(0.01)
        assert not scraper.browser_cleanup_pending()

    asyncio.run(check())
    saved = openpyxl.load_workbook(path)
    try:
        assert saved.active["C2"].value == 999
    finally:
        saved.close()


def test_threads_failed_final_save_still_closes_browser_and_workbook(tmp_path, monkeypatch):
    path = _threads_scan_fixture(tmp_path, monkeypatch)
    original = path.read_bytes()
    closed = {"browser": False, "workbook": False}

    class Browser:
        async def close(self):
            closed["browser"] = True

    class Chromium:
        async def launch(self, **_kwargs):
            return Browser()

    class Playwright(_NoBrowser):
        chromium = Chromium()

        async def __aenter__(self):
            return self

    async def browser_result(_browser, _url):
        return {"channel": "demo", "metrics": {"views": 999}, "error": ""}

    def failed_save(*_args):
        raise PermissionError("fixture locked destination")

    close = openpyxl.workbook.workbook.Workbook.close

    def track_close(book):
        closed["workbook"] = True
        close(book)

    monkeypatch.setattr(threads_scraper, "async_playwright", Playwright)
    monkeypatch.setattr(threads_scraper, "fetch_threads_browser", browser_result)
    monkeypatch.setattr(threads_scraper, "save_workbook_atomic", failed_save)
    monkeypatch.setattr(openpyxl.workbook.workbook.Workbook, "close", track_close)
    with pytest.raises(PermissionError):
        asyncio.run(threads_scraper.run_threads_scraper(path, sheet_name="Data", mode="browser"))
    assert path.read_bytes() == original
    assert closed == {"browser": True, "workbook": True}


def test_hybrid_does_not_send_confirmed_unavailable_pages_to_the_browser(monkeypatch):
    import asyncio
    import scraper

    statuses = {"https://www.tiktok.com/@a/video/1": "Error: Trang TikTok không khả dụng",
                "https://www.tiktok.com/@a/video/2": "Error: Không đọc được số liệu"}
    monkeypatch.setattr(scraper, "_run_request_scrape",
                        lambda _index, url, **_kwargs: (scraper.empty_metrics(), "", statuses[url], 1, ""))

    async def scenario():
        scrape_queue, browser_queue, result_queue = asyncio.Queue(), asyncio.Queue(), asyncio.Queue()
        for url in statuses:
            scrape_queue.put_nowait({"url": url, "row": 2})
        scrape_queue.put_nowait(None)
        await scraper.request_worker_loop(1, scrape_queue, browser_queue, result_queue, retries=1, browser_fallback=True)
        return [item["url"] for item in browser_queue._queue], [item["url"] for item in result_queue._queue]

    to_browser, finished = asyncio.run(scenario())
    # A deleted/unavailable page is final in Request mode; only unreadable pages get the slow browser retry.
    assert finished == ["https://www.tiktok.com/@a/video/1"]
    assert to_browser == ["https://www.tiktok.com/@a/video/2"]


def _rows_fixture(tmp_path, count):
    path = tmp_path / "rows.xlsx"
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Data"
    sheet.append(["Link", "Tên Kênh", "LƯỢT XEM", "TIM", "BÌNH LUẬN", "LƯỢT LƯU", "CHIA SẺ"])
    for index in range(count):
        sheet.append([f"https://www.tiktok.com/@demo/video/{100 + index}", "Demo", 10, 1, 0, 0, 0])
    book.save(path)
    book.close()
    return path


def _success_for_url(_index, url, **_kwargs):
    return {"Views": "999", "Likes": "10", "Comments": "1", "Saves": "2", "Shares": "3"}, "Demo", "Success", 1, url


def test_total_row_placement_does_constant_whole_sheet_scans(monkeypatch):
    # max_row/max_column walk every cell; one walk per row made TỔNG placement
    # quadratic (minutes on a 3,674 x 68 workbook, all of it on the event loop).
    from openpyxl.worksheet.worksheet import Worksheet

    scans = []
    for name in ("max_row", "max_column"):
        getter = getattr(Worksheet, name).fget

        def counted(sheet, _getter=getter):
            scans.append(1)
            return _getter(sheet)

        monkeypatch.setattr(Worksheet, name, property(counted))

    def scans_for(rows):
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "Data"
        sheet.append(["Link", "Tên Kênh", "LƯỢT XEM", "TIM", "BÌNH LUẬN", "LƯỢT LƯU", "CHIA SẺ", *[f"X{i}" for i in range(30)]])
        for index in range(rows):
            sheet.append([f"https://www.tiktok.com/@demo/video/{index + 1}", "Demo", 10, 1, 0, 0, 0])
        sheet.append(["Ghi chú", "footer"])
        # Styled but empty rows below the data must not push TỔNG further down.
        sheet.cell(row=rows + 5, column=3).fill = openpyxl.styles.PatternFill("solid", fgColor="FFFF00")
        scans.clear()
        scraper.append_sheet_total_rows(book, sheet_name="Data")
        assert sheet.cell(row=rows + 3, column=1).value == "TỔNG"
        assert sheet.cell(row=rows + 3, column=3).value == f"=SUM(C2:C{rows + 1})"
        return len(scans)

    small, large = scans_for(50), scans_for(400)
    assert small == large
    assert large < 100  # the old per-row walk made 400+ here


def test_scan_setup_and_finish_keep_the_event_loop_responsive(tmp_path, monkeypatch):
    path = _rows_fixture(tmp_path, 3)
    monkeypatch.setattr(scraper, "async_playwright", _NoBrowser)
    monkeypatch.setattr(scraper, "_run_request_scrape", _success_for_url)
    clear_totals, rebuild_summary = scraper.clear_existing_total_rows, scraper.rebuild_summary_sheet

    # Stand-ins for a large workbook's slow setup and finish steps.
    def slow_clear(*args, **kwargs):
        time.sleep(0.6)
        return clear_totals(*args, **kwargs)

    def slow_summary(*args, **kwargs):
        time.sleep(0.6)
        return rebuild_summary(*args, **kwargs)

    monkeypatch.setattr(scraper, "clear_existing_total_rows", slow_clear)
    monkeypatch.setattr(scraper, "rebuild_summary_sheet", slow_summary)

    async def check():
        lags = []
        scan = asyncio.create_task(scraper.run_scraper(path, worker_count=1, retries=0, sheet_name="Data", base_dir=tmp_path))
        while not scan.done():
            started = time.perf_counter()
            await asyncio.sleep(0.01)
            lags.append(time.perf_counter() - started)
        await scan
        return max(lags)

    assert asyncio.run(check()) < 0.3
    saved = openpyxl.load_workbook(path)
    try:
        assert [saved["Data"].cell(row, 3).value for row in range(2, 5)] == [999] * 3
        assert saved["Data"]["A5"].value == "TỔNG"
    finally:
        saved.close()


@pytest.mark.parametrize("links, expected_saves", [(3, 1), (5, 1), (7, 2)])
def test_completed_scan_writes_the_workbook_once_at_the_end(tmp_path, monkeypatch, links, expected_saves):
    # links=5: an autosave due on the last result would only repeat the final save.
    # links=7: the autosave at 5/7 stays for crash safety, then one final save.
    path = _rows_fixture(tmp_path, links)
    monkeypatch.setattr(scraper, "async_playwright", _NoBrowser)
    monkeypatch.setattr(scraper, "_run_request_scrape", _success_for_url)
    saves = []
    save = scraper.save_workbook_atomic

    def counted_save(workbook, destination):
        saves.append(destination)
        return save(workbook, destination)

    monkeypatch.setattr(scraper, "save_workbook_atomic", counted_save)
    asyncio.run(scraper.run_scraper(path, worker_count=1, retries=0, save_every=5, sheet_name="Data", base_dir=tmp_path))
    assert len(saves) == expected_saves
    saved = openpyxl.load_workbook(path)
    try:
        assert [saved["Data"].cell(row, 3).value for row in range(2, links + 2)] == [999] * links
        assert saved["Data"].cell(links + 2, 1).value == "TỔNG"
    finally:
        saved.close()


def test_scan_stopped_during_cleanup_after_last_result_still_saves(tmp_path, monkeypatch):
    path = _rows_fixture(tmp_path, 2)
    monkeypatch.setattr(scraper, "_run_request_scrape", _success_for_url)
    stopping = asyncio.Event()

    class StuckStop:
        async def __aenter__(self):
            return object()

        async def __aexit__(self, *_args):
            stopping.set()
            await asyncio.sleep(5)

    monkeypatch.setattr(scraper, "async_playwright", StuckStop)
    monkeypatch.setattr(scraper, "PLAYWRIGHT_STOP_TIMEOUT", 5)

    async def check():
        task = asyncio.create_task(scraper.run_scraper(path, worker_count=1, retries=0, sheet_name="Data", base_dir=tmp_path))
        await stopping.wait()
        task.cancel()
        with pytest.raises(asyncio.CancelledError):
            await task

    asyncio.run(check())
    saved = openpyxl.load_workbook(path)
    try:
        assert [saved["Data"].cell(row, 3).value for row in (2, 3)] == [999, 999]
        assert saved["Data"]["A4"].value == "TỔNG"
    finally:
        saved.close()


class _CountingPlaywright:
    def __init__(self):
        self.launches = 0
        fake = self

        class Browser:
            contexts = []

            async def close(self):
                pass

        class Chromium:
            async def launch(self, **_kwargs):
                fake.launches += 1
                await asyncio.sleep(0.05)
                return Browser()

        self.chromium = Chromium()

    def __call__(self):
        return self

    async def __aenter__(self):
        return self

    async def __aexit__(self, *_args):
        return False


def test_hybrid_without_fallback_never_launches_chromium(tmp_path, monkeypatch):
    path = _rows_fixture(tmp_path, 4)
    fake = _CountingPlaywright()
    monkeypatch.setattr(scraper, "async_playwright", fake)
    monkeypatch.setattr(scraper, "_run_request_scrape", _success_for_url)
    asyncio.run(scraper.run_scraper(path, worker_count=4, retries=0, sheet_name="Data", base_dir=tmp_path,
                                    use_request=True, browser_fallback=True))
    assert fake.launches == 0
    saved = openpyxl.load_workbook(path)
    try:
        assert [saved["Data"].cell(row, 3).value for row in range(2, 6)] == [999] * 4
    finally:
        saved.close()


def test_hybrid_fallback_workers_share_one_lazily_launched_chromium(tmp_path, monkeypatch):
    path = _rows_fixture(tmp_path, 4)
    fake = _CountingPlaywright()
    monkeypatch.setattr(scraper, "async_playwright", fake)
    _needs_browser_fallback(monkeypatch)
    asyncio.run(scraper.run_scraper(path, worker_count=4, retries=0, sheet_name="Data", base_dir=tmp_path,
                                    use_request=True, browser_fallback=True))
    # Four fallback links reach three browser workers; they start one Chromium together.
    assert fake.launches == 1
    saved = openpyxl.load_workbook(path)
    try:
        assert [saved["Data"].cell(row, 3).value for row in range(2, 6)] == [999] * 4
    finally:
        saved.close()
