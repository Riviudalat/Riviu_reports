"""Interrupted scans continue, transient failures get one retry pass, cookie sessions are checked first.

Synthetic workbooks in tmp_path and stubbed scanners/fetchers only: no network, no live scans.
See riviu/scan_journal.py and riviu/scan_service.py.
"""

import asyncio
import json
import os
import threading
import time

import openpyxl
import pytest

from riviu import app as backend
from riviu import scan_journal, scan_service, scan_state, server_context, server_mode
from riviu.platforms import tiktok as scraper
from riviu.platforms import tiktok_session
from riviu.platforms.tiktok.statuses import STATUS_TIKTOK_NOT_FOUND_UNCONFIRMED, STATUS_TIKTOK_NO_STATS

LINKS = [f"https://www.tiktok.com/@demo/video/{index}" for index in range(1, 5)]
SECRET_PROXY = "http://proxyuser:proxy-secret-pw@proxy.example:8080"


@pytest.fixture
def env(tmp_path, monkeypatch):
    monkeypatch.setattr(backend, "EXCEL_DIR", str(tmp_path))
    for name, value in (("SCRAPE_TASK", None), ("SCAN_STARTING", False), ("SOURCE_BUSY", False)):
        monkeypatch.setattr(scan_state, name, value)
    monkeypatch.setattr(scan_state, "manager", scan_state.ConnectionManager())
    monkeypatch.setattr(scan_state, "USER_CANCELLED_RUNS", set())
    monkeypatch.setattr(backend, "CURRENT_SELECTED_FILE", "")
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Data"
    sheet.append(["Link", "Tên Kênh", "Đối tác", "LƯỢT XEM", "TIM", "BÌNH LUẬN", "LƯỢT LƯU", "CHIA SẺ"])
    for link in LINKS:
        sheet.append([link, "Demo", "Partner", 7, 7, 7, 7, 7])
    book.save(tmp_path / "Book.xlsx")
    book.close()
    return tmp_path


def logs_of(monkeypatch):
    logs = []
    original = scan_state.manager.broadcast_log

    async def log(message, **kwargs):
        logs.append(message)
        await original(message, **kwargs)

    monkeypatch.setattr(scan_state.manager, "broadcast_log", log)
    return logs


def journal_file(base):
    return base / "data" / scan_journal.JOURNAL_FILENAME


def options(**overrides):
    return scan_service.scan_options(**{"platform": "tiktok", "file_id": "Book.xlsx", "sheet_name": "Data", **overrides})


def key(index, sheet="Data"):
    return (sheet, index + 2, LINKS[index])


def saved(index, status, kind=""):
    sheet, row, url = key(index)
    return {"sheet_name": sheet, "row": row, "url": url, "status": status, "kind": kind}


def write_journal(base, **fields):
    data = {
        "version": 1, "journalId": "fixture", "platform": "tiktok", "fileId": "Book.xlsx", "sheetName": "Data",
        "startedAt": "", "scheduled": False, "proxyFingerprint": None, "resumeAttempts": 0, "mainDone": False,
        "stopReason": "", "rows": [],
        "payload": {"platform": "tiktok", "file_id": "Book.xlsx", "sheet_name": "Data", "workers": 1, "scrape_mode": "request",
                    "use_proxy": False, "partners": [], "partner": "", "create_result_sheet": False,
                    "push_to_google": True, "retry_failed": False},
        **fields,
    }
    os.makedirs(base / "data", exist_ok=True)
    journal_file(base).write_text(json.dumps(data), encoding="utf-8")
    return data


# ---- Journal ---------------------------------------------------------------------------------

def test_journal_records_saved_rows_without_secrets_and_is_removed_when_the_run_ends(env, monkeypatch):
    seen = {}

    async def runner(target_path, events, opts):
        await opts["on_rows_saved"]([saved(0, "Success"), saved(1, STATUS_TIKTOK_NO_STATS)])
        seen["journal"] = journal_file(env).read_text(encoding="utf-8")
        await events.broadcast_status({"total": 4, "processed": 4, "success": 1, "hidden": 1, "error": 0, "done": True})

    monkeypatch.setattr(scan_service, "SCAN_RUNNERS", {**scan_service.SCAN_RUNNERS, "tiktok": runner})
    asyncio.run(scan_service.run_scraper_safely(str(env / "Book.xlsx"), options(use_proxy=True, proxy_text=SECRET_PROXY, retry_failed=False)))

    journal = json.loads(seen["journal"])
    assert journal["rows"] == [["Data", 2, LINKS[0], "s", False], ["Data", 3, LINKS[1], "h", False]]
    assert journal["payload"]["use_proxy"] is True and "proxy_text" not in journal["payload"]
    assert journal["proxyFingerprint"]["hash"]
    assert "proxy-secret-pw" not in seen["journal"] and "proxyuser" not in seen["journal"]
    assert not journal_file(env).exists()
    assert scan_state.manager.last_status["phase"] == "completed"


# ---- Resume after a restart ------------------------------------------------------------------

def fake_tiktok(monkeypatch, calls, block_url=None, blocked=None):
    class NoPlaywright:
        async def __aenter__(self):
            return object()

        async def __aexit__(self, *_args):
            return False

    def scrape(_worker_index, url, *_args, **_kwargs):
        calls.append(url)
        if url == block_url and blocked is not None:
            blocked.set()
            time.sleep(0.5)  # still loading when the server stops
        return ({"Views": "100", "Likes": "10", "Comments": "1", "Saves": "2", "Shares": "3"}, "Demo", "Success", 1, url)

    monkeypatch.setattr(scraper.runner, "async_playwright", lambda: NoPlaywright())
    monkeypatch.setattr(scraper.workers, "_run_request_scrape", scrape)


def test_a_scan_stopped_by_a_restart_resumes_with_only_the_unsaved_rows(env, monkeypatch):
    monkeypatch.setattr(server_mode, "ACTIVE", object())  # the admin hint in messages; resume is server-only
    calls = []
    blocked = threading.Event()
    fake_tiktok(monkeypatch, calls, block_url=LINKS[2], blocked=blocked)

    async def first_process():
        target, opts = await scan_service.prepare_scan_start({
            "platform": "tiktok", "file_id": "Book.xlsx", "sheet_name": "Data", "workers": 1, "push_to_google": False})
        task, _ = scan_service.start_scan_task(target, opts)
        while not blocked.is_set():
            assert not task.done(), scan_state.manager.last_status
            await asyncio.sleep(0.01)
        task.cancel()  # server shutdown: not a user cancel
        await asyncio.gather(task, return_exceptions=True)

    asyncio.run(first_process())
    assert calls == LINKS[:3]
    journal = json.loads(journal_file(env).read_text(encoding="utf-8"))
    assert [row[:3] for row in journal["rows"]] == [list(key(0)), list(key(1))]

    calls.clear()
    monkeypatch.setattr(scan_state, "manager", scan_state.ConnectionManager())
    monkeypatch.setattr(scan_state, "SCRAPE_TASK", None)

    async def second_process():
        task = await scan_service.resume_interrupted_scan()
        assert task is not None
        await task

    asyncio.run(second_process())
    assert calls == LINKS[2:]  # the two saved rows are not scanned again
    assert not journal_file(env).exists()
    final = scan_state.manager.last_status
    assert final["phase"] == "completed" and final["total"] == 4 and final["success"] == 4
    book = openpyxl.load_workbook(env / "Book.xlsx")
    try:
        assert [book["Data"].cell(row, 5).value for row in range(2, 6)] == [10, 10, 10, 10]  # TIM
    finally:
        book.close()


def test_a_resumed_scheduled_run_still_never_pushes(env, monkeypatch):
    write_journal(env, scheduled=True)
    seen = []

    async def runner(target_path, events, opts):
        seen.append(opts["push_to_google"])

    monkeypatch.setattr(scan_service, "SCAN_RUNNERS", {**scan_service.SCAN_RUNNERS, "tiktok": runner})
    monkeypatch.setattr(scan_service, "push_rows_to_new_sheet", lambda *_a, **_k: pytest.fail("never push"))

    async def go():
        await (await scan_service.resume_interrupted_scan())

    asyncio.run(go())
    assert seen == [False]


def test_user_cancel_drops_the_journal_and_nothing_resumes(env, monkeypatch):
    async def runner(target_path, events, opts):
        await opts["on_rows_saved"]([saved(0, "Success")])
        await asyncio.Event().wait()

    monkeypatch.setattr(scan_service, "SCAN_RUNNERS", {**scan_service.SCAN_RUNNERS, "tiktok": runner})

    async def go():
        task, _ = scan_service.start_scan_task(str(env / "Book.xlsx"), options())
        while not journal_file(env).exists() or "Data" not in journal_file(env).read_text(encoding="utf-8"):
            await asyncio.sleep(0.01)
        assert scan_service.user_cancel_scan()
        await asyncio.gather(task, return_exceptions=True)
        return await scan_service.resume_interrupted_scan()

    assert asyncio.run(go()) is None
    assert not journal_file(env).exists()


def test_a_threads_429_stop_is_never_resumed(env, monkeypatch):
    logs = logs_of(monkeypatch)
    write_journal(env, platform="threads", stopReason="rate_limit")
    assert asyncio.run(scan_service.resume_interrupted_scan()) is None
    assert not journal_file(env).exists() and any("HTTP 429" in log for log in logs)


def test_resume_attempts_are_capped(env, monkeypatch):
    starts = []

    async def runner(target_path, events, opts):
        starts.append(json.loads(journal_file(env).read_text(encoding="utf-8"))["resumeAttempts"])

    monkeypatch.setattr(scan_service, "SCAN_RUNNERS", {**scan_service.SCAN_RUNNERS, "tiktok": runner})

    async def go():
        task = await scan_service.resume_interrupted_scan()
        if task is not None:
            await task
        return task

    write_journal(env, resumeAttempts=scan_journal.MAX_RESUME_ATTEMPTS - 1)
    assert asyncio.run(go()) is not None
    # Counted on disk before the run starts, so a crash inside it uses the attempt up.
    assert starts == [scan_journal.MAX_RESUME_ATTEMPTS]
    write_journal(env, resumeAttempts=scan_journal.MAX_RESUME_ATTEMPTS)
    assert asyncio.run(go()) is None
    assert starts == [scan_journal.MAX_RESUME_ATTEMPTS] and not journal_file(env).exists()


def test_resume_is_refused_when_the_saved_proxy_list_is_another_route(env, monkeypatch):
    salt = "s"
    write_journal(env, payload={**write_journal(env)["payload"], "use_proxy": True},
                  proxyFingerprint={"salt": salt, "hash": scan_journal.fingerprint([{"host": "old"}], salt)})
    (env / "proxy_list.txt").write_text("new.example:8080\n", encoding="utf-8")
    monkeypatch.setattr(scan_service, "proxy_route", lambda *_args: [{"host": "new"}])
    monkeypatch.setattr(scan_service, "validate_proxy_start", lambda *_args, **_kwargs: None)
    monkeypatch.setattr(scan_service, "start_scan_task", lambda *_a: pytest.fail("never on another route"))
    assert asyncio.run(scan_service.resume_interrupted_scan()) is None
    assert not journal_file(env).exists()


# ---- Stall watchdog --------------------------------------------------------------------------

def stall_clock(monkeypatch):
    now = [0.0]
    monkeypatch.setattr(scan_service, "_clock", lambda: now[0])
    monkeypatch.setattr(scan_journal, "STALL_CHECK_SECONDS", 0.001)
    return now


def test_stall_watchdog_cancels_the_stuck_pass_and_continues_with_unsaved_rows(env, monkeypatch):
    now = stall_clock(monkeypatch)
    passes = []

    async def runner(target_path, events, opts):
        passes.append(opts.get("skip_rows"))
        if len(passes) == 1:
            await opts["on_rows_saved"]([saved(0, "Success")])
            now[0] += scan_journal.STALL_SECONDS + 1  # no progress for longer than the threshold
            await asyncio.Event().wait()
        await opts["on_rows_saved"]([saved(index, "Success") for index in (1, 2, 3)])
        await events.broadcast_status({"total": 3, "processed": 3, "success": 3, "hidden": 0, "error": 0, "done": True})

    monkeypatch.setattr(scan_service, "SCAN_RUNNERS", {**scan_service.SCAN_RUNNERS, "tiktok": runner})
    logs = logs_of(monkeypatch)
    asyncio.run(scan_service.run_scraper_safely(str(env / "Book.xlsx"), options()))

    assert passes == [None, {key(0)}]
    assert any("không có kết quả mới" in log and "1/2" in log for log in logs)
    final = scan_state.manager.last_status
    assert final["phase"] == "completed" and final["total"] == 4 and final["success"] == 4
    assert not journal_file(env).exists()


def test_a_scan_that_keeps_stalling_fails_after_the_resume_cap(env, monkeypatch):
    now = stall_clock(monkeypatch)
    passes = []

    async def runner(target_path, events, opts):
        passes.append(opts.get("skip_rows"))
        now[0] += scan_journal.STALL_SECONDS + 1
        await asyncio.Event().wait()

    monkeypatch.setattr(scan_service, "SCAN_RUNNERS", {**scan_service.SCAN_RUNNERS, "tiktok": runner})
    asyncio.run(scan_service.run_scraper_safely(str(env / "Book.xlsx"), options()))
    assert len(passes) == scan_journal.MAX_RESUME_ATTEMPTS + 1
    final = scan_state.manager.last_status
    assert final["phase"] == "failed" and "dừng hẳn" in final["message"]
    assert not journal_file(env).exists()


# ---- Retry pass ------------------------------------------------------------------------------

@pytest.mark.parametrize("status, retryable", [
    (STATUS_TIKTOK_NOT_FOUND_UNCONFIRMED, True),
    ("Error: HTTP 404", True),
    ("Error: <urlopen error [Errno 11001] getaddrinfo failed>", True),
    ("Error: [SSL: UNEXPECTED_EOF_WHILE_READING] EOF occurred in violation of protocol", True),
    ("Error: HTTP 503", True),
    ("Error: HTTP 400", True),
    ("Error: Worker đã dừng", True),
    ("Error: HTTP 403", False),
    ("Error: HTTP 429", False),
    ("Error: Trang TikTok không khả dụng", False),  # confirmed gone
    ("Error: Redirect không khớp video", False),
    (STATUS_TIKTOK_NO_STATS, False),
    ("Success", False),
])
def test_tiktok_retryable_statuses(status, retryable):
    assert scan_journal.tiktok_retryable(status) is retryable


def run_with_passes(env, monkeypatch, platform, main_rows, retry_rows, main_status=None, cooldown=None, **option_overrides):
    """retry_rows: the rows every retry pass saves (a list), or a function of that pass's only_rows."""
    passes, waits = [], []

    async def runner(target_path, events, opts):
        passes.append({name: opts.get(name) for name in ("skip_rows", "only_rows", "create_result_sheet", "threads_cookies")})
        rows = main_rows if len(passes) == 1 else retry_rows(opts["only_rows"]) if callable(retry_rows) else retry_rows
        await opts["on_rows_saved"](rows)
        await events.broadcast_status({"total": len(rows), "processed": len(rows), "done": True, **(main_status or {})})

    async def record_wait(seconds):
        waits.append(seconds)
        if cooldown:
            await cooldown(seconds)

    monkeypatch.setattr(scan_service, "SCAN_RUNNERS", {**scan_service.SCAN_RUNNERS, platform: runner})
    monkeypatch.setattr(scan_service, "_cooldown_sleep", record_wait)
    asyncio.run(scan_service.run_scraper_safely(str(env / "Book.xlsx"), options(platform=platform, **option_overrides)))
    return passes, waits


def test_retry_pass_rescans_only_rows_with_a_transient_tiktok_error(env, monkeypatch):
    logs = logs_of(monkeypatch)
    main = [saved(0, "Success"), saved(1, STATUS_TIKTOK_NOT_FOUND_UNCONFIRMED), saved(2, "Error: HTTP 403"), saved(3, "Error: HTTP 503")]
    passes, waits = run_with_passes(env, monkeypatch, "tiktok", main, [saved(1, "Success"), saved(3, "Error: HTTP 503")],
                                    create_result_sheet=True)
    assert len(passes) == 2
    assert passes[1] == {"skip_rows": None, "only_rows": {key(1), key(3)}, "create_result_sheet": False, "threads_cookies": None}
    assert waits == [scan_journal.RETRY_COOLDOWN_SECONDS["tiktok"]]
    assert any("Quét lại 2 link lỗi" in log for log in logs)
    assert any("1/2 link đọc được" in log for log in logs)
    final = scan_state.manager.last_status
    assert (final["total"], final["success"], final["error"], final["phase"]) == (4, 2, 2, "completed")


def test_retry_pass_can_be_switched_off(env, monkeypatch):
    passes, waits = run_with_passes(env, monkeypatch, "tiktok", [saved(0, "Error: HTTP 503")], [], retry_failed=False)
    assert len(passes) == 1 and waits == []


INVALID_POST = "Error: Link Threads không khả dụng (invalid_post)"
SESSION_COOKIES = ({"name": "sessionid", "value": "synthetic-threads-cookie", "domain": ".threads.com"},)


class FakeThreadsSession:
    """The stored Cookie Threads session as scan_service sees it (no store, no network)."""
    def __init__(self, configured=True, state="valid"):
        self.configured, self.state, self.invalidated = configured, state, []

    def status(self):
        return {"configured": self.configured, "state": self.state, "generation": "gen-1" if self.configured else "",
                "message": ""}

    def snapshot(self, generation=None):
        return SESSION_COOKIES

    def invalidate(self, generation, state="unknown"):
        self.invalidated.append((generation, state))


def threads_session(monkeypatch, **state):
    session = FakeThreadsSession(**state)
    monkeypatch.setattr(server_context, "threads_session_manager", lambda: session)
    return session


def test_threads_retry_pass_waits_out_a_route_429_and_skips_blocks(env, monkeypatch):
    threads_session(monkeypatch, configured=False)
    main = [saved(0, INVALID_POST, "invalid_post"),
            saved(1, "Error: Bài Threads bị giới hạn theo vị trí hiện tại", "geo_restricted"),
            saved(2, "Error: HTTP 429", "rate_limit"), saved(3, "Success")]
    passes, waits = run_with_passes(env, monkeypatch, "threads", main, [saved(2, "Success")])
    # invalid_post is never retried anonymously; without a cookie session it is reported instead.
    assert len(passes) == 2 and passes[1]["only_rows"] == {key(2)}
    # A 429 stretches the Threads cooldown to the owner's 10 minutes (not 15: a scheduled slot waits on it).
    assert waits == [10 * 60] and scan_service.THREADS_RATE_LIMIT_WAIT_MINUTES == 10
    assert scan_state.manager.last_status["loginRequired"] == 1


def test_threads_login_only_posts_without_a_cookie_end_the_run_at_once_with_a_hint(env, monkeypatch):
    threads_session(monkeypatch, state="expired")
    logs = logs_of(monkeypatch)
    main = [saved(0, INVALID_POST, "invalid_post"), saved(1, INVALID_POST, "invalid_post"), saved(2, "Success")]
    passes, waits = run_with_passes(env, monkeypatch, "threads", main, [])
    assert len(passes) == 1 and waits == []
    final = scan_state.manager.last_status
    assert (final["phase"], final["loginRequired"]) == ("completed", 2)
    assert final["notice"].startswith("2 bài Threads chỉ xem được khi đăng nhập")
    assert "Cookie Threads" in final["notice"] and final["notice"] in logs
    assert not journal_file(env).exists()


def test_threads_login_only_posts_are_retried_through_the_cookie_session_after_a_short_wait(env, monkeypatch):
    session = threads_session(monkeypatch)
    main = [saved(0, INVALID_POST, "invalid_post"), saved(1, "Error: timeout", "transport"), saved(2, "Success")]
    by_pass = {frozenset({key(1)}): [saved(1, "Success")], frozenset({key(0)}): [saved(0, "Success")]}
    passes, waits = run_with_passes(env, monkeypatch, "threads", main, lambda only: by_pass[frozenset(only)])
    # Transient rows first (anonymous, as before), then the login-only rows on the stored session.
    assert [(p["only_rows"], p["threads_cookies"]) for p in passes[1:]] == [({key(1)}, None), ({key(0)}, SESSION_COOKIES)]
    assert waits == [scan_journal.RETRY_COOLDOWN_SECONDS["threads"]]
    final = scan_state.manager.last_status
    assert (final["success"], final["error"], "notice" in final) == (3, 0, False)
    assert session.invalidated == []

    passes, waits = run_with_passes(env, monkeypatch, "threads", main[:1] + main[2:], [saved(0, "Success")])
    assert waits == [scan_journal.LOGIN_RETRY_COOLDOWN_SECONDS] and passes[1]["threads_cookies"] == SESSION_COOKIES


def test_threads_login_only_posts_are_not_retried_when_the_run_already_used_the_cookie(env, monkeypatch):
    threads_session(monkeypatch)
    main = [saved(0, INVALID_POST, "invalid_post"), saved(1, "Success")]
    passes, waits = run_with_passes(env, monkeypatch, "threads", main, [], threads_cookies=SESSION_COOKIES)
    assert len(passes) == 1 and waits == []
    final = scan_state.manager.last_status
    assert final["phase"] == "completed" and "notice" not in final and "loginRequired" not in final


def test_a_cookie_rejected_in_the_login_retry_keeps_the_run_completed(env, monkeypatch):
    session = threads_session(monkeypatch)

    def rejected(only):
        raise scan_service.SessionError("invalid")

    passes, waits = run_with_passes(env, monkeypatch, "threads", [saved(0, INVALID_POST, "invalid_post")], rejected)
    final = scan_state.manager.last_status
    assert final["phase"] == "completed" and final["error"] == 1
    assert session.invalidated == [("gen-1", "invalid")]
    assert not journal_file(env).exists()


def test_skip_retry_ends_the_cooldown_as_a_completed_run_with_the_main_pass_results(env, monkeypatch):
    async def admin_skips(seconds):
        assert scan_service.skip_retry_wait() and not scan_service.skip_retry_wait()
        await asyncio.Event().wait()  # the cooldown itself would never end

    main = [saved(0, "Success"), saved(1, "Error: HTTP 503")]
    passes, waits = run_with_passes(env, monkeypatch, "tiktok", main, [saved(1, "Success")], cooldown=admin_skips)
    assert len(passes) == 1 and len(waits) == 1
    final = scan_state.manager.last_status
    assert (final["phase"], final.get("cancelled"), final["success"], final["error"]) == ("completed", None, 1, 1)
    assert "retryPass" not in final and "retryWaitSeconds" not in final
    assert scan_state.RETRY_SKIP is None and not scan_service.skip_retry_wait()
    assert not journal_file(env).exists()


def test_no_retry_pass_after_a_threads_429_stop(env, monkeypatch):
    main = [saved(0, "Error: HTTP 429", "rate_limit"), saved(1, "Error: Link Threads không khả dụng (invalid_post)", "invalid_post")]
    passes, waits = run_with_passes(env, monkeypatch, "threads", main, [], main_status={"stopReason": "rate_limit"})
    assert len(passes) == 1 and waits == []
    assert scan_state.manager.last_status["stopReason"] == "rate_limit"
    assert not journal_file(env).exists()


# ---- Cookie check before scanning ------------------------------------------------------------

class MemoryStore:
    def __init__(self, cookies=None):
        self.saved = cookies

    def load(self):
        return self.saved

    def save(self, cookies):
        self.saved = cookies

    def delete(self):
        self.saved = None


def tiktok_cookie(value="synthetic-session-value"):
    return [{"domain": ".tiktok.com", "name": "sessionid", "value": value, "path": "/",
             "expirationDate": time.time() + 86400, "secure": True, "httpOnly": True, "hostOnly": False}]


def test_tiktok_cookie_is_verified_and_an_expired_login_fails_before_scanning(env, monkeypatch):
    monkeypatch.setattr(server_mode, "ACTIVE", object())
    store = MemoryStore(tiktok_session.normalize_tiktok_cookies(tiktok_cookie()))
    monkeypatch.setattr(server_context, "tiktok_session_manager",
                        lambda: tiktok_session.TikTokSession("unused", store=store, verifier=lambda _cookies: "invalid"))
    with pytest.raises(ValueError) as caught:
        asyncio.run(scan_service.prepare_scan_start({"platform": "tiktok", "file_id": "Book.xlsx", "use_tiktok_session": True}))
    message = str(caught.value)
    assert "Cookie TikTok đã hết hạn hoặc chưa nạp" in message and "/admin › Cookie TikTok" in message
    assert "synthetic-session-value" not in message
    assert not scan_state.SCAN_STARTING


def test_threads_cookie_missing_fails_fast_with_the_admin_hint(env, monkeypatch):
    monkeypatch.setattr(server_mode, "ACTIVE", object())
    from riviu.platforms.threads_session import ThreadsSession

    session = ThreadsSession("unused", store=MemoryStore())
    monkeypatch.setattr(server_context, "threads_session_manager", lambda: session)
    with pytest.raises(ValueError, match="Cookie Threads đã hết hạn hoặc chưa nạp .* vào /admin › Cookie Threads để nạp lại"):
        asyncio.run(scan_service.prepare_scan_start({
            "platform": "threads", "file_id": "Book.xlsx", "use_threads_session": True,
            "threads_session_generation": session.status()["generation"]}))
