"""Threads transport, bounded fallback and proxy routing regressions."""
import asyncio
import json
import threading
import time
import urllib.error
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

import openpyxl
import pytest
from playwright.async_api import async_playwright

from riviu import proxy_utils
from riviu.platforms import threads
from riviu.workbook_utils import worksheet_find_column_index

URL = "https://www.threads.com/@demo/post/Fixture"


def html_fixture(shares=2):
    payload = {"require": [
        ["BarcelonaLoggedOutExpansionGating", [], {"view_counts": 15}, 7623],
        ["RelayPrefetchedStreamCache", "next", [], ["adp_BarcelonaPostPageTargetQueryRelayPreloader_fixture",
         {"__bbox": {"result": {"data": {"media": {
             "code": "Fixture", "user": {"username": "demo"}, "like_count": 0,
             "like_and_view_counts_disabled": False,
             "text_post_app_info": {"direct_reply_count": 0, "repost_count": 0, "reshare_count": shares},
         }}}}}]],
    ]}
    return '<html><body><article><a href="/@demo/post/Fixture">demo</a><button><svg title="Like"></svg></button><button><svg title="Comment"></svg></button></article><script type="application/json">' + json.dumps(payload) + '</script></body></html>'


def metrics(**values):
    return {"channel": "demo", "metrics": {**threads.empty_metrics(), **values}, "error": ""}


class NoBrowser:
    async def __aenter__(self):
        return self

    async def __aexit__(self, *_args):
        pass


class Events:
    def __init__(self):
        self.logs = []
        self.data = []

    async def broadcast_log(self, message):
        self.logs.append(message)

    async def broadcast_data(self, data):
        self.data.append(data)

    async def broadcast_status(self, _data):
        pass

    async def broadcast_duplicates(self, _data):
        pass


def workbook(tmp_path, count=1, profile=False):
    path = tmp_path / "threads.xlsx"
    book = openpyxl.Workbook()
    book.active.title = "Data"
    book.active.append(["Link", "TIM"])
    for index in range(count):
        book.active.append([URL if count == 1 else URL + str(index)])
    if profile:
        book.active.append(["https://www.threads.com/@profile?xmt=sample", 99])
    book.save(path)
    book.close()
    return path


def test_http_sends_accept_and_explicit_route(monkeypatch):
    config = proxy_utils.normalize_proxy_config({"host": "proxy.example", "port": 8080})
    seen = []

    class Response:
        def __enter__(self): return self
        def __exit__(self, *_args): pass
        def read(self): return html_fixture().encode()
        def geturl(self): return URL

    def open_request(request, route, timeout):
        seen.append((request.get_header("Accept"), route, timeout))
        return Response()

    monkeypatch.setattr(proxy_utils, "urlopen_with_config", open_request)
    result = threads.fetch_threads_http(URL, proxy_config=config)
    assert result["metrics"] == {"views": 15, "likes": 0, "comments": 0, "reposts": 0, "shares": 2}
    assert seen == [(threads.HTTP_HEADERS["Accept"], config, 20)]


@pytest.mark.parametrize("code,kind", [(403, "access_denied"), (429, "rate_limit"), (407, "proxy_auth")])
def test_http_terminal_error_does_not_expose_credentials(monkeypatch, code, kind):
    def rejected(*_args, **_kwargs):
        raise urllib.error.HTTPError("http://private-user:private-pass@proxy.example", code, "private-pass", {}, None)
    monkeypatch.setattr(proxy_utils, "urlopen_with_config", rejected)
    result = threads.fetch_threads_http(URL)
    assert result["error_kind"] == kind
    assert "private" not in json.dumps(result)


def test_http_transport_error_is_credential_safe(monkeypatch):
    def failed(*_args, **_kwargs):
        raise urllib.error.URLError("http://user:secret@proxy.example")
    monkeypatch.setattr(proxy_utils, "urlopen_with_config", failed)
    result = threads.fetch_threads_http(URL, proxy_config={"enabled": True})
    assert result["error_kind"] == "transport"
    assert "secret" not in json.dumps(result)


def test_gate_and_caption_are_not_confused():
    gate = "<h1>This content isn't available to everyone</h1><p>It can't be seen by certain audiences.</p>"
    assert threads.parse_threads_http(gate, URL, URL)["error_kind"] == "audience"
    assert threads.parse_threads_http(html_fixture() + gate, URL, URL)["error"] == ""
    assert threads.parse_threads_http('<p>My caption says invalid_post</p>' + html_fixture(), URL, URL)["metrics"]["likes"] == 0
    assert threads.parse_threads_http("", URL, "https://www.threads.com/?error=invalid_post")["error_kind"] == "invalid_post"


@pytest.mark.parametrize("kind", ["audience", "invalid_post", "access_denied", "rate_limit", "proxy_auth"])
def test_hybrid_terminal_errors_skip_retry_and_browser(tmp_path, monkeypatch, kind):
    path = workbook(tmp_path)
    calls = []
    monkeypatch.setattr(threads, "async_playwright", NoBrowser)
    monkeypatch.setattr(threads, "fetch_threads_http", lambda url: calls.append(url) or threads.failure("fixture terminal", kind))
    asyncio.run(threads.run_threads_scraper(path, mode="hybrid"))
    assert calls == [URL]


def test_retry_and_hybrid_keep_same_local_proxy_and_http_values(tmp_path, monkeypatch):
    path = workbook(tmp_path)
    config = proxy_utils.normalize_proxy_config({"host": "proxy.example", "port": 8080})
    calls = []
    results = iter([metrics(views=15), metrics(views=99, likes=0)])
    monkeypatch.setattr(threads, "resolve_threads_proxies", lambda *_args: [config])
    monkeypatch.setattr(threads, "async_playwright", NoBrowser)
    sentinel_pool = proxy_utils.get_session_proxies()

    def http(url, proxy_config):
        calls.append(("http", proxy_config))
        return next(results)

    class Browser:
        async def close(self): calls.append(("close", None))
    class Chromium:
        async def launch(self, **kwargs):
            assert kwargs["proxy"] == {"server": "per-context"}
            return Browser()
    NoBrowser.chromium = Chromium()

    async def browser(_browser, url, proxy_config):
        calls.append(("browser", proxy_config))
        return metrics(views=100, likes=8, comments=0, reposts=0, shares=2)

    monkeypatch.setattr(threads, "fetch_threads_http", http)
    monkeypatch.setattr(threads, "fetch_threads_browser", browser)
    asyncio.run(threads.run_threads_scraper(path, mode="hybrid", use_proxy=True))
    assert [kind for kind, _ in calls] == ["http", "http", "browser", "close"]
    assert all(route is config for kind, route in calls if kind != "close")
    assert proxy_utils.get_session_proxies() == sentinel_pool
    saved = openpyxl.load_workbook(path)
    try:
        assert saved.active.cell(2, worksheet_find_column_index(saved.active, ["LƯỢT XEM"])).value == 15
        assert saved.active["B2"].value == 0
    finally: saved.close()


def test_request_retry_ceiling_profile_warning_and_lazy_browser(tmp_path, monkeypatch):
    path = workbook(tmp_path, profile=True)
    calls = []
    monkeypatch.setattr(threads, "async_playwright", NoBrowser)
    monkeypatch.setattr(threads, "fetch_threads_http", lambda url: calls.append(url) or metrics(views=15))
    events = Events()
    asyncio.run(threads.run_threads_scraper(path, events, mode="request"))
    assert len(calls) == 2
    assert any("link hồ sơ" in line and "dòng 3" in line for line in events.logs)
    saved = openpyxl.load_workbook(path)
    try: assert saved.active["B3"].value == 99
    finally: saved.close()


def test_malformed_threads_post_links_are_reported_counted_and_left_untouched(tmp_path, monkeypatch):
    path = tmp_path / "threads.xlsx"
    book = openpyxl.Workbook()
    book.active.title = "Data"
    book.active.append(["Link", "TIM"])
    book.active.append([URL])
    malformed = [
        "https://www.threads.com/@user/post/",
        "https://www.threads.com/@user/post/CODE/media?token=private-value",
        "threads.net/t/CODE",
    ]
    for link in malformed:
        book.active.append([link, 99])
    book.active.append(["https://www.threads.com/@profile", 99])
    book.active.append(["https://example.com/@user/post/", 99])
    book.save(path)
    book.close()
    calls = []
    monkeypatch.setattr(threads, "async_playwright", NoBrowser)
    monkeypatch.setattr(threads, "fetch_threads_http", lambda url: calls.append(url) or metrics(views=15, likes=4))
    events = Events()
    asyncio.run(threads.run_threads_scraper(path, events, mode="request", base_dir=str(tmp_path)))

    assert calls[0] == URL and set(calls) == {URL}
    for row in (3, 4, 5):
        assert any(f"dòng {row}:" in line and "sai định dạng" in line for line in events.logs), row
    assert any("dòng 6:" in line and "link hồ sơ" in line for line in events.logs)
    assert not any("dòng 7:" in line for line in events.logs)  # non-Threads hosts are not Threads rows
    assert not any("private-value" in line for line in events.logs)
    assert any("bỏ qua 4 link Threads không hợp lệ" in line for line in events.logs)
    history = json.loads((tmp_path / "data" / "threads_scan_history.json").read_text(encoding="utf-8"))["history"]
    assert history[0]["skippedProfiles"] == 1
    assert history[0]["skippedInvalidLinks"] == 3
    saved = openpyxl.load_workbook(path)
    try:
        sheet = saved.active
        assert sheet["B2"].value == 4
        assert [sheet.cell(row, 2).value for row in range(3, 8)] == [99] * 5
        views = worksheet_find_column_index(sheet, ["LƯỢT XEM"])
        assert [sheet.cell(row, views).value for row in range(3, 8)] == [None] * 5
    finally:
        saved.close()


def test_zero_with_missing_shares_counts_as_success_not_hidden(tmp_path, monkeypatch):
    path = workbook(tmp_path)
    monkeypatch.setattr(threads, "async_playwright", NoBrowser)
    monkeypatch.setattr(threads, "fetch_threads_http", lambda url: metrics(views=0, likes=0, comments=0, reposts=0))
    class Manager(Events):
        statuses = []
        async def broadcast_status(self, data): self.statuses.append(data)
    events = Manager()
    asyncio.run(threads.run_threads_scraper(path, events, mode="request"))
    assert events.data[0]["status"] == "Success"
    assert events.data[0]["shares"] is None
    assert events.data[0]["missingMetrics"] == ["CHIA SẺ"]
    assert events.statuses[-1]["success"] == 1
    assert events.statuses[-1]["hidden"] == events.statuses[-1]["error"] == 0


def test_http_success_does_not_launch_chromium(tmp_path, monkeypatch):
    path = workbook(tmp_path)
    monkeypatch.setattr(threads, "async_playwright", NoBrowser)
    monkeypatch.setattr(threads, "fetch_threads_http", lambda url: metrics(views=15, likes=0, comments=0, reposts=0, shares=2))
    asyncio.run(threads.run_threads_scraper(path, mode="hybrid"))


def test_browser_retains_navigation_json_removed_by_page():
    html = html_fixture() + '<script>document.querySelector("script[type=\\"application/json\\"]").remove();</script>'
    async def check():
        async with async_playwright() as p:
            browser = await p.chromium.launch(headless=True)
            original = browser.new_page
            async def page_factory():
                page = await original()
                await page.route("**/*", lambda route: route.fulfill(body=html, content_type="text/html"))
                return page
            browser.new_page = page_factory
            try: return await threads.fetch_threads_browser(browser, URL)
            finally: await browser.close()
    result = asyncio.run(check())
    assert result["metrics"] == {"views": 15, "likes": 0, "comments": 0, "reposts": 0, "shares": 2}


def test_real_http_and_chromium_use_local_proxy():
    received = []
    body = html_fixture().encode()
    class Proxy(BaseHTTPRequestHandler):
        def do_GET(self):
            received.append((self.path, self.headers.get("User-Agent")))
            self.send_response(200)
            self.send_header("Content-Type", "text/html; charset=utf-8")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)
        def log_message(self, *_args): pass
    server = ThreadingHTTPServer(("127.0.0.1", 0), Proxy)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    config = proxy_utils.normalize_proxy_config({"host": "127.0.0.1", "port": server.server_port})
    url = URL.replace("https:", "http:")
    async def check():
        async with async_playwright() as p:
            browser = await p.chromium.launch(headless=True, proxy={"server": "per-context"})
            try:
                result = await threads.fetch_threads_browser(browser, url, proxy_config=config)
                assert browser.contexts == []
                return result
            finally: await browser.close()
    try:
        http = threads.fetch_threads_http(url, proxy_config=config)
        browser = asyncio.run(check())
        assert http["metrics"] == browser["metrics"] == {"views": 15, "likes": 0, "comments": 0, "reposts": 0, "shares": 2}
        relevant = [agent for path, agent in received if "/@demo/post/Fixture" in path]
        assert "Mozilla/5.0" in relevant
        assert any("Chrome" in agent for agent in relevant)
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=5)


def test_proxy_assignment_is_run_local_round_robin(tmp_path, monkeypatch):
    path = workbook(tmp_path, count=4)
    configs = [proxy_utils.normalize_proxy_config({"host": host, "port": 8080}) for host in ("a.example", "b.example")]
    routes = {}
    monkeypatch.setattr(threads, "resolve_threads_proxies", lambda *_args: configs)
    monkeypatch.setattr(threads, "async_playwright", NoBrowser)
    def request(url, proxy_config):
        routes[url] = proxy_config
        return metrics(views=15, likes=0, comments=0, reposts=0, shares=None)
    monkeypatch.setattr(threads, "fetch_threads_http", request)
    asyncio.run(threads.run_threads_scraper(path, mode="request", use_proxy=True))
    assert [routes[URL + str(i)] for i in range(4)] == configs * 2


def test_retry_rate_limit_never_starts_browser(tmp_path, monkeypatch):
    path = workbook(tmp_path)
    values = iter([metrics(views=15), threads.failure("HTTP 429", "rate_limit")])
    monkeypatch.setattr(threads, "async_playwright", NoBrowser)
    monkeypatch.setattr(threads, "fetch_threads_http", lambda url: next(values))
    asyncio.run(threads.run_threads_scraper(path, mode="hybrid"))


def test_browser_polls_parse_each_snapshot_once_off_the_event_loop(monkeypatch):
    # Missing views keeps the poll running until the deadline on an unchanged page.
    html = html_fixture().replace('"view_counts": 15', '"view_counts": null')
    parsed_on = []
    original = threads.parse_threads_http
    def counted(*args, **kwargs):
        parsed_on.append(threading.current_thread() is threading.main_thread())
        return original(*args, **kwargs)
    monkeypatch.setattr(threads, "parse_threads_http", counted)
    monkeypatch.setattr(threads, "BROWSER_POLL_SECONDS", 0.6, raising=False)
    polls = []
    class Response:
        status = 200
        async def text(self): return html
    class Page:
        url = URL
        async def goto(self, *_args, **_kwargs): return Response()
        async def content(self):
            polls.append(True)
            return html
        async def evaluate(self, *_args): return {}
        async def close(self): pass
    class Browser:
        async def new_page(self): return Page()
    result = asyncio.run(threads.fetch_threads_browser(Browser(), URL))
    assert result["metrics"]["likes"] == 0 and result["metrics"]["views"] is None
    assert len(polls) > 1
    # The navigation response plus one rendered snapshot, never once per poll.
    assert len(parsed_on) == 2
    assert not any(parsed_on)


def test_session_redirect_off_threads_is_a_redirect_result(monkeypatch):
    def blocked(request, config, timeout, cookiejar):
        return proxy_utils.SessionRedirectHandler().redirect_request(request, None, 302, "Found", {}, "https://example.com/")
    monkeypatch.setattr(proxy_utils, "urlopen_with_config", blocked)
    result = threads.fetch_threads_http(URL, session_cookies=[{"name": "sessionid", "value": "synthetic", "domain": ".threads.com", "path": "/"}])
    assert result["error_kind"] == "redirect"


def test_browser_errors_preserve_navigation_and_close_owned_context():
    closed = []
    class Response:
        status = 200
        async def text(self): return html_fixture()
    class Page:
        url = URL
        async def goto(self, *_args, **_kwargs): return Response()
        async def content(self): raise RuntimeError("secret proxy credential")
    class Context:
        async def new_page(self): return Page()
        async def close(self): closed.append(True)
    class Browser:
        async def new_context(self, **kwargs):
            assert kwargs["proxy"]["password"] == "secret"
            return Context()
    config = proxy_utils.normalize_proxy_config({"host": "proxy.example", "port": 8080, "username": "user", "password": "secret"})
    result = asyncio.run(threads.fetch_threads_browser(Browser(), URL, proxy_config=config))
    assert result["metrics"]["likes"] == 0
    assert result["metrics"]["shares"] == 2
    assert "secret" not in json.dumps(result)
    assert closed == [True]


def test_browser_launch_failure_keeps_http_metrics(tmp_path, monkeypatch):
    path = workbook(tmp_path)
    class Chromium:
        async def launch(self, **kwargs): raise RuntimeError("Chromium missing")
    class Playwright(NoBrowser):
        chromium = Chromium()
    monkeypatch.setattr(threads, "async_playwright", Playwright)
    # A missing core metric (REPOST) sends the URL to Chromium; shares alone would not.
    monkeypatch.setattr(threads, "fetch_threads_http", lambda url: metrics(views=15, likes=0, comments=0))
    events = Events()
    asyncio.run(threads.run_threads_scraper(path, events, mode="hybrid"))
    assert events.data[0]["views"] == 15
    assert events.data[0]["status"] == "Success"
    assert events.data[0]["missingMetrics"] == ["REPOST", "CHIA SẺ"]
    assert any("Không khởi chạy được Chromium" in line for line in events.logs)


def test_invalid_proxy_fails_before_network(tmp_path, monkeypatch):
    path = workbook(tmp_path)
    def forbidden(*_args, **_kwargs): raise AssertionError("network must not run")
    monkeypatch.setattr(threads, "fetch_threads_http", forbidden)
    with pytest.raises(ValueError, match="proxy"):
        asyncio.run(threads.run_threads_scraper(path, mode="hybrid", use_proxy=True, proxy_text="malformed"))


def test_cancel_waits_for_owned_http_thread(tmp_path, monkeypatch):
    path = workbook(tmp_path)
    entered = threading.Event()
    release = threading.Event()
    finished = threading.Event()
    def request(url):
        entered.set()
        assert release.wait(5)
        finished.set()
        return metrics(views=15)
    monkeypatch.setattr(threads, "fetch_threads_http", request)
    monkeypatch.setattr(threads, "async_playwright", NoBrowser)
    async def check():
        task = asyncio.create_task(threads.run_threads_scraper(path, mode="request"))
        assert await asyncio.to_thread(entered.wait, 5)
        task.cancel()
        await asyncio.sleep(0.05)
        assert not task.done()
        release.set()
        with pytest.raises(asyncio.CancelledError): await task
        assert finished.is_set()
    try: asyncio.run(check())
    finally: release.set()


class LaunchableChromium:
    """A browser that launches: a Hybrid fallback really reaches fetch_threads_browser."""

    def __call__(self):
        return self

    async def __aenter__(self):
        return self

    async def __aexit__(self, *_args):
        pass

    @property
    def chromium(self):
        return self

    async def launch(self, **_kwargs):
        return self

    async def close(self):
        pass


def test_hybrid_shares_only_gap_stays_on_http(tmp_path, monkeypatch):
    # Live (2026-10-08): when HTTP has reshare_count null, the rendered page has it null
    # too and the Share button shows no number; 0 of 4 fallbacks ever added shares.
    path = workbook(tmp_path)
    calls = []
    monkeypatch.setattr(threads, "async_playwright", LaunchableChromium())
    monkeypatch.setattr(threads, "fetch_threads_http", lambda url: calls.append("http") or metrics(views=15, likes=0, comments=0, reposts=0))

    async def browser(*_args, **_kwargs):
        calls.append("browser")
        return metrics(views=15, likes=0, comments=0, reposts=0)

    monkeypatch.setattr(threads, "fetch_threads_browser", browser)
    events = Events()
    asyncio.run(threads.run_threads_scraper(path, events, mode="hybrid"))
    assert calls == ["http"]
    assert events.data[0]["worker"] == "Request"
    assert events.data[0]["status"] == "Success"
    assert events.data[0]["missingMetrics"] == ["CHIA SẺ"]


def test_slow_hybrid_fallbacks_leave_http_slots_draining(tmp_path, monkeypatch):
    # Live profile: 4 slow fallbacks held 4 of 5 slots and HTTP ran one at a time.
    path = workbook(tmp_path, count=12)
    gaps = {URL + str(index) for index in range(4)}  # sheet order: the fallbacks come first
    lock = threading.Lock()
    flight = {"gap": 0, "plain": 0, "browser": 0}
    peak = {"plain_beside_browser": 0, "browser": 0, "total": 0}
    http_done = threading.Event()
    plain_left = [12 - len(gaps)]

    def move(kind, delta):
        with lock:
            flight[kind] += delta
            if flight["browser"]:
                peak["plain_beside_browser"] = max(peak["plain_beside_browser"], flight["plain"])
            peak["browser"] = max(peak["browser"], flight["browser"])
            peak["total"] = max(peak["total"], sum(flight.values()))
            if kind == "plain" and delta < 0:
                plain_left[0] -= 1
                if not plain_left[0]:
                    http_done.set()

    def http(url):
        kind = "gap" if url in gaps else "plain"
        move(kind, 1)
        try:
            if kind == "gap":
                return threads.failure("HTTP 500", "http")  # not terminal, no retry: needs the browser
            time.sleep(0.1)
            return metrics(views=1, likes=1, comments=1, reposts=1, shares=1)
        finally:
            move(kind, -1)

    async def browser(_browser, url):
        move("browser", 1)
        try:
            # Slow Chromium: it waits until every plain HTTP URL is done (or gives up).
            await asyncio.to_thread(http_done.wait, 10)
            return metrics(likes=2, comments=0, reposts=0)
        finally:
            move("browser", -1)

    monkeypatch.setattr(threads, "async_playwright", LaunchableChromium())
    monkeypatch.setattr(threads, "fetch_threads_http", http)
    monkeypatch.setattr(threads, "fetch_threads_browser", browser)
    events = Events()
    started = time.monotonic()
    asyncio.run(threads.run_threads_scraper(path, events, worker_count=5, mode="hybrid"))
    assert http_done.is_set()
    assert peak["plain_beside_browser"] == 3  # HTTP kept the other slots while the browsers waited
    assert peak["browser"] == 2  # BROWSER_FALLBACK_LIMIT
    assert peak["total"] <= 5  # one route never carries more than the worker count
    assert time.monotonic() - started < 5
    assert sorted(item["worker"] for item in events.data) == ["Hybrid"] * 4 + ["Request"] * 8
