"""Report text safety, stable snapshots, transport guards and durable preferences."""
import io
import json
from pathlib import Path
import urllib.request
import zipfile

import openpyxl
import pytest
from fastapi.testclient import TestClient

import app as backend
import proxy_utils
import reports
import threads_scraper as threads
from test_backend_review_fixes import isolated_backend


def test_report_untrusted_strings_are_literal_excel_text():
    row = {'NGÀY AIR':'=1+1','TÊN KÊNH':'=HYPERLINK("https://example.invalid","x")',
           'LINK AIR':'=1+1','LƯỢT XEM':100,'TIM':0,'BÌNH LUẬN':0,'REPOST':0,'CHIA SẺ':''}
    payload = reports.build_partner_report('Fixture', [row], platform='threads')
    book = openpyxl.load_workbook(io.BytesIO(payload))
    try:
        sheet = book.active
        for address in ('A4','B4','C4'):
            assert sheet[address].data_type == 's'
            assert sheet[address].value.startswith('=')
        assert sheet['D4'].value == 100
    finally: book.close()


def test_multi_partner_export_uses_single_snapshot(tmp_path, monkeypatch):
    path = tmp_path / 'source.xlsx'
    def write(views):
        book = openpyxl.Workbook()
        book.active.title='Data'
        book.active.append(['Link','Tên Kênh','LƯỢT XEM','TIM','Đối tác','Đối tác 2'])
        book.active.append(['https://www.threads.com/@demo/post/Fixture','demo',views,0,'A','B'])
        book.save(path); book.close()
    write(10)
    original = reports.build_partner_report_rows
    def read(snapshot, **kwargs):
        # A scan replaces the workbook after the partner list was read, before the rows are.
        write(999)
        return original(snapshot, **kwargs)
    monkeypatch.setattr(reports,'build_partner_report_rows',read)
    result = reports.build_export_payload(path,['A','B'],False,0,'Data','threads')
    with zipfile.ZipFile(io.BytesIO(result['content'])) as archive:
        values=[]
        for name in archive.namelist():
            book=openpyxl.load_workbook(io.BytesIO(archive.read(name)))
            values.append(book.active['D4'].value);book.close()
    assert values == [10,10]


def test_threads_exact_unknown_clears_stale_views_and_channel_is_literal():
    book=openpyxl.Workbook();sheet=book.active
    sheet.append(['Link','Tên Kênh','LƯỢT XEM','TIM']);sheet.append(['https://www.threads.com/@demo/post/Test','demo',15100,5])
    result={'channel':'=1+1','metrics':{'views':None,'likes':None},'error':'','metricSources':{'views':'view_query_unknown'}}
    threads.write_threads_result(sheet,2,result)
    assert sheet['C2'].value is None and sheet['D2'].value==5
    assert sheet['B2'].value=='=1+1' and sheet['B2'].data_type=='s'
    book.close()


def test_summary_unicode_sources_do_not_collide():
    from workbook_utils import summary_sheet_title_for_data_sheet,data_sheet_name_for_summary_title
    first=summary_sheet_title_for_data_sheet('Straße');second=summary_sheet_title_for_data_sheet('STRASSE')
    assert first.lower()!=second.lower()
    assert data_sheet_name_for_summary_title(['Straße','STRASSE'],first)=='Straße'
    assert data_sheet_name_for_summary_title(['Straße','STRASSE'],second)=='STRASSE'


def test_preference_gid_fragment_round_trips():
    result=backend.normalize_source_preferences({'sources':{'threads':{'url':'https://docs.google.com/spreadsheets/d/ABC/edit#gid=42'}}})
    assert result['sources']['threads']['url'].endswith('?gid=42')


def test_tiktok_all_disabled_proxy_is_rejected(tmp_path):
    disabled=json.dumps({'enabled':False,'host':'proxy.example','port':8080})
    assert backend.validate_proxy_start(True,disabled,str(tmp_path)) is not None


def test_generic_redirect_validator_checks_before_following():
    validator=lambda url: url.startswith('https://www.tiktok.com/')
    guard=proxy_utils.ValidatedRedirectHandler(validator)
    req=urllib.request.Request('https://www.tiktok.com/@demo/video/123')
    with pytest.raises(ValueError):guard.redirect_request(req,None,302,'Found',{},'http://127.0.0.1/?tiktok.com')
    assert guard.redirect_request(req,None,302,'Found',{},'https://www.tiktok.com/@demo/video/456').full_url.endswith('456')
    with pytest.raises(ValueError):proxy_utils.urlopen_with_config(urllib.request.Request('http://127.0.0.1/?tiktok.com'),None,redirect_validator=validator)


def test_later_exact_unknown_invalidates_rounded_views_but_not_exact():
    display={'channel':'demo','metrics':{'views':15100,'likes':0},'error':'','metricSources':{'views':'header_display'}}
    unknown={'channel':'demo','metrics':{'views':None},'error':'','metricSources':{'views':'view_query_unknown'}}
    result=threads.merge_results(display,unknown)
    assert result['metrics']['views'] is None and result['metricSources']['views']=='view_query_unknown'
    exact={**display,'metricSources':{'views':'view_query_exact'}}
    assert threads.merge_results(exact,unknown)['metrics']['views']==15100


@pytest.fixture
def preference_backend(isolated_backend):
    return isolated_backend


def test_preferences_persist_across_ports_and_ignore_secrets(preference_backend):
    payload={'sources':{'tiktok':{'fileId':'data/A.xlsx','displaySheet':'Data','scanSheet':'Data','pushSheet':'Data','url':'https://docs.google.com/spreadsheets/d/ABC/edit?cookie=private&gid=12'},'threads':{}},'cookie':'private'}
    with TestClient(backend.app,base_url='http://127.0.0.1:1231') as client:
        client.get('/')
        response=client.post('/source-preferences',json=payload)
        assert response.status_code==200
    with TestClient(backend.app,base_url='http://127.0.0.1:4567') as client:
        client.get('/')
        response=client.get('/source-preferences')
        assert response.headers['cache-control']=='no-store'
        result=response.json()
        assert result['sources']['tiktok']['fileId']=='data/A.xlsx'
        assert result['sources']['tiktok']['url'].endswith('gid=12')
        assert 'private' not in response.text
    raw=(preference_backend/'data/source_preferences.json').read_text(encoding='utf-8')
    assert 'private' not in raw


@pytest.mark.parametrize('bad', ['../secret.xlsx','C:/secret.xlsx','/secret.xlsx','data//secret.xlsx'])
def test_preferences_reject_unsafe_file_ids(preference_backend,bad):
    with TestClient(backend.app,base_url='http://127.0.0.1:1231') as client:
        client.get('/')
        response=client.post('/source-preferences',json={'sources':{'threads':{'fileId':bad}}})
        assert response.status_code==400


def test_preferences_busy_and_foreign_origin_rejected(preference_backend,monkeypatch):
    with TestClient(backend.app,base_url='http://127.0.0.1:1231') as client:
        client.get('/')
        assert client.post('/source-preferences',json={'sources':{}},headers={'Origin':'https://foreign.example'}).status_code==403
        monkeypatch.setattr(backend,'SOURCE_BUSY',True)
        assert client.post('/source-preferences',json={'sources':{}}).status_code==409


def test_preference_write_failure_preserves_previous(preference_backend,monkeypatch):
    before={'sources':{'tiktok':{'fileId':'A.xlsx'},'threads':{}}}
    backend.save_source_preferences(before)
    path=preference_backend/'data/source_preferences.json'
    old=path.read_bytes()
    def failed(*_args):raise PermissionError('locked')
    monkeypatch.setattr(backend.os,'replace',failed)
    with pytest.raises(PermissionError):backend.save_source_preferences({'sources':{}})
    assert path.read_bytes()==old
    assert list(path.parent.glob('*.json')) == [path.parent/'google_sheet_sources.json',path] or set(p.name for p in path.parent.glob('*.json'))=={'google_sheet_sources.json','source_preferences.json'}
