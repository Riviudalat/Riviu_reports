import asyncio
import io
import threading
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from starlette.datastructures import UploadFile
from starlette.websockets import WebSocketDisconnect

from riviu import app as backend
from riviu.workbook_utils import google_sheet_source_for_file, register_google_sheet_source


def make_workbook(path, marker="original"):
    book = Workbook()
    sheet = book.active
    sheet.title = "Data"
    sheet.append(["Link", "Tên Kênh", "Đối tác", "Marker"])
    sheet.append(["https://www.tiktok.com/@demo/video/123", "Demo", "Partner", marker])
    book.save(path)
    book.close()


@pytest.fixture
def isolated_backend(tmp_path, monkeypatch):
    monkeypatch.setattr(backend, "EXCEL_DIR", str(tmp_path))
    monkeypatch.setattr(backend, "CURRENT_SELECTED_FILE", "A.xlsx")
    monkeypatch.setattr(backend, "CURRENT_SELECTED_SHEET", "Data")
    monkeypatch.setattr(backend, "CURRENT_SCAN_SHEET", "Data")
    monkeypatch.setattr(backend, "SCRAPE_TASK", None)
    monkeypatch.setattr(backend, "manager", backend.ConnectionManager())
    monkeypatch.setattr(backend, "oauth_status", lambda _: {"valid": False, "configured": False})
    for name in ("A", "B"):
        make_workbook(tmp_path / f"{name}.xlsx")
        register_google_sheet_source(str(tmp_path), f"{name}.xlsx", f"https://docs.google.com/spreadsheets/d/ID_{name}/edit")
    return tmp_path


def test_select_and_push_uses_workbook_source_without_rebinding_override(isolated_backend, monkeypatch):
    pushed = []
    monkeypatch.setattr(backend, "push_rows_to_new_sheet", lambda base, sid, rows, **kw: pushed.append(sid) or "Result")
    asyncio.run(backend.select_file({"filename": "B.xlsx"}))
    state = asyncio.run(backend.list_files())
    assert state["googleSheetUrl"].endswith("/ID_B/edit")
    asyncio.run(backend.push_google_sheet({"sourceSheet": "Data"}))
    asyncio.run(backend.push_google_sheet({"sourceSheet": "Data", "url": "https://docs.google.com/spreadsheets/d/EXPLICIT/edit"}))
    assert pushed == ["ID_B", "EXPLICIT"]
    assert google_sheet_source_for_file(str(isolated_backend), "B.xlsx")["spreadsheetId"] == "ID_B"


def test_invalid_upload_never_destroys_existing_workbook(isolated_backend):
    path = isolated_backend / "A.xlsx"
    before = path.read_bytes()
    response = asyncio.run(backend.upload_excel(UploadFile(filename="A.xlsx", file=io.BytesIO(b"not excel"))))
    assert not response.get("success")
    assert path.read_bytes() == before
    assert backend.CURRENT_SELECTED_FILE == "A.xlsx"


def test_valid_same_name_upload_gets_new_filename(isolated_backend):
    path = isolated_backend / "A.xlsx"
    before = path.read_bytes()
    new = isolated_backend / "new.xlsx"
    make_workbook(new, "replacement")
    response = asyncio.run(backend.upload_excel(UploadFile(filename="A.xlsx", file=io.BytesIO(new.read_bytes()))))
    assert response["success"]
    assert response["filename"] != "A.xlsx"
    assert path.read_bytes() == before
    saved = load_workbook(isolated_backend / response["filename"])
    assert saved.active["D2"].value == "replacement"
    saved.close()


def test_sync_same_title_same_minute_creates_distinct_files(isolated_backend, monkeypatch):
    monkeypatch.setattr(backend, "fetch_google_spreadsheet_title", lambda _: "Same Name")
    monkeypatch.setattr(backend, "format_filename_datetime", lambda: "25-09-2026-10-30")
    monkeypatch.setattr(backend, "sync_download_google_sheet", lambda base, url, sid, dest: make_workbook(dest, sid))
    first = asyncio.run(backend.sync_google_sheet({"url": "https://docs.google.com/spreadsheets/d/SOURCE_A/edit"}))
    first_bytes = (isolated_backend / first["file"]).read_bytes()
    second = asyncio.run(backend.sync_google_sheet({"url": "https://docs.google.com/spreadsheets/d/SOURCE_B/edit"}))
    assert first["file"] != second["file"]
    assert (isolated_backend / first["file"]).read_bytes() == first_bytes


def test_rejects_foreign_origin_websocket_and_http(isolated_backend):
    with TestClient(backend.app, base_url="http://127.0.0.1:1231") as client:
        assert client.get("/").status_code == 200
        with pytest.raises(WebSocketDisconnect):
            with client.websocket_connect("ws://127.0.0.1:1231/ws", headers={"Origin": "https://untrusted.example"}):
                pass
        response = client.post("/select-file", headers={"Origin": "https://untrusted.example"}, json={"filename": "B.xlsx"})
        assert response.status_code == 403
        assert backend.CURRENT_SELECTED_FILE == "A.xlsx"
        assert client.get("/list-files", headers={"Host": "rebound.example:1231"}).status_code == 403


def test_local_origin_requires_session_and_gets_run_snapshot(isolated_backend):
    with TestClient(backend.app, base_url="http://127.0.0.1:1231") as client:
        assert client.get("/list-files").status_code == 403
        assert client.get("/").status_code == 200
        context = backend.manager.begin_run("threads", "A.xlsx", "Data")
        with client.websocket_connect("ws://127.0.0.1:1231/ws", headers={"Origin": "http://127.0.0.1:1231"}) as ws:
            snapshot = ws.receive_json()
        assert snapshot["type"] == "session"
        assert snapshot["data"]["platform"] == "threads"
        assert snapshot["data"]["runId"] == context["runId"]
        assert snapshot["data"]["running"]


def test_no_source_mutation_during_active_scan(isolated_backend):
    before = (isolated_backend / "A.xlsx").read_bytes()
    backend.manager.begin_run("threads", "A.xlsx", "Data")
    selected = asyncio.run(backend.select_file({"filename": "B.xlsx"}))
    assert selected.status_code == 409
    uploaded = asyncio.run(backend.upload_excel(UploadFile(filename="A.xlsx", file=io.BytesIO(b"bad"))))
    assert not uploaded["success"]
    assert backend.CURRENT_SELECTED_FILE == "A.xlsx"
    assert (isolated_backend / "A.xlsx").read_bytes() == before


def test_snapshot_metadata_and_terminal_failure(isolated_backend, monkeypatch):
    class Connection:
        def __init__(self): self.messages = []
        async def send_json(self, message): self.messages.append(message)
    connection = Connection()
    backend.manager.active_connections.append(connection)
    async def fail_after_partial(path, events, **kwargs):
        await events.broadcast_data({"id": 1, "views": None, "saves": None})
        await events.broadcast_status({"total": 1, "processed": 1, "success": 0, "hidden": 1, "error": 0, "done": False})
        raise RuntimeError("fixture save failed")
    monkeypatch.setattr(backend, "run_threads_scraper", fail_after_partial)
    asyncio.run(backend.run_scraper_safely(str(isolated_backend / "A.xlsx"), backend.scan_options(platform="threads", sheet_name="Data")))
    snapshot = backend.manager.snapshot()
    assert not snapshot["running"]
    assert snapshot["status"]["phase"] == "failed"
    assert snapshot["results"][0]["platform"] == "threads"
    assert snapshot["results"][0]["views"] is None
    assert all(message["runId"] == snapshot["runId"] for message in connection.messages)
    assert not any(message.get("data", {}).get("phase") == "completed" for message in connection.messages)


def test_tiktok_legacy_phase_is_normalized_to_running(isolated_backend):
    backend.manager.begin_run("tiktok", "A.xlsx", "Data")
    events = backend.RunEvents(backend.manager)
    asyncio.run(events.broadcast_status({"phase": "scanning", "total": 2, "processed": 1, "done": False}))
    assert backend.manager.snapshot()["status"]["phase"] == "running"
    assert backend.manager.running


def test_cancel_before_runner_starts_releases_run(isolated_backend):
    async def scenario():
        context = backend.manager.begin_run("threads", "A.xlsx", "Data")
        task = asyncio.create_task(backend.run_scraper_safely(str(isolated_backend / "A.xlsx"), backend.scan_options(platform="threads"), run_context=context))
        task.add_done_callback(lambda finished: backend.finalize_unstarted_run(finished, context["runId"]))
        task.cancel()
        await asyncio.gather(task, return_exceptions=True)
        await asyncio.sleep(0)
        assert not backend.manager.running
        assert backend.manager.last_status["phase"] == "cancelled"
    asyncio.run(scenario())


def test_cancel_blocking_write_waits_for_actual_finish():
    entered, release, finished = threading.Event(), threading.Event(), threading.Event()
    def write():
        entered.set()
        release.wait(2)
        finished.set()
    async def scenario():
        task = asyncio.create_task(backend.complete_blocking(write))
        await asyncio.to_thread(entered.wait, 1)
        task.cancel()
        await asyncio.sleep(.02)
        assert not task.done()
        release.set()
        with pytest.raises(asyncio.CancelledError):
            await task
        assert finished.is_set()
    asyncio.run(scenario())


def test_slow_preview_remains_bound_to_original_file(isolated_backend, monkeypatch):
    entered, release = threading.Event(), threading.Event()
    def delayed_preview(path, **kwargs):
        entered.set()
        assert release.wait(2)
        return {"currentSheet": "A Sheet", "columns": [], "data": []}
    monkeypatch.setattr(backend, "read_sheet_preview", delayed_preview)
    async def scenario():
        task = asyncio.create_task(backend.preview_excel("Data"))
        await asyncio.to_thread(entered.wait, 1)
        backend.CURRENT_SELECTED_FILE = "B.xlsx"
        backend.CURRENT_SELECTED_SHEET = "B Sheet"
        release.set()
        response = await task
        assert response["file"] == "A.xlsx"
        assert backend.CURRENT_SELECTED_SHEET == "B Sheet"
    asyncio.run(scenario())


def test_oauth_does_not_block_other_coroutines(isolated_backend, monkeypatch):
    entered, release = threading.Event(), threading.Event()
    def waiting_auth(_):
        entered.set()
        assert release.wait(2)
    monkeypatch.setattr(backend, "authorize_google", waiting_auth)

    async def scenario():
        task = asyncio.create_task(backend.google_oauth_login())
        try:
            await asyncio.wait_for(asyncio.to_thread(entered.wait, 1), 1.5)
            assert not task.done(), "OAuth blocked the event loop"
            # An abandoned login tab must not hold the Google I/O lock used by previews, pushes and scans.
            assert backend.GOOGLE_IO_LOCK.acquire(timeout=0.2), "OAuth wait holds GOOGLE_IO_LOCK"
            backend.GOOGLE_IO_LOCK.release()
            assert (await backend.google_oauth_login()).status_code == 409
        finally:
            release.set()
        assert (await task)["success"]
    asyncio.run(scenario())


def test_scan_target_stays_bound_if_selection_changes(isolated_backend, monkeypatch):
    observed = []
    async def fake_scan(path, manager, **kwargs):
        observed.append(Path(path).name)
        backend.CURRENT_SELECTED_FILE = "B.xlsx"
        await manager.broadcast_status({"total": 1, "processed": 1, "success": 1, "error": 0, "done": True})
    monkeypatch.setattr(backend, "run_scraper", fake_scan)
    monkeypatch.setattr(backend, "push_rows_to_new_sheet", lambda base, sid, rows, **kw: observed.append(sid) or "Result")
    asyncio.run(backend.run_scraper_safely(str(isolated_backend / "A.xlsx"), backend.scan_options(sheet_name="Data", push_to_google=True, file_id="A.xlsx")))
    assert observed == ["A.xlsx", "ID_A"]
    assert backend.manager.last_status["phase"] == "completed"


@pytest.mark.parametrize("failure", [RuntimeError("quota exceeded"), backend.GoogleLoginRequired()])
def test_push_failure_after_saved_scan_is_a_warning_not_a_failed_scan(isolated_backend, monkeypatch, failure):
    async def fake_scan(path, manager, **kwargs):
        await manager.broadcast_status({"total": 1, "processed": 1, "success": 1, "error": 0, "done": True})
    def failing_push(*_args, **_kwargs):
        raise failure
    monkeypatch.setattr(backend, "run_scraper", fake_scan)
    monkeypatch.setattr(backend, "push_rows_to_new_sheet", failing_push)
    logs = []
    original_log = backend.manager.broadcast_log
    async def capture(message, **kwargs):
        logs.append(message)
        await original_log(message, **kwargs)
    monkeypatch.setattr(backend.manager, "broadcast_log", capture)
    asyncio.run(backend.run_scraper_safely(str(isolated_backend / "A.xlsx"), backend.scan_options(
        sheet_name="Data", push_to_google=True)))
    assert backend.manager.last_status["phase"] == "completed"
    assert any("đã lưu xong" in message for message in logs)


def test_websocket_start_scans_the_workbook_the_client_names(isolated_backend, monkeypatch):
    calls = []
    async def fake_runner(path, options, **kwargs):
        calls.append((Path(path).name, options["file_id"]))
    monkeypatch.setattr(backend, "run_scraper_safely", fake_runner)
    backend.CURRENT_SELECTED_FILE = "A.xlsx"
    asyncio.run(run_fixture_websocket_start(monkeypatch, {"platform": "tiktok", "file_id": "B.xlsx"}))
    assert calls == [("B.xlsx", "B.xlsx")]


def test_broadcast_reaches_every_tab_when_one_disconnects_mid_send():
    manager = backend.ConnectionManager()
    received = []
    class Leaving:
        async def send_json(self, message):
            manager.disconnect(self)
            raise RuntimeError("closed")
    class Staying:
        async def send_json(self, message):
            received.append(message["type"])
    manager.active_connections.extend([Leaving(), Staying()])
    asyncio.run(manager.broadcast_log("hello"))
    assert received == ["log"]
    assert len(manager.active_connections) == 1


@pytest.mark.parametrize("mode", ["request", "browser", "hybrid"])
@pytest.mark.parametrize("use_proxy", [False, True])
def test_threads_dispatch_forwards_proxy_settings(isolated_backend, monkeypatch, mode, use_proxy):
    calls = []
    async def fake_threads(path, events, **kwargs):
        calls.append((path, kwargs))
    monkeypatch.setattr(backend, "run_threads_scraper", fake_threads)
    text = "http://fixture-user:fixture-password@proxy.example:8080"
    target = str(isolated_backend / "A.xlsx")
    asyncio.run(backend.run_scraper_safely(target, backend.scan_options(
        worker_count=3, platform="threads", sheet_name="Data", partners=["Partner"],
        scrape_mode=mode, use_proxy=use_proxy, proxy_text=text,
    )))
    assert calls == [(target, {
        "worker_count": 3, "selected_partners": ["Partner"], "sheet_name": "Data", "mode": mode,
        "base_dir": str(isolated_backend), "use_proxy": use_proxy, "proxy_text": text,
    })]
    assert backend.manager.last_status["phase"] == "completed"


def save_fixture_proxy_list(base_dir, text):
    path = base_dir / "data" / "proxy_list.txt"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")


@pytest.mark.parametrize("inline,saved,valid", [
    ("", "proxy.example:8080", True),
    ("proxy.example:8080", "", True),
    ("", "", False),
    ("", '{"enabled":false,"host":"proxy.example","port":8080}', False),
    ('{"enabled":false,"host":"proxy.example","port":8080}', "proxy.example:8080", False),
    ("malformed-fixture-password", "proxy.example:8080", False),
    ("proxy.example:8080\nmalformed-fixture-password", "", False),
    ("", "malformed-fixture-password", False),
    ("http://fixture-user:fixture-password@proxy.example:bad-port", "proxy.example:8080", False),
])
def test_threads_proxy_validation_strict_inline_or_saved(tmp_path, inline, saved, valid):
    save_fixture_proxy_list(tmp_path, saved)
    error = backend.validate_proxy_start(True, inline, str(tmp_path), platform="threads")
    assert (error is None) == valid
    if error:
        assert "fixture-user" not in error
        assert "fixture-password" not in error
        assert "proxy.example" not in error


@pytest.mark.parametrize("mode,valid", [("request", True), ("browser", False), ("hybrid", False)])
@pytest.mark.parametrize("saved", [False, True])
def test_threads_authenticated_socks_proxy_mode_validation(tmp_path, mode, valid, saved):
    text = "socks5://fixture-user:fixture-password@proxy.example:1080"
    save_fixture_proxy_list(tmp_path, text if saved else "")
    error = backend.validate_proxy_start(True, "" if saved else text, str(tmp_path), platform="threads", mode=mode)
    assert (error is None) == valid
    if error:
        assert "fixture-user" not in error
        assert "fixture-password" not in error


def test_proxy_validation_off_skips_threads_resolver(tmp_path, monkeypatch):
    def unexpected(*args):
        pytest.fail("Proxy-off validation must not resolve proxy configuration")
    monkeypatch.setattr(backend, "resolve_threads_proxies", unexpected)
    assert backend.validate_proxy_start(False, "malformed", str(tmp_path), platform="threads", mode="hybrid") is None


def test_threads_proxy_validation_returns_safe_resolver_error(tmp_path, monkeypatch):
    calls = []
    def reject(base_dir, proxy_text, mode):
        calls.append((base_dir, proxy_text, mode))
        raise ValueError("Fixture safe proxy validation error")
    monkeypatch.setattr(backend, "resolve_threads_proxies", reject)
    error = backend.validate_proxy_start(True, "fixture input", str(tmp_path), platform="threads", mode="hybrid")
    assert error == "Fixture safe proxy validation error"
    assert calls == [(str(tmp_path), "fixture input", "hybrid")]


def test_proxy_validation_preserves_tiktok_defaults_and_fallback(tmp_path):
    assert backend.validate_proxy_start(False, "", str(tmp_path)) is None
    assert backend.validate_proxy_start(True, "", str(tmp_path))
    save_fixture_proxy_list(tmp_path, "proxy.example:8080")
    # TikTok intentionally retains its existing lenient inline parsing and saved fallback.
    assert backend.validate_proxy_start(True, "malformed", str(tmp_path)) is None
    assert backend.validate_proxy_start(True, "proxy.example:8080\nmalformed", str(tmp_path)) is None
    assert backend.validate_proxy_start(True, "socks5://fixture-user:fixture-password@proxy.example:1080",
                                        str(tmp_path), platform="tiktok", mode="hybrid") is None


async def run_fixture_websocket_start(monkeypatch, payload):
    class FixtureSocket:
        def __init__(self):
            self.messages = []
            self.received = False
        async def accept(self):
            pass
        async def send_json(self, message):
            self.messages.append(message)
        async def receive_text(self):
            if not self.received:
                self.received = True
                return backend.json.dumps({"action": "start", "sheet_name": "Data", **payload})
            await asyncio.sleep(0)
            raise WebSocketDisconnect()
    monkeypatch.setattr(backend, "trusted_local_request", lambda *args, **kwargs: True)
    monkeypatch.setattr(backend, "valid_local_session", lambda *args: True)
    socket = FixtureSocket()
    await backend.websocket_endpoint(socket)
    if backend.SCRAPE_TASK:
        await backend.SCRAPE_TASK
    return socket.messages


@pytest.mark.parametrize("platform", ["threads", "tiktok"])
@pytest.mark.parametrize("input_mode,expected_mode", [
    ("request", "request"), ("browser", "browser"), ("hybrid", "hybrid"), ("invalid", "request"),
])
def test_websocket_start_passes_platform_and_normalized_mode(isolated_backend, monkeypatch, platform, input_mode, expected_mode):
    validation_calls, dispatch_calls = [], []
    def validate(use_proxy, proxy_text, base_dir, platform="tiktok", mode="request"):
        validation_calls.append((use_proxy, proxy_text, base_dir, platform, mode))
    async def fake_runner(path, options, **kwargs):
        dispatch_calls.append(options)
    monkeypatch.setattr(backend, "validate_proxy_start", validate)
    monkeypatch.setattr(backend, "run_scraper_safely", fake_runner)
    asyncio.run(run_fixture_websocket_start(monkeypatch, {
        "platform": platform, "scrape_mode": input_mode, "use_proxy": True, "proxy_text": "proxy.example:8080",
    }))
    assert validation_calls == [(True, "proxy.example:8080", str(isolated_backend), platform, expected_mode)]
    assert len(dispatch_calls) == 1
    assert dispatch_calls[0]["platform"] == platform
    assert dispatch_calls[0]["use_proxy"] is True
    assert dispatch_calls[0]["proxy_text"] == "proxy.example:8080"
    assert dispatch_calls[0]["scrape_mode"] == expected_mode


@pytest.mark.parametrize("text,mode", [
    ("malformed-fixture-password", "request"),
    ('{"enabled":false,"host":"proxy.example","port":8080}', "request"),
    ("socks5://fixture-user:fixture-password@proxy.example:1080", "hybrid"),
])
def test_websocket_rejects_invalid_threads_proxy_before_dispatch(isolated_backend, monkeypatch, text, mode):
    calls = []
    async def fake_runner(*args, **kwargs):
        calls.append(kwargs)
    monkeypatch.setattr(backend, "run_scraper_safely", fake_runner)
    save_fixture_proxy_list(isolated_backend, "proxy.example:8080")
    messages = asyncio.run(run_fixture_websocket_start(monkeypatch, {
        "platform": "threads", "scrape_mode": mode, "use_proxy": True, "proxy_text": text,
    }))
    assert not calls
    assert backend.SCRAPE_TASK is None
    assert not backend.manager.running
    assert any(message.get("data", {}).get("phase") == "failed" for message in messages)
    output = backend.json.dumps(messages)
    assert "fixture-user" not in output
    assert "fixture-password" not in output
